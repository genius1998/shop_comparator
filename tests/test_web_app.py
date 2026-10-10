"""Web API, login handling, logging, recovery and exports (mock data only)."""
import csv
import io
import json
import re
import threading
import time
import unittest
from unittest import mock

from openpyxl import load_workbook

from shop_comparator.core import app_logging as L
from shop_comparator.services import comparison_service as cmp
from shop_comparator.services import exports
from shop_comparator.services import monitoring
from shop_comparator.shops import site_login
from shop_comparator.storage import store
from shop_comparator.core.crawl_errors import CrawlError
from helpers import TempDbTestCase, option, options_payload, seed_product
from shop_comparator.services.jobs import JobManager
from shop_comparator.web.web_app import create_app, validate_settings

SECRET = "S3cret-Pass!word"


class ApiCase(TempDbTestCase):
    def setUp(self):
        super().setUp()
        self.jobs = JobManager(persist=True)
        self.app = create_app(self.jobs)
        self.client = self.app.test_client()
        self.token = re.search(r'name="csrf-token" content="([^"]+)"', self.client.get("/").text)[1]

    def post(self, path, body=None):
        return self.client.post(path, json=body or {}, headers={"X-CSRF-Token": self.token})

    def finish(self, response):
        self.assertEqual(response.status_code, 202, response.get_json())
        return self.jobs.wait(response.get_json()["job"], 5)


class LoginTests(ApiCase):
    def test_each_shop_requires_its_own_account_and_password(self):
        for site in ("musinsa", "kream", "poizon"):
            self.assertEqual(self.post(f"/api/login/{site}", {"account": "", "password": ""}).status_code, 400)
            self.assertEqual(self.post(f"/api/login/{site}", {"account": "user@example.com", "password": ""}).status_code, 400)
        self.assertEqual(self.post("/api/login/other", {"account": "a", "password": "b"}).status_code, 400)

    def test_credentials_reach_the_login_function_for_that_shop_only(self):
        seen = {}
        def fake(site):
            def login(account, password, cancel=None, wait_seconds=300, remember=False):
                seen[site] = (account, password, remember)
                store.save_session(site, "logged_in", L.mask_account(account), "인증 확인", verified=True)
            return login
        with mock.patch.object(site_login, "login_musinsa", fake("musinsa")), mock.patch.object(site_login, "login_kream", fake("kream")):
            job = self.finish(self.post("/api/login/musinsa", {"account": "me@example.com", "password": SECRET, "remember": True}))
        self.assertEqual(job["state"], "done")
        self.assertEqual(seen, {"musinsa": ("me@example.com", SECRET, True)})
        sessions = self.client.get("/api/state").get_json()["sessions"]
        self.assertEqual(sessions["musinsa"]["state"], "logged_in")
        self.assertEqual(sessions["musinsa"]["account"], "me***@example.com")
        self.assertEqual(sessions["kream"]["state"], "unknown")  # the other shop is untouched

    def test_failed_login_of_one_shop_leaves_the_other_logged_in(self):
        store.save_session("musinsa", "logged_in", "me***@example.com", "ok", verified=True)
        def fail(account, password, cancel=None, wait_seconds=300, remember=False):
            store.save_session("kream", "failed", None, "이메일 또는 비밀번호를 확인해주세요")
            raise CrawlError("KREAM", "auth", "KREAM 로그인 실패: 이메일 또는 비밀번호를 확인해주세요")
        with mock.patch.object(site_login, "login_kream", fail):
            job = self.finish(self.post("/api/login/kream", {"account": "me@example.com", "password": SECRET}))
        self.assertEqual((job["state"], job["errorKind"]), ("failed", "auth"))
        sessions = self.client.get("/api/state").get_json()["sessions"]
        self.assertEqual((sessions["musinsa"]["state"], sessions["kream"]["state"]), ("logged_in", "failed"))

    def test_password_never_appears_in_state_logs_or_database(self):
        def leaky(account, password, cancel=None, wait_seconds=300, remember=False):
            L.register_secret(password)
            L.log(f"debug: typed password={password} for {account}", shop="musinsa")
            L.log(f"Authorization: Bearer abcdefghijklmnop.qrstuvwxyz and {password}")
            raise CrawlError("무신사", "auth", f"로그인 실패 ({password})")
        with mock.patch.object(site_login, "login_musinsa", leaky):
            job = self.finish(self.post("/api/login/musinsa", {"account": "me@example.com", "password": SECRET}))
        exposed = json.dumps([self.client.get("/api/state").get_json(), self.client.get("/api/logs").get_json(),
                              self.client.get("/api/jobs/history").get_json()], ensure_ascii=False)
        self.assertNotIn(SECRET, exposed)
        self.assertNotIn("me@example.com", exposed)
        self.assertNotIn("abcdefghijklmnop", exposed)
        self.assertNotIn(SECRET.encode(), self.db_path.read_bytes())
        self.assertEqual(job["state"], "failed")

    def test_env_login_uses_env_values_without_exposing_them(self):
        seen = []
        env = {"KREAM_ID": "env-user@example.com", "KREAM_PASSWORD": SECRET}
        with mock.patch.dict("os.environ", env), mock.patch("shop_comparator.matching.product_matcher.load_local_env"), \
             mock.patch.object(site_login, "login_kream", lambda a, p, cancel=None, wait_seconds=300, remember=False: seen.append((a, p))):
            self.finish(self.post("/api/login/kream", {"env": True}))
            state = self.client.get("/api/state").get_json()
            self.assertEqual(state["envAccounts"]["kream"], True)
            self.assertEqual(self.post("/api/login/musinsa", {"env": True}).status_code, 400)  # not set for this shop
        self.assertEqual(seen, [("env-user@example.com", SECRET)])
        self.assertNotIn(SECRET, json.dumps(state))
        self.assertNotIn("env-user", json.dumps(state))

    def test_manual_login_needs_no_credentials(self):
        seen = []
        with mock.patch.object(site_login, "login_kream", lambda a, p, cancel=None, wait_seconds=300, remember=False: seen.append((a, p))):
            self.finish(self.post("/api/login/kream", {"manual": True}))
        self.assertEqual(seen, [("", "")])

    def test_session_validity_needs_an_authenticated_answer(self):
        def session(status, body="{}"):
            response = mock.Mock(status_code=status)
            response.json.side_effect = lambda: json.loads(body)
            return mock.Mock(get=mock.Mock(return_value=response), observed={"authorization": "Bearer x"})
        self.assertFalse(site_login.verify_musinsa(session(401))[0])   # cookie present but API refuses
        self.assertTrue(site_login.verify_musinsa(session(200))[0])
        with self.assertRaises(CrawlError):
            site_login.verify_musinsa(session(503))                      # outage is not "expired"
        self.assertFalse(site_login.verify_kream(session(401))[0])
        self.assertFalse(site_login.verify_kream(session(200, "{}"))[0])  # 200 without sales data is not proof
        self.assertTrue(site_login.verify_kream(session(200, '{"items": [], "total": 0}'))[0])
        anonymous = session(200, '{"items": []}')
        anonymous.observed = {}
        self.assertFalse(site_login.verify_kream(anonymous)[0])            # no member token at all
        anonymous.get.assert_not_called()

    def test_remembered_credentials_live_in_memory_only_and_can_be_dropped(self):
        site_login.vault.put("kream", "me@example.com", SECRET)
        self.addCleanup(site_login.vault.drop, "kream")
        self.assertTrue(self.client.get("/api/state").get_json()["sessions"]["kream"]["remembered"])
        self.assertNotIn(SECRET.encode(), self.db_path.read_bytes())
        self.assertEqual(L.redact(f"x {SECRET} y"), "x *** y")
        self.post("/api/session/kream/forget")
        self.assertFalse(site_login.vault.has("kream"))

    def test_stale_login_state_from_a_previous_process_is_not_shown_as_in_progress(self):
        store.save_session("musinsa", "needs_user", None, "열린 Chrome 창에서 추가 인증을 완료해 주세요.")
        self.assertEqual(site_login.status()["musinsa"]["state"], "unknown")


class JobTests(ApiCase):
    def test_shops_run_independently_and_stop_only_their_own_job(self):
        gate = threading.Event()
        started = {"musinsa": threading.Event(), "kream": threading.Event()}
        def work(shop):
            def run(ctx):
                started[shop].set()
                while not gate.is_set():
                    ctx.wait(0.05)
            return run
        musinsa, _ = self.jobs.start("musinsa", "무신사 수집", ["musinsa"], work("musinsa"))
        kream, conflict = self.jobs.start("kream", "KREAM 가격 수집", ["kream"], work("kream"))
        self.assertTrue(kream and not conflict)
        self.assertTrue(started["musinsa"].wait(1) and started["kream"].wait(1))
        blocked, conflict = self.jobs.start("match", "KREAM 매칭", ["kream"], work("kream"))
        self.assertIsNone(blocked)
        self.assertEqual(conflict, "KREAM 가격 수집")
        self.assertEqual(self.jobs.start("full", "전체", ["musinsa", "kream"], work("musinsa"))[0], None)
        self.post("/api/stop", {"id": kream})
        self.assertEqual(self.jobs.wait(kream, 3)["state"], "cancelled")
        self.assertEqual(self.jobs.get(musinsa)["state"], "running")
        gate.set()
        self.assertEqual(self.jobs.wait(musinsa, 3)["state"], "done")

    def test_job_can_be_run_again_after_a_stop(self):
        def run(ctx):
            ctx.wait(5)
        first, _ = self.jobs.start("musinsa", "무신사 수집", ["musinsa"], run)
        time.sleep(0.05)
        self.jobs.stop(first)
        self.assertEqual(self.jobs.wait(first, 3)["state"], "cancelled")
        second, _ = self.jobs.start("musinsa", "무신사 수집", ["musinsa"], lambda ctx: ctx.count("ok"))
        self.assertEqual(self.jobs.wait(second, 3)["state"], "done")

    def test_failure_kind_and_job_id_are_traceable_in_logs(self):
        def run(ctx):
            with L.bind(shop="kream", goods="42", stage="kream-price"):
                L.log("가격 조회")
            raise CrawlError("KREAM", "rate_limit", "요청 한도에 도달했습니다 (429).")
        job_id, _ = self.jobs.start("kream", "KREAM 가격 수집", ["kream"], run)
        job = self.jobs.wait(job_id, 3)
        self.assertEqual((job["state"], job["errorKind"]), ("failed", "rate_limit"))
        events = self.client.get(f"/api/logs?job={job_id}").get_json()["events"]
        traced = [e for e in events if e["goods"] == "42"][0]
        self.assertEqual((traced["job"], traced["shop"], traced["stage"]), (job_id, "kream", "kream-price"))
        failure = [e for e in events if e["error"] == "rate_limit"][0]
        self.assertEqual((failure["level"], failure["shop"]), ("error", "kream"))
        self.assertTrue(all(e["ts"] for e in events))
        self.assertEqual(self.client.get("/api/logs?job=nope").get_json()["events"], [])

    def test_jobs_left_running_by_a_dead_server_are_marked_interrupted(self):
        store.save_job({"id": "old", "kind": "musinsa", "title": "무신사 수집", "state": "running", "startedAt": store.now(), "summary": {}})
        store.interrupt_stale_jobs()
        self.assertEqual(store.recent_jobs()[0]["state"], "interrupted")

    def test_stage_without_prerequisite_fails_with_explanation(self):
        response = self.post("/api/jobs/compare")          # nothing collected yet
        self.assertEqual(response.status_code, 409)
        self.assertIn("무신사만 수집", response.get_json()["error"])
        job = self.finish(self.post("/api/jobs/compare", {"scope": "all"}))
        self.assertEqual((job["state"], job["errorKind"]), ("failed", "prerequisite"))
        self.assertIn("무신사 수집", job["error"])

    def test_follow_up_jobs_default_to_the_last_collected_batch_only(self):
        from shop_comparator.services import pipeline
        for goods_no in ("1", "2", "3"):
            seed_product(goods_no, f"상품 {goods_no}")
        store.set_setting("last_collection", {"goods": ["2", "3"], "at": store.now()})
        seen = []
        with mock.patch.object(pipeline, "match_kream", lambda ids, ctx, force=False: seen.append(("match", ids))), \
             mock.patch.object(pipeline, "collect_kream", lambda ids, ctx, match_missing=False: seen.append(("kream", ids)) or self.assertTrue(match_missing)):
            self.finish(self.post("/api/jobs/match"))                       # default: last batch
            self.finish(self.post("/api/jobs/kream", {"ids": ["1"]}))       # explicit selection wins
            self.finish(self.post("/api/jobs/match", {"scope": "all"}))     # whole list only when asked
        self.assertEqual(seen, [("match", ["2", "3"]), ("kream", ["1"]), ("match", None)])
        shown = self.client.get("/api/products").get_json()                 # default view = last batch only
        self.assertEqual(([p["goodsNo"] for p in shown["products"]], shown["total"]), (["2", "3"], 3))
        self.assertEqual(len(self.client.get("/api/products?scope=all").get_json()["products"]), 3)
        self.assertEqual(self.client.get("/api/state").get_json()["lastCollection"]["count"], 2)


class DataTests(ApiCase):
    def seed(self):
        seed_product("1", "러너", sizes=(("10", "250", 90000), ("11", "260", 95000), ("12", "270", None)))
        store.set_manual_match("1", "500")
        store.save_kream_quotes("500", [
            {"size": "250", "instantBuy": 100000, "instantBuyDelivery": "일반배송", "instantSell": 80000},
            {"size": "260", "instantBuy": None, "bidOnly": True, "instantSell": 85000},
            {"size": "280", "instantBuy": 120000, "instantSell": None}])

    def test_results_survive_a_server_restart(self):
        self.seed()
        fresh = create_app(JobManager(persist=False)).test_client()   # new process, same database
        product = fresh.get("/api/products").get_json()["products"][0]
        self.assertEqual(product["match"]["status"], "직접 지정")
        sizes = {s["size"] or "K" + s["kreamSize"]: s for s in product["sizes"]}
        self.assertEqual(set(sizes), {"250", "260", "270", "K280"})
        self.assertEqual((sizes["250"]["musinsaPrice"], sizes["250"]["kreamBuy"], sizes["250"]["difference"]), (90000, 100000, 10000))
        self.assertEqual(sizes["250"]["sellDifference"], -10000)
        self.assertEqual(sizes["260"]["status"], "즉시구매 불가(입찰만)")
        self.assertIsNone(sizes["260"]["kreamBuy"])
        self.assertEqual(sizes["260"]["kreamSell"], 85000)   # sell side kept separate from buy side
        self.assertEqual(sizes["270"]["status"], "품절")
        self.assertEqual(sizes["K280"]["status"], "무신사에 없는 사이즈")
        self.assertTrue(all(s["musinsaCheckedAt"] or s["kreamCheckedAt"] for s in product["sizes"]))

    def test_hiding_a_product_keeps_its_data_and_can_be_undone(self):
        self.seed()
        self.post("/api/products/active", {"ids": ["1"], "active": False})
        self.assertEqual(self.client.get("/api/products").get_json()["products"], [])
        self.assertEqual(len(store.variants("1")), 3)
        self.assertEqual(len(self.client.get("/api/products?hidden=1").get_json()["products"]), 1)
        self.post("/api/products/active", {"ids": ["1"], "active": True})
        self.assertEqual(len(self.client.get("/api/products").get_json()["products"]), 1)

    def test_manual_match_approve_and_clear(self):
        seed_product("1")
        from shop_comparator.storage import database
        database.cache_kream_search("1", "600", "https://kream.co.kr/products/600", match_verified=False)
        row = self.client.get("/api/matches").get_json()["matches"][0]
        self.assertEqual((row["status"], row["usable"]), ("미검증", False))
        self.post("/api/matches/approve", {"ids": ["1"]})
        self.assertTrue(self.client.get("/api/matches").get_json()["matches"][0]["usable"])
        response = self.post("/api/matches", {"goodsNo": "1", "kream": "https://kream.co.kr/products/74130"})
        self.assertEqual(response.status_code, 200)
        row = self.client.get("/api/matches").get_json()["matches"][0]
        self.assertEqual((row["kreamId"], row["status"], row["confidence"]), ("74130", "직접 지정", 1.0))
        self.assertEqual(self.post("/api/matches", {"goodsNo": "1", "kream": "https://evil.example/products/1"}).status_code, 400)
        self.post("/api/matches/clear", {"ids": ["1"]})
        self.assertEqual(self.client.get("/api/matches").get_json()["matches"][0]["status"], "매칭 없음")

    def test_downloads_contain_every_size_with_source_and_time(self):
        self.seed()
        sheet = load_workbook(io.BytesIO(self.client.get("/api/export/comparison.xlsx").data)).active
        header = [c.value for c in sheet[1]]
        rows = [dict(zip(header, [c.value for c in row])) for row in sheet.iter_rows(min_row=2)]
        self.assertEqual(len(rows), 4)
        first = next(r for r in rows if r["무신사 사이즈(원본)"] == "250")
        self.assertEqual((first["무신사 회원가"], first["KREAM 즉시구매 최저가"], first["가격 차이 (KREAM 즉시구매가-무신사)"]), (90000, 100000, 10000))
        self.assertEqual((first["회원가 출처"], first["상품 매칭"], first["정규화 사이즈"]), ("주문서(해당 옵션)", "직접 지정", "MM:250"))
        self.assertTrue(first["무신사 조회 시각(UTC)"] and first["KREAM 조회 시각(UTC)"])
        bid_only = next(r for r in rows if r["무신사 사이즈(원본)"] == "260")
        self.assertIsNone(bid_only["KREAM 즉시구매 최저가"])
        text = self.client.get("/api/export/comparison.csv").data.decode("utf-8-sig")
        self.assertEqual(len(list(csv.reader(io.StringIO(text)))), 5)
        gui = list(csv.DictReader(io.StringIO(self.client.get("/api/export/gui.csv?source=kream").data.decode("utf-8-sig"))))
        self.assertEqual(gui[0]["price_diff"], "-10000")   # GUI sign: 무신사 − KREAM
        self.assertEqual(gui[0]["kream_link"], "https://kream.co.kr/products/500")

    def test_price_check_download_has_one_row_per_shared_size_and_highlights(self):
        self.assertEqual(self.client.get("/api/export/price-check.xlsx").status_code, 409)       # nothing collected yet
        self.seed()
        seed_product("2", "부츠", sizes=(("20", "250", 100000),))
        store.set_manual_match("2", "600")
        store.save_kream_quotes("600", [{"size": "250", "instantBuy": 104000, "instantBuyDelivery": "일반배송", "instantSell": 90000}])
        store.set_setting("last_collection", {"goods": ["1", "2"], "at": store.now()})
        from shop_comparator.storage import database
        conn = database.get_db_conn()
        conn.execute("UPDATE products SET kream_monthly_transactions = 30 WHERE goodsNo IN ('1', '2')")
        conn.commit()
        book = load_workbook(io.BytesIO(self.client.get("/api/export/price-check.xlsx").data))
        sheet = book.worksheets[0]
        self.assertEqual(sheet.title, "사이즈별 가격 비교")
        header = [c.value for c in sheet[1]]
        rows = {(r[header.index("무신사 상품ID")].value, r[header.index("무신사 사이즈")].value): r for r in sheet.iter_rows(min_row=2)}
        # sizes Musinsa priced AND KREAM lists: 270 has no Musinsa price, 280 exists only on KREAM
        self.assertEqual(set(rows), {("1", "250"), ("1", "260"), ("2", "250")})

        def cell(key, name):
            return rows[key][header.index(name)]

        self.assertEqual((cell(("1", "250"), "무신사 회원가").value, cell(("1", "250"), "KREAM 즉시구매가").value), (90000, 100000))
        self.assertAlmostEqual(cell(("1", "250"), "KREAM÷무신사").value, 100000 / 90000)
        self.assertTrue(cell(("1", "250"), "KREAM 즉시구매가").fill.fgColor.rgb.endswith("F4511E"))   # 111%: strong
        self.assertTrue(cell(("2", "250"), "KREAM 즉시구매가").fill.fgColor.rgb.endswith("FFD966"))   # 104%: soft
        self.assertIsNone(cell(("1", "260"), "KREAM 즉시구매가").value)                              # bid only: no guessed price
        self.assertIsNone(cell(("1", "260"), "KREAM 즉시구매가").fill.fill_type)
        only_two = load_workbook(io.BytesIO(self.client.get("/api/export/price-check.xlsx?ids=2").data)).worksheets[0]
        self.assertEqual(only_two.max_row, 2)

    def test_bid_sheet_uses_each_sizes_own_price_and_only_matched_sizes(self):
        self.seed()
        rows = exports.bid_rows(10.0, 2, "30일")
        self.assertEqual([(r["옵션명"], r["판매 희망가"]) for r in rows], [("250", 99000), ("260", 105000)])  # ceil to 1,000
        self.assertTrue(all(r["상품번호"] == "500" and r["수량"] == "2" and r["입찰기한"] == "30일" for r in rows))
        self.assertEqual(self.client.get("/api/export/bid.xlsx?margin=abc").status_code, 400)
        self.assertEqual(self.client.get("/api/export/bid.xlsx?deadline=내일").status_code, 400)

    def test_csv_export_neutralises_formulas(self):
        seed_product("1", "=HYPERLINK(\"http://x\")")
        text = self.client.get("/api/export/comparison.csv").data.decode("utf-8-sig")
        self.assertIn("'=HYPERLINK", text)

    def test_webhook_is_validated_and_masked(self):
        self.assertEqual(self.post("/api/settings", {"webhook_url": "https://example.com/api/webhooks/1/abc"}).status_code, 400)
        url = "https://discord.com/api/webhooks/123456/abcdefSECRETTOKEN"
        self.assertEqual(self.post("/api/settings", {"webhook_url": url, "alert_margin": "5000"}).status_code, 200)
        data = self.client.get("/api/settings").get_json()
        self.assertTrue(data["webhookSet"])
        self.assertNotIn("SECRETTOKEN"[:6], json.dumps(data))
        self.assertEqual(data["alertMargin"], 5000)
        self.assertNotIn("abcdefSECRET", L.redact(f"sent to {url}"))

    def test_collection_settings_validation(self):
        self.assertEqual(validate_settings({"mode": "ranking", "rankings": ["shoes_ranking", "x"], "limit": "100"})["rankings"], ["shoes_ranking"])
        self.assertEqual(validate_settings({"mode": "category", "categories": ["103", "001"], "limit": 5})["categories"], ["103", "001"])
        for bad in ({"mode": "ranking", "rankings": []}, {"mode": "category", "categories": ["999"]}, {"mode": "products", "products": "1", "limit": 0},
                    {"mode": "products", "products": "1", "price_mode": "guess"}, {"mode": "weird"}, {"mode": "selected", "ids": []}):
            with self.assertRaises(ValueError):
                validate_settings(bad)


class SizeAndCompareTests(unittest.TestCase):
    def test_size_keys_never_cross_systems(self):
        self.assertEqual(cmp.size_key("250"), cmp.size_key("250mm"))
        self.assertEqual(cmp.size_key("W250"), cmp.size_key("250"))     # KREAM gender prefix on a mm size
        self.assertEqual(cmp.size_key("227.5"), "MM:227.5")
        self.assertEqual(cmp.size_key("2XL"), cmp.size_key("XXL"))
        self.assertEqual(cmp.size_key("free"), cmp.size_key("ONE SIZE"))
        self.assertNotEqual(cmp.size_key("US 8"), cmp.size_key("UK 8"))
        self.assertNotEqual(cmp.size_key("US 8"), cmp.size_key("260"))  # no US -> mm conversion
        self.assertNotEqual(cmp.size_key("8"), cmp.size_key("US 8"))    # bare number has no system
        self.assertNotEqual(cmp.size_key("M"), cmp.size_key("L"))
        self.assertEqual(cmp.size_key("L(100)"), "=:L(100)")            # unknown format: exact match only

    def test_no_first_size_price_is_copied(self):
        match = {"kream_product_id": "5", "match_verified": 1, "match_method": "style_code", "match_confidence": 0.95}
        variants = [{"size": "250", "optionNo": "1", "finalPrice": 90000, "priceStatus": "verified"},
                    {"size": "260", "optionNo": "2", "finalPrice": None, "priceStatus": "error", "error": "주문서 오류"},
                    {"size": "270", "optionNo": "3", "priceStatus": "not_requested"}]
        quotes = [{"size": s, "instantBuy": 100000, "instantSell": None} for s in ("250", "260", "270")]
        rows = {r["size"]: r for r in cmp.compare_product({"goodsNo": 1}, variants, match, quotes)}
        self.assertEqual(rows["250"]["difference"], 10000)
        self.assertIsNone(rows["260"]["musinsaPrice"])
        self.assertIsNone(rows["260"]["difference"])
        self.assertEqual(rows["260"]["status"], "회원가 확인 실패")
        self.assertEqual(rows["270"]["status"], "회원가 미조회")
        self.assertEqual((rows["250"]["matchStatus"], rows["250"]["matchConfidence"]), ("자동 검증", 0.95))

    def test_ambiguous_kream_sizes_are_not_matched(self):
        match = {"kream_product_id": "5", "match_verified": 1, "match_method": "manual"}
        quotes = [{"size": "250", "instantBuy": 100000}, {"size": "W250", "instantBuy": 90000}]
        row = cmp.compare_product({"goodsNo": 1}, [{"size": "250", "finalPrice": 80000, "priceStatus": "verified"}], match, quotes)[0]
        self.assertEqual(row["status"], "KREAM 사이즈 중복")
        self.assertIsNone(row["difference"])


class MonitorTests(unittest.TestCase):
    def stock(self, **sizes):
        return options_payload(*[option(i, size, soldout=not available) for i, (size, available) in enumerate(sizes.items(), 1)])

    def test_restock_detection(self):
        from shop_comparator.shops import crawler
        with mock.patch.object(crawler, "fetch_options_any", return_value=self.stock(**{"250": False, "260": False})):
            status, state, _, event = monitoring.restock_check(None, "1", {})
        self.assertEqual((status, event), ("품절", None))                    # first look never alerts
        with mock.patch.object(crawler, "fetch_options_any", return_value=self.stock(**{"250": True, "260": False})):
            status, state, flag, event = monitoring.restock_check(None, "1", state)
        self.assertEqual((status, flag, event), ("재입고", "good", "restock"))
        with mock.patch.object(crawler, "fetch_options_any", return_value=self.stock(**{"250": True, "260": True})):
            status, state, _, event = monitoring.restock_check(None, "1", state)
        self.assertEqual((status, event), ("새 입고: 260", "restock"))
        self.assertEqual(state["text"], "재고있음: 250, 260")
        with mock.patch.object(crawler, "fetch_options_any", return_value=self.stock(**{"250": True, "260": True})):
            self.assertIsNone(monitoring.restock_check(None, "1", state)[3])   # unchanged: no repeat alert

    def test_price_drop_detection_uses_member_price(self):
        from shop_comparator.shops import crawler
        with mock.patch.object(crawler, "fetch_options_any", return_value=self.stock(**{"250": True})), \
             mock.patch.object(crawler, "price_option", side_effect=[{"finalPrice": 90000}, {"finalPrice": 90000}, {"finalPrice": 85000}]):
            _, state, _, event = monitoring.price_check(None, {}, "1", {})
            self.assertEqual((state["original"], event), (90000, None))
            self.assertEqual(monitoring.price_check(None, {}, "1", state)[0], "변동 없음")
            status, state, flag, event = monitoring.price_check(None, {}, "1", state)
        self.assertEqual((status, state["current"], event), ("가격 변동 감지", 85000, "price"))

    def test_purchase_limit_release_detection(self):
        from shop_comparator.shops import crawler
        with mock.patch.object(crawler, "fetch_purchase_limit_info", side_effect=[{"isLimitedQuantity": True}, {"isLimitedQuantity": False}]):
            status, state, _, event, _ = monitoring.limit_check(None, "1", {})
            self.assertEqual((status, event), ("구매 제한", None))
            status, state, _, event, _ = monitoring.limit_check(None, "1", state)
        self.assertEqual((status, event), ("구매제한 해제!", "limit"))

    def test_auto_event_interval_rule(self):
        self.assertEqual(monitoring.AutoEvent.interval("5", 300), 30)     # GUI minimum
        self.assertEqual(monitoring.AutoEvent.interval("abc", 300), 300)
        self.assertEqual(monitoring.AutoEvent.interval("600", 300), 600)


class LoggingTests(unittest.TestCase):
    def test_redaction_patterns(self):
        text = L.redact('cookie: app_atk=abc123def456; password="hunter2"; Authorization: Bearer abcdefgh12345678 user someone@example.com')
        for leaked in ("abc123def456", "hunter2", "abcdefgh12345678", "someone@example.com"):
            self.assertNotIn(leaked, text)

    def test_rotation_policy_is_configured(self):
        self.assertEqual((L.MAX_BYTES, L.BACKUP_COUNT), (5 * 1024 * 1024, 5))
        self.assertTrue(L.LOG_FILE.endswith("app.jsonl"))


if __name__ == "__main__":
    unittest.main()

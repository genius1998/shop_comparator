"""Capture the real screens of the demo app (and of the portfolio page) with Playwright.

    python scripts/portfolio/build_demo.py          # once: isolated demo database
    python scripts/portfolio/capture.py app         # app screens  -> docs/portfolio/assets/*.png
    python scripts/portfolio/capture.py page        # the portfolio page itself (desktop, mobile, click check)

Uses the installed Google Chrome (playwright channel "chrome").  The app runs on
the demo folder through scripts/portfolio/demo_server.py: stored real data is
replayed, no shop and no API is contacted, and the failure shown in the
recovery screens is injected for the test.  Every image gets an entry in
docs/portfolio/assets/captures.json (what it shows, where the data came from).
"""
import html
import io
import json
import os
import subprocess
import sys
import time
import urllib.request

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DEMO = os.path.join(ROOT, "portfolio_demo")
ASSETS = os.path.join(ROOT, "docs", "portfolio", "assets")
PORT = 8799
BASE = f"http://127.0.0.1:{PORT}"
REPLAY = "저장된 실데이터 재생·네트워크 미호출"
captures = []


def note(name, title, feature, data, detail="", kind="replay"):
    captures.append({"file": f"assets/{name}.png", "title": title, "feature": feature, "dataSource": data, "kind": kind,
                     "detail": detail, "capturedAt": time.strftime("%Y-%m-%dT%H:%M:%S%z")})
    print("captured", name)


def wait_idle(page, timeout=120):
    end = time.time() + timeout
    while time.time() < end:
        jobs = page.evaluate("fetch('/api/state').then(r => r.json()).then(s => s.jobs)")
        if not any(j["state"] in ("running", "stopping") for j in jobs):
            return jobs
        time.sleep(0.4)
    raise RuntimeError("job did not finish")


def start_server(fail):
    process = subprocess.Popen([sys.executable, os.path.join(ROOT, "scripts", "portfolio", "demo_server.py"), "--dir", DEMO,
                                "--port", str(PORT), "--fail", ",".join(fail)], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    for _ in range(60):
        try:
            urllib.request.urlopen(BASE + "/api/state", timeout=2)
            return process
        except OSError:
            time.sleep(0.5)
    process.terminate()
    raise RuntimeError("demo server did not start")


def xlsx_preview(data, limit=26):
    """HTML rendering of the first sheet of a real export (cells, fills, thumbnails), for a readable screenshot."""
    import base64
    from openpyxl import load_workbook
    sheet = load_workbook(io.BytesIO(data)).worksheets[0]
    images = {}
    for image in sheet._images:
        anchor = image.anchor._from
        images[(anchor.row + 1, anchor.col + 1)] = base64.b64encode(image._data()).decode("ascii")
    rows = []
    for r, row in enumerate(sheet.iter_rows(min_row=1, max_row=limit), 1):
        cells = []
        for cell in [c for c in row[:17] if c.column != 6]:      # the brand column is left out so the price columns fit
            style = ""
            if cell.fill and cell.fill.fill_type == "solid":
                rgb = cell.fill.fgColor.rgb[-6:]
                style = f"background:#{rgb};" + ("color:#fff;font-weight:700;" if rgb in ("163A35", "F4511E") else "")
            value = cell.value
            if (r, cell.column) in images:
                content = f'<img src="data:image/png;base64,{images[(r, cell.column)]}" alt="">'
            elif isinstance(value, float) and 0 < value < 5:
                content = f"{value * 100:.1f}%"
            elif isinstance(value, (int, float)) and cell.column >= 10:
                content = f"{value:,.0f}"
            else:
                content = html.escape("" if value is None else str(value))
            cells.append(f'<{"th" if r == 1 else "td"} style="{style}">{content}</{"th" if r == 1 else "td"}>')
        rows.append("<tr>" + "".join(cells) + "</tr>")
    return ("<!doctype html><meta charset='utf-8'><style>body{font:13px 'Malgun Gothic',sans-serif;margin:16px;background:#fff}"
            "h1{font-size:16px}p{color:#555}table{border-collapse:collapse}td,th{border:1px solid #cfd6d4;padding:4px 7px;white-space:nowrap;"
            "max-width:200px;overflow:hidden;text-overflow:ellipsis}img{width:64px;height:64px;object-fit:contain}</style>"
            f"<h1>price_check.xlsx · 시트 '{html.escape(sheet.title)}' (앞 {limit - 1}행)</h1>"
            f"<p>웹의 [사이즈별 가격 비교 엑셀] 버튼으로 받은 실제 파일을 표로 옮겨 그린 미리보기입니다. {REPLAY}. 전체 {sheet.max_row - 1}행.</p>"
            "<table>" + "".join(rows) + "</table>")


def capture_app():
    from playwright.sync_api import sync_playwright
    os.makedirs(ASSETS, exist_ok=True)
    with open(os.path.join(DEMO, "DEMO_INFO.json"), encoding="utf-8") as f:
        info = json.load(f)
    import sqlite3
    db = sqlite3.connect(f"file:{os.path.join(DEMO, 'musinsa_data.sqlite')}?mode=ro", uri=True)
    coded = [r[0] for r in db.execute("SELECT musinsa_goods_no FROM kream_search_cache WHERE match_method = 'model_code' AND match_verified = 1")]
    db.close()
    price_ids = [g for g in info["batch"] if g in coded][:10]
    fail = price_ids[3:6]
    server = start_server(fail)
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(channel="chrome")
            context = browser.new_context(viewport={"width": 1440, "height": 1000}, device_scale_factor=1.25, locale="ko-KR", accept_downloads=True)
            page = context.new_page()
            data = f"데모 DB의 {info['products']}개 상품 ({info['source']})"

            # 1. products and per-size price comparison
            page.goto(BASE + "/#products")
            page.wait_for_selector("#product-rows tr.size")
            page.screenshot(path=os.path.join(ASSETS, "01-products.png"))
            note("01-products", "상품 목록과 사이즈별 가격 비교", "price-compare", data, f"{REPLAY}. 무신사 회원가와 KREAM 즉시구매가·즉시판매가를 같은 사이즈끼리 비교한 화면.")
            page.fill("#product-search", coded and price_ids[0] or "")
            page.wait_for_timeout(300)
            page.locator("#product-table").screenshot(path=os.path.join(ASSETS, "01b-products-zoom.png"))
            note("01b-products-zoom", "사이즈별 가격 비교 확대", "price-compare", data, f"{REPLAY}. 한 상품의 사이즈 행. 가격 차이 = KREAM 즉시구매가 − 무신사 회원가이며 수수료·배송비는 반영하지 않았다.")
            page.fill("#product-search", "")

            # 2. VLM review job in progress (stored verdicts replayed, paced so the bar is visible)
            page.goto(BASE + "/#review")
            page.wait_for_selector("#review-metrics .metric")
            page.screenshot(path=os.path.join(ASSETS, "02a-review-before.png"))
            note("02a-review-before", "VLM 검토 탭: 실행 전", "vlm-review", data, f"{REPLAY}. 아직 판정을 불러오기 전이라 검수 대기가 0이다.")
            page.click("#review-stored")
            page.wait_for_selector("#jobbar .job progress", timeout=15000)
            page.wait_for_timeout(1800)
            page.screenshot(path=os.path.join(ASSETS, "02-vlm-progress.png"), clip={"x": 0, "y": 0, "width": 1440, "height": 330})
            note("02-vlm-progress", "VLM 검토 작업 진행 상태", "jobs", data,
                 f"{REPLAY}. 저장된 gpt-5.4-mini 판정을 불러오는 작업. 진행 표시가 보이도록 데모 서버가 상품당 0.25초 지연을 넣었다. 모델을 새로 호출하지 않았다.")
            jobs = wait_idle(page)
            page.reload()
            page.wait_for_selector("#review-list details")
            page.screenshot(path=os.path.join(ASSETS, "02c-review-queue.png"))
            note("02c-review-queue", "검수 대기 목록과 집계", "vlm-review", data, f"{REPLAY}. 판정 요약: {json.dumps(jobs[0]['summary'], ensure_ascii=False)}")

            def open_card(goods_no, state=""):
                page.select_option("#review-filter", state)
                page.fill("#review-search", goods_no)
                card = page.locator(f"details[data-goods='{goods_no}']")
                card.wait_for()
                if not card.evaluate("el => el.open"):      # a card left open earlier stays open
                    card.locator("summary").click()
                page.wait_for_timeout(700)
                return card

            # 3-4. side-by-side images, reason and the three decisions
            card = open_card("3976350")
            card.screenshot(path=os.path.join(ASSETS, "03-review-card.png"))
            note("03-review-card", "원본과 후보 이미지를 나란히 보는 검토 화면", "vlm-review", "무신사 3976350 (에어포스 1 '07 화이트)과 KREAM 후보",
                 f"{REPLAY}. 같은 모델번호를 일반판과 협업판이 함께 써서 규칙으로 확정하지 못한 상품. VLM 판정은 2026-10-09에 저장된 것이다.")
            card.locator(".review-main").screenshot(path=os.path.join(ASSETS, "04-review-decision.png"))
            note("04-review-decision", "판단 사유와 승인·거절·보류", "human-review", "무신사 3976350",
                 f"{REPLAY}. 판정, 모델명, 실행 시각, 사유, 모델 출력 신뢰도(확률 아님 표시), 충돌 안내, 후보 선택, 승인·거절·보류 버튼.")

            # 5. conflict with the model-code rule
            card = open_card("5042975")
            card.screenshot(path=os.path.join(ASSETS, "05-conflict.png"))
            note("05-conflict", "규칙과 VLM이 충돌하는 경우", "vlm-review", "무신사 5042975 (아디제로 EVO SL)",
                 f"{REPLAY}. 모델번호 규칙이 확정한 후보와 VLM이 고른 후보가 다르다. 같은 모델의 다른 색상을 고른 사례.")

            # 6. a wrong suggestion, then reject it
            card = open_card("4313969")
            card.screenshot(path=os.path.join(ASSETS, "06-wrong-suggestion.png"))
            note("06-wrong-suggestion", "대표 오판 사례: 닮은 다른 모델을 동일로 판정", "evaluation", "무신사 4313969 (팀버랜드)",
                 f"{REPLAY}. VLM이 모델번호 계열이 다른 상품을 '동일'로 판정한 저장 결과. 충돌 안내에 모델번호 불일치가 표시된다.")
            card.locator("input[type=text]").fill("모델번호 계열이 다름")
            card.get_by_role("button", name="거절 (같은 상품 없음)").click()
            page.wait_for_timeout(900)

            # 7. approve, then see it in the price comparison
            card = open_card("3976350")
            card.get_by_role("button", name="승인 (선택한 후보)").click()
            page.wait_for_timeout(900)
            open_card("3976350", "approved").screenshot(path=os.path.join(ASSETS, "07a-approved-card.png"))
            note("07a-approved-card", "승인 후 검수 이력", "human-review", "무신사 3976350",
                 "캡처 스크립트가 데모 DB에서 승인 버튼을 누른 결과. 사람의 실제 검수가 아니며 검토 시간도 사람의 업무 시간이 아니다.", kind="automation")
            page.fill("#review-search", "")
            page.select_option("#review-filter", "")
            page.locator("#review-metrics").screenshot(path=os.path.join(ASSETS, "08-review-metrics.png"))
            note("08-review-metrics", "검수 업무량 집계", "measurement", data,
                 "대상, 규칙 확정, VLM 판정, 검수 대기, 승인·거절·보류. 승인 1건·거절 1건은 캡처 스크립트가 누른 것이다(사람 측정 아님).", kind="automation")
            page.goto(BASE + "/#products")
            page.wait_for_selector("#product-rows tr")
            page.fill("#product-search", "3976350")
            page.wait_for_timeout(500)
            page.screenshot(path=os.path.join(ASSETS, "07-approved-in-comparison.png"), clip={"x": 0, "y": 0, "width": 1440, "height": 900})
            note("07-approved-in-comparison", "검수 승인이 가격 비교에 반영된 화면", "price-compare", "무신사 3976350 ↔ KREAM 12831",
                 f"{REPLAY}. 승인 뒤 매칭 상태가 '검수 승인'으로 바뀌고 저장돼 있던 KREAM 사이즈별 시세와 비교된다. 승인은 캡처 스크립트가 눌렀다.", kind="automation")

            # 9-10. injected failure, stop, resume of the unfinished part only
            started = page.evaluate("ids => api('/api/jobs/kream', {ids})", price_ids)
            wait_idle(page)
            page.goto(BASE + "/#export")
            page.wait_for_selector("#history-rows tr")
            page.wait_for_timeout(600)
            page.locator("#history-rows").locator("xpath=ancestor::div[1]").screenshot(path=os.path.join(ASSETS, "09-failed-job.png"))
            note("09-failed-job", "중단된 작업과 실패·미처리 항목", "recovery", f"데모 DB의 모델번호 매칭 상품 {len(price_ids)}개",
                 f"테스트 장애 주입: 상품 3개의 KREAM 가격 조회가 시간 초과되도록 했다(실제 쇼핑몰 미접속, 저장된 시세 재생). 작업은 '일부 실패'로 끝났고 "
                 "완료·실패·미처리 건수와 [실패·미처리만 다시 실행] 버튼이 보인다.", kind="fault-injection")
            page.get_by_role("button", name="실패·미처리").first.click()
            page.wait_for_timeout(800)
            wait_idle(page)
            page.reload()
            page.wait_for_selector("#history-rows tr")
            page.wait_for_timeout(600)
            page.locator("#history-rows").locator("xpath=ancestor::div[1]").screenshot(path=os.path.join(ASSETS, "10-resumed.png"))
            note("10-resumed", "실패·미처리 항목만 다시 실행한 결과", "recovery", f"앞 작업 {started['job']}의 남은 항목",
                 "테스트 장애 주입 해제 후 재실행. 이미 끝난 상품은 다시 조회하지 않았다(저장된 시세 재생).", kind="fault-injection")
            page.goto(BASE + "/#logs")
            page.wait_for_selector("#log-rows tr, #log-list > *", timeout=8000)
            page.wait_for_timeout(500)
            page.screenshot(path=os.path.join(ASSETS, "11-logs.png"))
            note("11-logs", "작업별 구조화 로그", "jobs", "위 작업들의 로그", "테스트 장애 주입으로 생긴 오류 기록이 작업 ID·상품번호와 함께 남는다.", kind="fault-injection")

            # 12. a real export, rendered for a readable picture
            response = page.request.get(BASE + "/api/export/price-check.xlsx")
            preview = os.path.join(ASSETS, "_export_preview.html")
            with open(preview, "w", encoding="utf-8") as f:
                f.write(xlsx_preview(response.body()))
            viewer = context.new_page()
            viewer.set_viewport_size({"width": 1500, "height": 900})
            viewer.goto("file:///" + preview.replace("\\", "/"))
            viewer.screenshot(path=os.path.join(ASSETS, "12-export.png"))
            viewer.close()
            os.remove(preview)
            note("12-export", "사이즈별 가격 비교 엑셀 내보내기 결과", "export", data,
                 f"{REPLAY}. 실제로 내려받은 price_check.xlsx의 첫 시트를 표로 옮겨 그린 미리보기(엑셀 프로그램 화면이 아님). 노랑·주황은 강조 규칙.")

            # other tabs referenced by the diagram
            for name, tab, selector, title, feature in (
                    ("14-collect", "collect", "#mode", "수집 대상과 작업 버튼", "collect"),
                    ("15-match", "match", "#match-rows tr", "매칭 상태와 근거 목록", "rules"),
                    ("16-monitor", "monitor", "#auto-status", "모니터링·자동화 설정", "monitor")):
                page.goto(BASE + "/#" + tab)
                page.wait_for_selector(selector)
                page.wait_for_timeout(500)
                page.screenshot(path=os.path.join(ASSETS, name + ".png"))
                note(name, title, feature, data, f"{REPLAY}. 화면 구성 확인용이며 이 화면에서 수집이나 알림을 실행하지 않았다.")

            # mobile width of the review screen
            phone = browser.new_context(viewport={"width": 390, "height": 844}, device_scale_factor=2, locale="ko-KR")
            mobile = phone.new_page()
            mobile.goto(BASE + "/#review")
            mobile.wait_for_selector("#review-list details")
            mobile.locator("details[data-goods='3157671'] summary").click()
            mobile.wait_for_timeout(700)
            mobile.locator("details[data-goods='3157671']").screenshot(path=os.path.join(ASSETS, "13-review-mobile.png"))
            note("13-review-mobile", "모바일 폭의 검토 화면", "vlm-review", "무신사 3157671 (리복 클럽 C 85)", f"{REPLAY}. 390px 폭.")
            browser.close()
    finally:
        server.terminate()
    with open(os.path.join(ASSETS, "captures.json"), "w", encoding="utf-8", newline="\n") as f:
        json.dump({"demo": {k: info[k] for k in ("builtAt", "source", "products", "reopened", "removed", "network")},
                   "faultInjection": {"products": fail, "what": "KREAM 가격 조회 시간 초과(처음 두 번)"}, "captures": captures}, f, ensure_ascii=False, indent=1)


def capture_page():
    """The portfolio page: desktop, an opened panel, mobile, and a keyboard/click check."""
    from playwright.sync_api import sync_playwright
    index = "file:///" + os.path.join(ROOT, "docs", "portfolio", "index.html").replace("\\", "/")
    checks = {}
    with sync_playwright() as p:
        browser = p.chromium.launch(channel="chrome")
        page = browser.new_page(viewport={"width": 1440, "height": 1000}, device_scale_factor=1.25)
        errors = []
        page.on("pageerror", lambda e: errors.append(str(e)))
        page.goto(index)
        page.wait_for_selector("[data-node]")
        page.screenshot(path=os.path.join(ASSETS, "page-desktop.png"))
        nodes = page.locator("svg [data-node]").evaluate_all("els => els.map(e => e.dataset.node)")
        opened = []
        for node in nodes:
            page.locator(f"svg [data-node='{node}']").click()
            page.wait_for_selector("#panel:not([hidden])")
            title = page.locator("#panel-title").inner_text()
            opened.append({"node": node, "title": title, "hash": page.evaluate("location.hash")})
            if node == "vlm":
                page.wait_for_timeout(400)
                page.screenshot(path=os.path.join(ASSETS, "page-panel-vlm.png"))
            page.keyboard.press("Escape")
            page.wait_for_selector("#panel", state="hidden")
        checks["nodesOpened"] = opened
        page.locator("svg [data-node]").first.focus()
        page.keyboard.press("Enter")
        checks["keyboardOpens"] = page.locator("#panel").is_visible()
        figure = page.locator("#panel figure img").first
        if figure.count():
            figure.click()
            checks["imageZoom"] = page.locator("#lightbox").is_visible()
            page.keyboard.press("Escape")
        page.keyboard.press("Escape")
        page.goto(index + "#storage")
        page.wait_for_selector("#panel:not([hidden])")
        checks["deepLink"] = page.locator("#panel-title").inner_text()
        page.screenshot(path=os.path.join(ASSETS, "page-panel-storage.png"))
        checks["pageErrors"] = errors
        phone = browser.new_context(viewport={"width": 390, "height": 844}, device_scale_factor=2)
        mobile = phone.new_page()
        mobile.goto(index)
        mobile.wait_for_selector(".node-list [data-node]")      # the SVG is replaced by a list at this width
        mobile.screenshot(path=os.path.join(ASSETS, "page-mobile.png"))
        mobile.locator(".node-list [data-node='vlm']").click()
        mobile.wait_for_selector("#panel:not([hidden])")
        mobile.screenshot(path=os.path.join(ASSETS, "page-mobile-panel.png"))
        checks["mobileHorizontalScroll"] = mobile.evaluate("document.documentElement.scrollWidth > window.innerWidth + 1")
        # the text must stay readable when images cannot be loaded
        blocked = browser.new_context(viewport={"width": 1200, "height": 900})
        blocked.route("**/*.png", lambda route: route.abort())
        plain = blocked.new_page()
        plain.goto(index + "#vlm")
        plain.wait_for_selector("#panel:not([hidden])")
        checks["readableWithoutImages"] = len(plain.locator("#panel").inner_text()) > 400
        browser.close()
    with open(os.path.join(ASSETS, "page-check.json"), "w", encoding="utf-8", newline="\n") as f:
        json.dump(checks, f, ensure_ascii=False, indent=1)
    print(json.dumps(checks, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    mode = sys.argv[1] if len(sys.argv) > 1 else "app"
    capture_app() if mode == "app" else capture_page()

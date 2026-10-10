# -*- coding: utf-8 -*- 

from __future__ import annotations

import faulthandler
import sys

# Redirect faulthandler output to a file to prevent crash on executable
try:
    if sys.stderr is None:
        faulthandler_log = open('faulthandler_log.txt', 'w')
        faulthandler.enable(file=faulthandler_log)
    else:
        faulthandler.enable()
except Exception:
    pass

import tkinter as tk
from tkinter import ttk, messagebox, scrolledtext
import webbrowser
import threading
import json
import queue
import csv
import uuid
import pandas as pd
import openpyxl
import math
import io
import sqlite3
import re
from typing import List
from datetime import datetime, timedelta, timezone
import subprocess
from shop_comparator.core.size_converter import find_ugg_match, get_mm_size, normalize_size

# 이미지 표시
from PIL import Image, ImageTk
import requests

# 모듈화된 파일 임포트
from shop_comparator.core import config
from shop_comparator.core import utils
from shop_comparator.storage import database
from shop_comparator.shops import crawler
from shop_comparator.shops import kream_crawler
from shop_comparator.shops import kream_login
from shop_comparator.desktop import kream_filter
from shop_comparator.desktop import rotate_module
from shop_comparator.core import webhook
from shop_comparator.desktop import poizon_search
from shop_comparator.desktop import crawl_poizon_price
from shop_comparator.desktop import poizon_login
import traceback
from shop_comparator.shops import get_kream_price6
from shop_comparator.shops import poizon_seller_crawler
from shop_comparator.shops.kream_crawler import KreamAuthException
from shop_comparator.core.utils import resource_path, get_persistent_path

import random
import time
import os
from urllib.parse import parse_qs, urlparse

class App:
    def __init__(self, root):
        self.root = root
        root.title("무신사-크림 가격 비교 및 엑셀 생성기")
        root.geometry("1450x900")

        self.items_in_tree = {}
        self.musinsa_cookie_header = None
        self.musinsa_auth_payload = {}
        self.kream_cookies = None
        self.kream_login_lock = threading.Lock()
        self._promo_groups_cache = None
        self.last_crawl_source = None
        self.poizon_crawler = None
        
        self.poizon_queue = queue.Queue()
        threading.Thread(target=self._poizon_worker_loop, daemon=True).start()

        self.auto_crawl_active = threading.Event()
        self.auto_event_active = threading.Event()
        self.auto_event_thread = None
        self.auto_event_last_prices = {}
        self.auto_event_last_mode = None
        self.auto_event_mode_override = None
        self.auto_event_seed_brands = []
        self.auto_event_seed_count = 0
        self.auto_event_force_reseed = False
        self.auto_event_use_current_selection = False

        self.KREAM_COOKIES_PATH = utils.get_persistent_path(config.KREAM_COOKIES_FILE)
        self.KREAM_TOKEN_PATH = utils.get_persistent_path(config.KREAM_TOKEN_FILE)
        self.POIZON_COOKIES_PATH = utils.get_persistent_path("poizon_cookies.json")
        
        self.categories = {
            "상의": "001", "아우터": "002", "바지": "003", "가방": "004",
            "신발": "103", "원피스/스커트": "100", "패션소품": "101", "속옷/홈웨어": "026",
            "스포츠/레저": "017", "디지털/라이프": "102", "뷰티": "104", "키즈": "106", "아울렛": "107"
        }
        self.ranking_categories = {
            "신발 랭킹": "shoes_ranking",
            "부티크 - 지갑/가방": "boutique_wallet_bag_ranking",
            "부티크 - 악세서리": "boutique_accessory_ranking"
        }
        self.category_vars = {}
        self.ranking_vars = {}

        self.product_list_window = None
        self.tree = None # Will be initialized in open_product_list_window

        self.build_widgets()
        self.load_initial_data()

    def _poizon_worker_loop(self):
        while True:
            try:
                task = self.poizon_queue.get()
                if task is None:
                    break
                self._process_poizon_task(task)
                self.poizon_queue.task_done()
            except Exception as e:
                self.log(f"[ERROR] Poizon Worker Error: {e}")

    def _process_poizon_task(self, task):
        task_type = task.get('type')
        
        if task_type == 'LOGIN':
            email = task.get('email')
            pw = task.get('pw')
            callback = task.get('callback') # Optional callback for UI update
            
            try:
                if self.poizon_crawler:
                    self.poizon_crawler.close()
                
                # Crawler is instantiated INSIDE this thread
                self.poizon_crawler = poizon_seller_crawler.PoizonSellerCrawler(log_callback=self.log)
                self.poizon_crawler.start_browser(headless=True, cookies_path=self.POIZON_COOKIES_PATH)
                success = self.poizon_crawler.login(email, pw, cookies_path=self.POIZON_COOKIES_PATH)

                if success and self.poizon_crawler:
                    self.poizon_crawler.close()
                    self.poizon_crawler = None
                
                if callback:
                    self.root.after(0, callback, success)
            except Exception as e:
                self.log(f"[POIZON WORKER] Login Error: {e}")
                if callback:
                    self.root.after(0, callback, False)

        elif task_type == 'CRAWL_LIST':
            items = task.get('items', {})
            email = task.get('email')
            pw = task.get('pw')

            self.log("[INFO] Poizon Seller 가격 크롤링 및 UI 업데이트를 시작합니다.")
            self.last_crawl_source = 'poizon'
            self.root.after(0, self._change_tree_headings, "poizon")
            
            # Ensure crawler is initialized and browser is running
            if not self.poizon_crawler or not self.poizon_crawler.page:
                self.log("[INFO] 브라우저를 시작하고 쿠키 기반 로그인을 시도합니다.")
                try:
                    if self.poizon_crawler:
                        self.poizon_crawler.close()
                    self.poizon_crawler = poizon_seller_crawler.PoizonSellerCrawler(log_callback=self.log)
                    self.poizon_crawler.start_browser(headless=True, cookies_path=self.POIZON_COOKIES_PATH)
                except Exception as e:
                    self.log(f"[ERROR] 브라우저 시작 중 오류: {e}")
                    return

            # Perform login check / execution
            if not self.poizon_crawler.login(email, pw, cookies_path=self.POIZON_COOKIES_PATH):
                 self.log("[ERROR] 로그인 실패. 아이디/비밀번호를 확인하거나 잠시 후 다시 시도하세요.")
                 return

            self.log(f"[INFO] Poizon Seller item queue: {len(items)}")

            # Load Cache
            cache_file = "poizon_musinsa_overlapping_products.csv"
            cache_map = {}
            if os.path.exists(cache_file):
                try:
                    cache_df = pd.read_csv(cache_file, encoding='utf-8-sig')
                    # Create a map: musinsa_product_id -> poizon_url (or stored link)
                    # The file has columns: musinsa_product_id, ..., poizon_link (if updated previously) or poizon_url
                    # We prioritize 'poizon_link' if available, otherwise 'poizon_url'
                    
                    # Ensure columns exist
                    cols = cache_df.columns
                    
                    duplicated_targets = set()
                    target_index = {}

                    for _, row in cache_df.iterrows():
                        mid = str(row.get('musinsa_product_id', ''))
                        plink = row.get('poizon_link') if 'poizon_link' in cols else None
                        purl = row.get('poizon_url') if 'poizon_url' in cols else None
                        cached_serial = str(row.get('musinsa_serial_number', '') or '').strip()
                        
                        target_link = plink if pd.notna(plink) and plink else purl
                        if mid and target_link and pd.notna(target_link):
                            parsed = urlparse(str(target_link))
                            query = parse_qs(parsed.query)
                            cache_identity = query.get('globalSpuIds', [''])[0] or str(target_link)
                            if cache_identity in target_index and target_index[cache_identity] != mid:
                                duplicated_targets.add(cache_identity)
                            target_index[cache_identity] = mid
                            cache_map[mid] = {
                                "link": str(target_link),
                                "serial": cached_serial,
                                "identity": cache_identity,
                            }

                    if duplicated_targets:
                        stale_mids = [
                            mid for mid, cache_entry in cache_map.items()
                            if cache_entry.get("identity") in duplicated_targets
                        ]
                        for mid in stale_mids:
                            cache_map.pop(mid, None)
                        self.log(
                            f"[WARN] 포이즌 캐시 충돌 감지: {len(stale_mids)}개 항목은 새 검색으로 재시도합니다."
                        )
                    
                    self.log(f"[INFO] 포이즌 캐시 로드 완료: {len(cache_map)}개 항목")
                except Exception as e:
                    self.log(f"[WARN] 포이즌 캐시 파일 로드 실패: {e}")

            # Load Blacklist
            blacklist_file = "poizon_search_blacklist.csv"
            blacklist_set = set()
            if os.path.exists(blacklist_file):
                try:
                    with open(blacklist_file, 'r', encoding='utf-8') as f:
                        reader = csv.reader(f)
                        for row in reader:
                            if row:
                                blacklist_set.add(row[0].strip())
                    self.log(f"[INFO] 포이즌 블랙리스트 로드 완료: {len(blacklist_set)}개 항목")
                except Exception as e:
                    self.log(f"[WARN] 포이즌 블랙리스트 파일 로드 실패: {e}")

            new_cache_entries = []

            for iid, item in items.items():
                # Check for cancellation or stop logic if needed (optional)
                
                musinsa_name = item.get("goodsName", "").strip()
                serial_number = str(item.get("serial_number") or "").strip()
                musinsa_goods_no = str(item.get("goodsNo", ""))

                if not musinsa_goods_no:
                    continue

                # Check Blacklist
                if musinsa_goods_no in blacklist_set:
                    self.log(f"[SKIP] 블랙리스트에 등록된 상품입니다: {musinsa_goods_no}")
                    self.root.after(0, self.update_poizon_prices, iid, item, {}, None)
                    continue
                
                # ... (Serial number extraction logic skipped for brevity, assuming standard flow) ...
                if not serial_number:
                    if "/" in musinsa_name:
                        parts = musinsa_name.rsplit("/", 1)
                        if len(parts) > 1:
                            potential = parts[-1].strip()
                            if len(potential) >= 3:
                                serial_number = potential
                                self.log(f"  -> [자동보정] 상품명에서 품번 추출: {serial_number}")
                    
                    if not serial_number and "(" in musinsa_name and musinsa_name.endswith(")"):
                        m = re.search(r'\(([^)]+)\)$', musinsa_name)
                        if m:
                            potential = m.group(1).strip()
                            if len(potential) >= 3:
                                serial_number = potential
                                self.log(f"  -> [자동보정] 상품명 괄호에서 품번 추출: {serial_number}")

                if not serial_number:
                    self.log(f"[SKIP] 상품 '{musinsa_name}'의 품번(Serial Number)이 없어 건너뜁니다.")
                    continue

                self.log(f"\n--- [Poizon Seller] 상품: {musinsa_name} (품번: {serial_number}) ---")

                try:
                    # Check cache first
                    cache_entry = cache_map.get(musinsa_goods_no)
                    cached_url = cache_entry.get("link") if cache_entry else None
                    prices = {}
                    poizon_link = ""

                    if cached_url:
                        cached_serial = str(cache_entry.get("serial") or "").strip()
                        if cached_serial and cached_serial != serial_number:
                            self.log(
                                f"  -> [CACHE] 시리얼 불일치로 캐시 무시: saved={cached_serial}, current={serial_number}"
                            )
                            cached_url = None

                    if cached_url:
                        self.log(f"  -> [CACHE] 캐시된 URL 사용: {cached_url}")
                        prices, poizon_link = self.poizon_crawler.scrape_direct_url(cached_url)

                    if not cached_url or not prices:
                        if cached_url and not prices:
                            self.log("  -> [CACHE] 캐시 URL 결과가 비어 새 검색으로 재시도합니다.")
                        prices, poizon_link = self.poizon_crawler.search_and_scrape(serial_number)
                        # Only update cache if we found a valid link and it wasn't there before
                        if poizon_link:
                             new_entry = {
                                'musinsa_product_id': musinsa_goods_no,
                                'musinsa_brand_name': item.get("brandName", ""),
                                'musinsa_product_name': musinsa_name,
                                'musinsa_serial_number': serial_number,
                                'poizon_product_id': '', # Can be extracted from URL if needed, but link is prioritized
                                'poizon_product_name': '',
                                'poizon_url': '', # Legacy
                                'poizon_image_url': '',
                                'poizon_link': poizon_link
                            }
                             new_cache_entries.append(new_entry)

                    if prices:
                        self.log(f"  -> 가격 수집 성공: {len(prices)}개 사이즈")
                    else:
                        self.log(f"  -> 가격 수집 실패 또는 결과 없음.")
                        if not poizon_link and not cached_url: # Only blacklist if we tried searching fresh and failed
                            self.log(f"  -> 검색 결과 없음. 블랙리스트에 추가: {musinsa_goods_no}")
                            blacklist_set.add(musinsa_goods_no)
                            try:
                                with open(blacklist_file, 'a', newline='', encoding='utf-8') as f:
                                    writer = csv.writer(f)
                                    writer.writerow([musinsa_goods_no])
                            except Exception as e:
                                self.log(f"[ERROR] 블랙리스트 저장 실패: {e}")
                    
                    # Update UI in main thread
                    self.root.after(0, self.update_poizon_prices, iid, item, prices, poizon_link)
                except Exception as e:
                    self.log(f"[ERROR] 크롤링 중 오류 발생: {e}")

                time.sleep(1)
            
            # Save new cache entries
            if new_cache_entries:
                try:
                    # Append to CSV
                    header = not os.path.exists(cache_file)
                    fieldnames = ['musinsa_product_id', 'musinsa_brand_name', 'musinsa_product_name', 
                                  'musinsa_serial_number', 'poizon_product_id', 'poizon_product_name', 
                                  'poizon_url', 'poizon_image_url', 'poizon_link']
                    
                    with open(cache_file, 'a', newline='', encoding='utf-8-sig') as f:
                        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction='ignore')
                        if header:
                            writer.writeheader()
                        writer.writerows(new_cache_entries)
                    self.log(f"[INFO] 포이즌 캐시 업데이트 완료: {len(new_cache_entries)}개 신규 항목 저장")
                except Exception as e:
                    self.log(f"[ERROR] 포이즌 캐시 저장 실패: {e}")

            self.log("[OK] Poizon Seller 가격 정보 크롤링 및 UI 업데이트 완료.")

    def _execute_kream_api(self, func, *args, **kwargs):
        try:
            return func(*args, **kwargs)
        except KreamAuthException as e:
            self.log(f"[KREAM AUTH] API 인증 실패: {e}. 자동 재로그인을 시도합니다.")
            with self.kream_login_lock:
                uid = self.kream_id.get().strip()
                pw = self.kream_pw.get().strip()
                if not uid or not pw:
                    self.log("[KREAM AUTH] KREAM 아이디/비밀번호가 입력되지 않아 자동 재로그인을 중단합니다.")
                    raise e

                self.root.after(0, lambda: self.kream_status_label.config(text="토큰 만료. 재로그인 시도 중...", foreground="orange"))
                success = kream_login.login_and_save_kream_cookies(uid, pw, self.KREAM_COOKIES_PATH, log_callback=self.log)
                
                if success:
                    self.log("[KREAM AUTH] 재로그인 성공. 이전 작업을 재시도합니다.")
                    try:
                        with open(self.KREAM_COOKIES_PATH, 'r', encoding='utf-8') as f:
                            self.kream_cookies = json.load(f)
                        self.root.after(0, lambda: self.kream_status_label.config(text="재로그인 성공!", foreground="green"))
                        return func(*args, **kwargs)
                    except Exception as load_err:
                        self.log(f"[ERROR] 재로그인 후 쿠키/토큰 파일 로드 실패: {load_err}")
                        raise e
                else:
                    self.log("[KREAM AUTH] 자동 재로그인에 실패했습니다. 수동으로 로그인해주세요.")
                    self.root.after(0, lambda: self.kream_status_label.config(text="자동 재로그인 실패", foreground="red"))
                    raise e

    def build_widgets(self):
        controls_container = ttk.Frame(self.root, padding=8)
        controls_container.pack(fill="x", side="top", pady=5)

        # Row 1: Musinsa + KREAM
        row1_frame = ttk.Frame(controls_container)
        row1_frame.pack(fill="x", pady=2)

        musinsa_frame = ttk.LabelFrame(row1_frame, text="무신사", padding=6)
        musinsa_frame.pack(side="left", fill="x", expand=True, padx=2)
        ttk.Label(musinsa_frame, text="아이디").grid(row=0, column=0, sticky="e", padx=4, pady=2)
        self.musinsa_id = ttk.Entry(musinsa_frame, width=15)
        self.musinsa_id.grid(row=0, column=1, sticky="w", padx=4)
        ttk.Label(musinsa_frame, text="비밀번호").grid(row=0, column=2, sticky="e", padx=4)
        self.musinsa_pw = ttk.Entry(musinsa_frame, width=15, show="*")
        self.musinsa_pw.grid(row=0, column=3, sticky="w", padx=4)
        ttk.Button(musinsa_frame, text="쿠키 얻기", command=self.on_click_get_musinsa_cookie).grid(row=0, column=4, padx=6)
        ttk.Label(musinsa_frame, text="TopN").grid(row=1, column=0, sticky="e", padx=4)
        self.entry_topn = ttk.Entry(musinsa_frame, width=6)
        self.entry_topn.insert(0, "100")
        self.entry_topn.grid(row=1, column=1, sticky="w")
        ttk.Button(musinsa_frame, text="무신사 크롤링", command=self.on_click_musinsa_collect).grid(row=1, column=2, padx=6, pady=4)
        ttk.Button(musinsa_frame, text="브랜드 필터링", command=self.open_brand_selection_window).grid(row=1, column=3, padx=6, pady=4)

        kream_frame = ttk.LabelFrame(row1_frame, text="KREAM", padding=6)
        kream_frame.pack(side="left", fill="x", expand=True, padx=2)
        ttk.Label(kream_frame, text="아이디").grid(row=0, column=0, sticky="e", padx=4, pady=2)
        self.kream_id = ttk.Entry(kream_frame, width=15)
        self.kream_id.grid(row=0, column=1, sticky="w", padx=4)
        ttk.Label(kream_frame, text="비밀번호").grid(row=0, column=2, sticky="e", padx=4)
        self.kream_pw = ttk.Entry(kream_frame, width=15, show="*")
        self.kream_pw.grid(row=0, column=3, sticky="w", padx=4)
        ttk.Button(kream_frame, text="KREAM 로그인", command=self.on_click_kream_login).grid(row=0, column=4, padx=6)
        ttk.Label(kream_frame, text="최소 거래건수").grid(row=1, column=0, sticky="e", padx=4)
        self.entry_min_total = ttk.Entry(kream_frame, width=8)
        self.entry_min_total.insert(0, "0")
        self.entry_min_total.grid(row=1, column=1, sticky="w")
        ttk.Label(kream_frame, text="최소 한달 거래수").grid(row=1, column=2, sticky="e", padx=4)
        self.entry_min_monthly_total = ttk.Entry(kream_frame, width=8)
        self.entry_min_monthly_total.insert(0, "0")
        self.entry_min_monthly_total.grid(row=1, column=3, sticky="w")
        self.kream_extract_ids_button = ttk.Button(kream_frame, text="크림 상품번호 추출", command=self.on_click_extract_kream_ids)
        self.kream_extract_ids_button.grid(row=2, column=0, pady=4)
        
        self.kream_extract_prices_button = ttk.Button(kream_frame, text="크림 가격 추출", command=self.on_click_extract_kream_prices)
        self.kream_extract_prices_button.grid(row=2, column=1, pady=4)

        self.kream_filter_button = ttk.Button(kream_frame, text="크림 필터링", command=self.on_click_kream_filter)
        self.kream_filter_button.grid(row=2, column=2, padx=6)

        self.full_crawl_button = ttk.Button(kream_frame, text="무신사+크림 한번에", command=self.on_click_full_crawl)
        self.full_crawl_button.grid(row=2, column=3, padx=6)

        self.kream_status_label = ttk.Label(kream_frame, text="", foreground="red")
        self.kream_status_label.grid(row=3, column=0, columnspan=5, pady=2)

        # Row 2: POIZON + Excel
        row2_frame = ttk.Frame(controls_container)
        row2_frame.pack(fill="x", pady=2)

        poizon_frame = ttk.LabelFrame(row2_frame, text="POIZON (Seller)", padding=6)
        poizon_frame.pack(side="left", fill="x", expand=True, padx=2)
        
        ttk.Label(poizon_frame, text="아이디").grid(row=0, column=0, sticky="e", padx=4, pady=2)
        self.poizon_id = ttk.Entry(poizon_frame, width=15)
        self.poizon_id.grid(row=0, column=1, sticky="w", padx=4)
        
        ttk.Label(poizon_frame, text="비밀번호").grid(row=0, column=2, sticky="e", padx=4)
        self.poizon_pw = ttk.Entry(poizon_frame, width=15, show="*")
        self.poizon_pw.grid(row=0, column=3, sticky="w", padx=4)

        ttk.Button(poizon_frame, text="로그인 & 브라우저 시작", command=self.on_click_poizon_login_auto).grid(row=0, column=4, padx=6)
        
        self.poizon_extract_prices_button = ttk.Button(poizon_frame, text="포이즌 가격 추출 (Seller)", command=self.on_click_extract_poizon_prices)
        self.poizon_extract_prices_button.grid(row=1, column=0, columnspan=2, pady=4, padx=2)

        self.full_poizon_crawl_button = ttk.Button(poizon_frame, text="무신사+포이즌 통합", command=self.on_click_full_poizon_crawl)
        self.full_poizon_crawl_button.grid(row=1, column=2, columnspan=2, pady=4, padx=2)

        self.poizon_status_label = ttk.Label(poizon_frame, text="", foreground="red")
        self.poizon_status_label.grid(row=2, column=0, columnspan=5, pady=2)

        excel_frame = ttk.LabelFrame(row2_frame, text="엑셀 작업", padding=6)
        excel_frame.pack(side="left", fill="x", expand=True, padx=2)
        ttk.Label(excel_frame, text="마진율(%)").grid(row=0, column=0, sticky="e", padx=4, pady=2)
        self.margin_rate = ttk.Entry(excel_frame, width=10)
        self.margin_rate.insert(0, "10.0")
        self.margin_rate.grid(row=0, column=1, sticky="w")
        ttk.Label(excel_frame, text="수량").grid(row=1, column=0, sticky="e", padx=4, pady=2)
        self.quantity = ttk.Entry(excel_frame, width=10)
        self.quantity.insert(0, "1")
        self.quantity.grid(row=1, column=1, sticky="w")
        ttk.Label(excel_frame, text="입찰기한").grid(row=1, column=2, sticky="e", padx=4, pady=2)
        self.deadline = ttk.Combobox(excel_frame, values=["무기한", "60일", "30일", "7일"], width=8)
        self.deadline.set("무기한")
        self.deadline.grid(row=1, column=3, sticky="w")
        
        # Restored Buttons
        ttk.Button(excel_frame, text="엑셀 파일 생성", command=self.on_click_generate_excel).grid(row=0, column=2, columnspan=2, padx=6)
        ttk.Button(excel_frame, text="리스탁 모니터링", command=self.open_monitoring_window).grid(row=2, column=0, columnspan=2, pady=4)
        ttk.Button(excel_frame, text="비교 CSV 저장", command=self.on_click_save_comparison_csv).grid(row=2, column=2, columnspan=2, pady=4)
        ttk.Button(excel_frame, text="가격변동 모니터링", command=self.open_price_monitoring_window).grid(row=3, column=0, columnspan=2, pady=4)
        ttk.Button(excel_frame, text="구매제한 모니터링", command=self.open_purchase_limit_monitoring_window).grid(row=3, column=2, columnspan=2, pady=4)

        # Discord Webhook Frame
        discord_frame = ttk.LabelFrame(controls_container, text="디스코드 알림 설정", padding=6)
        discord_frame.pack(fill="x", pady=5)
        ttk.Label(discord_frame, text="웹훅 URL:").pack(side="left", padx=(0, 5))
        self.discord_webhook_url_entry = ttk.Entry(discord_frame, width=100)
        self.discord_webhook_url_entry.pack(side="left", fill="x", expand=True)
        ttk.Label(discord_frame, text="알림 마진:").pack(side="left", padx=(10, 5))
        self.alert_margin_entry = ttk.Entry(discord_frame, width=10)
        self.alert_margin_entry.insert(0, "0")
        self.alert_margin_entry.pack(side="left", padx=(0, 5))

        auto_event_frame = ttk.LabelFrame(controls_container, text="이벤트 트리거 자동화", padding=6)
        auto_event_frame.pack(fill="x", pady=5)
        ttk.Label(auto_event_frame, text="낮 감시 주기(초)").grid(row=0, column=0, sticky="e", padx=4)
        self.auto_event_day_interval_entry = ttk.Entry(auto_event_frame, width=8)
        self.auto_event_day_interval_entry.insert(0, "0")
        self.auto_event_day_interval_entry.grid(row=0, column=1, sticky="w", padx=4)
        ttk.Label(auto_event_frame, text="밤 전체 갱신 주기(초)").grid(row=0, column=2, sticky="e", padx=4)
        self.auto_event_night_interval_entry = ttk.Entry(auto_event_frame, width=8)
        self.auto_event_night_interval_entry.insert(0, "0")
        self.auto_event_night_interval_entry.grid(row=0, column=3, sticky="w", padx=4)
        self.start_auto_event_button = ttk.Button(
            auto_event_frame,
            text="자동 감지 시작",
            command=self.on_click_start_auto_event_mode,
        )
        self.start_auto_event_button.grid(row=0, column=4, padx=6)
        self.start_day_event_button = ttk.Button(
            auto_event_frame,
            text="낮 모드 시작",
            command=self.on_click_start_day_event_mode,
        )
        self.start_day_event_button.grid(row=0, column=5, padx=6)
        self.start_night_event_button = ttk.Button(
            auto_event_frame,
            text="밤 모드 시작",
            command=self.on_click_start_night_full_mode,
        )
        self.start_night_event_button.grid(row=0, column=6, padx=6)
        self.stop_auto_event_button = ttk.Button(
            auto_event_frame,
            text="자동 이벤트 모드 중지",
            command=self.on_click_stop_auto_event_mode,
            state="disabled",
        )
        self.stop_auto_event_button.grid(row=0, column=7, padx=6)
        self.auto_event_status_label = ttk.Label(auto_event_frame, text="대기", foreground="gray")
        self.auto_event_status_label.grid(row=0, column=8, sticky="w", padx=6)

        # --- Category Selection Frame (Restored) ---
        category_frame = ttk.LabelFrame(controls_container, text="카테고리 선택", padding=6)
        category_frame.pack(fill="x", pady=5)
        self.setup_category_checkboxes(category_frame)

        # --- Product List & Log Section ---
        list_control_frame = ttk.Frame(self.root, padding=8)
        list_control_frame.pack(fill="x")
        
        ttk.Button(list_control_frame, text="상품 목록 및 정보 창 열기 (새 창)", command=self.open_product_list_window).pack(fill="x", pady=5)

        frame_log = ttk.Frame(self.root, padding=6, relief="solid", borderwidth=1)
        frame_log.pack(fill="both", expand=True, padx=8, pady=4)
        ttk.Label(frame_log, text="진행 로그").pack(anchor="w")
        self.log_box = scrolledtext.ScrolledText(frame_log, height=10, state="disabled")
        self.log_box.pack(fill="both", expand=True)

    def open_product_list_window(self):
        if self.product_list_window is not None and self.product_list_window.winfo_exists():
            self.product_list_window.lift()
            return

        self.product_list_window = tk.Toplevel(self.root)
        self.product_list_window.title("상품 목록 및 상세 정보")
        self.product_list_window.geometry("1400x800")
        self.product_list_window.protocol("WM_DELETE_WINDOW", self.on_product_window_close)

        list_frame = ttk.Frame(self.product_list_window, padding=6)
        list_frame.pack(fill="both", expand=True)

        cols = ("rank", "name", "brand", "size", "musinsa_price", "kream_buy_price", "kream_sell_price", "diff")
        self.tree = ttk.Treeview(list_frame, columns=cols, show="headings")
        self.setup_treeview(cols)
        
        vsb = ttk.Scrollbar(list_frame, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        
        self.tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")

        # Buttons in the separate window
        btns = ttk.Frame(self.product_list_window, padding=8)
        btns.pack(fill="x")
        ttk.Button(btns, text="선택 상세보기", command=self.open_selected_detail).pack(side="left", padx=6)
        ttk.Button(btns, text="상품 페이지 열기", command=self.open_selected_link).pack(side="left", padx=6)

        # Reload data into the new tree
        products = database.load_products()
        if products:
            self.populate_tree(products)

    def on_product_window_close(self):
        self.tree = None
        self.product_list_window.destroy()
        self.product_list_window = None

    def setup_category_checkboxes(self, parent_frame):
        self.category_vars = {}
        row, col = 0, 0
        for name, code in self.categories.items():
            var = tk.BooleanVar()
            chk = ttk.Checkbutton(parent_frame, text=name, variable=var)
            chk.grid(row=row, column=col, sticky="w", padx=2)
            self.category_vars[code] = var
            col += 1
            if col > 6:
                col = 0
                row += 1
        
        self.ranking_vars = {}
        for name, code in self.ranking_categories.items():
            var = tk.BooleanVar()
            chk = ttk.Checkbutton(parent_frame, text=name, variable=var)
            chk.grid(row=row, column=col, sticky="w", padx=2)
            self.ranking_vars[code] = var
            col += 1
            if col > 6:
                col = 0
                row += 1

    def setup_treeview(self, cols):
        headings = {"rank": "랭킹", "name": "상품명", "brand": "브랜드", "size": "사이즈", 
                    "musinsa_price": "무신사 최종가", "kream_buy_price": "크림 즉시구매가", 
                    "kream_sell_price": "크림 즉시판매가", "diff": "차액"}
        widths = {"rank": 40, "name": 350, "brand": 120, "size": 80, "musinsa_price": 100, 
                  "kream_buy_price": 100, "kream_sell_price": 100, "diff": 100}
        for col in cols:
            self.tree.heading(col, text=headings.get(col, col))
            self.tree.column(col, width=widths.get(col, 100), anchor="center" if col != 'name' else 'w')

    def log(self, text: str):
        self.root.after(0, utils.append_log, self.log_box, text)

    def load_initial_data(self):
        self.log("[INFO] 데이터베이스 초기화 확인 중…")
        database.init_db()
        musinsa_session = database.load_auth_session()
        self.musinsa_cookie_header = musinsa_session.get("cookie_header")
        self.musinsa_auth_payload = musinsa_session.get("auth_payload") or {}
        crawler.set_auth_payload(self.musinsa_auth_payload)
        if self.musinsa_cookie_header:
            self.log("[INFO] 저장된 무신사 쿠키를 불러왔습니다.")
        
        # Check environment variable for Kream Bearer token first
        if os.environ.get("KREAM_BEARER"):
            self.kream_cookies = {"source": "env"}
            self.log("[INFO] 환경변수 KREAM_BEARER가 감지되었습니다.")
            self.kream_status_label.config(text="토큰(Env) 로드 완료", foreground="green")
        else:
            try:
                with open(self.KREAM_COOKIES_PATH, 'r', encoding='utf-8') as f:
                    self.kream_cookies = json.load(f)
                self.log("[INFO] kream_cookies.json 파일을 불러왔습니다.")
                self.kream_status_label.config(text="KREAM 쿠키 로드 완료", foreground="green")
            except FileNotFoundError:
                self.log("[WARN] kream_cookies.json 파일을 찾을 수 없습니다.")
                self.kream_status_label.config(text="KREAM 쿠키 없음. 로그인 필요.", foreground="red")
            except Exception as e:
                self.log(f"[ERROR] 크림 쿠키 로드 실패: {e}")
                self.kream_status_label.config(text="쿠키 파일 오류", foreground="red")

        products = database.load_products()
        if products:
            self.populate_tree(products)
            self.log(f"[INFO] 저장된 상품 {len(products)}개를 불러왔습니다.")

    def populate_tree(self, items: list[dict]):
        self.items_in_tree.clear()
        if self.tree is not None:
            self.tree.delete(*self.tree.get_children())
        for item in items:
            if self.tree is None:
                iid = f"goods_{item.get('goodsNo')}"
                self.items_in_tree[iid] = item
                continue
            iid = self.tree.insert("", "end", values=(
                item.get("rank", ""), item.get("goodsName", ""), item.get("brandName", ""),
                "", f'{utils.krw(item.get("finalPrice") or 0)}원' if item.get("finalPrice") is not None else "",
                "", "", ""
            ), tags=('product',))
            self.items_in_tree[iid] = item

    def update_kream_prices(self, parent_iid, item_data, prices_by_size, prices_by_size2, all_kream_sizes):
        if self.tree is None:
            return
        musinsa_price = item_data.get("finalPrice") or 0
        
        musinsa_sizes_str = item_data.get("availableSizes", "")
        musinsa_available_sizes = sorted(list(set([s.strip() for s in musinsa_sizes_str.split(',') if s.strip()])))

        if not self.tree.exists(parent_iid):
             return

        for child_iid in self.tree.get_children(parent_iid):
            self.tree.delete(child_iid)
        self.tree.item(parent_iid, open=True)

        # KREAM 사이즈 기준으로 루프 (크림 데이터가 있으면 무조건 표시)
        # 만약 크림 데이터가 아예 없으면 무신사 사이즈라도 보여주기 위해 합집합 사용
        
        # 1. KREAM 사이즈 + 무신사 사이즈 합치기 (중복 제거)
        combined_sizes = set(all_kream_sizes or [])
        combined_sizes.update(str(size).strip() for size in prices_by_size.keys() if str(size).strip())
        combined_sizes.update(str(size).strip() for size in prices_by_size2.keys() if str(size).strip())
        # 무신사 사이즈도 KREAM 포맷과 비교를 위해 정규화 등의 과정이 필요할 수 있으나, 일단 원본 추가
        # (매칭을 위해 임시로 매핑 테이블 생성)
        
        # 매칭 로직: KREAM size -> Musinsa size
        # 1:1 매칭이 되면 Musinsa 가격 표시, 아니면 KREAM 가격만 표시
        
        # Display logic: Iterate KREAM sizes first as they have the prices we just fetched
        displayed_sizes = set()
        
        # 정렬: 숫자 변환 가능한 것은 숫자로, 아니면 문자열로 (튜플 반환으로 타입 에러 방지)
        def size_sort_key(s):
            s_str = str(s)
            s_clean = s_str.replace('.', '', 1)
            if s_clean.isdigit():
                return (0, float(s_str))
            return (1, s_str)

        sorted_kream_sizes = sorted(combined_sizes, key=size_sort_key)

        for k_size in sorted_kream_sizes:
            kream_buy_price = prices_by_size.get(k_size)
            kream_sell_price = prices_by_size2.get(k_size)
            
            # Find matching Musinsa size
            matched_m_size = None
            for m_size in musinsa_available_sizes:
                # Match strategy 1: Numeric/Shoe size comparison
                is_m_size_numeric = m_size.replace('.', '', 1).isdigit()
                converted_k_size_mm = get_mm_size(k_size)
                
                if is_m_size_numeric and converted_k_size_mm and m_size == converted_k_size_mm:
                    matched_m_size = m_size
                    break

                # Match strategy 2: Normalized string comparison
                normalized_m_size = normalize_size(m_size)
                normalized_k_size = normalize_size(k_size)
                if normalized_m_size and normalized_k_size and normalized_m_size == normalized_k_size:
                    matched_m_size = m_size
                    break
            
            # Display logic
            diff_str = ""
            if kream_buy_price is not None and musinsa_price:
                diff = musinsa_price - kream_buy_price
                diff_str = f"{utils.krw(diff)}원"
                if diff > 0: diff_str = f"▲{diff_str}"
                elif diff < 0: diff_str = f"▼{utils.krw(abs(diff))}원"

            # 무신사 매칭 사이즈가 있으면 그 사이즈명도 같이 표기해주거나, 그냥 KREAM 사이즈명 사용
            display_size_name = k_size
            if matched_m_size and matched_m_size != k_size:
                display_size_name = f"{k_size} ({matched_m_size})"

            self.tree.insert(parent_iid, "end", values=(
                "", "", "", display_size_name, 
                f"{utils.krw(musinsa_price)}원" if musinsa_price else "", 
                f"{utils.krw(kream_buy_price)}원" if kream_buy_price is not None else "",
                f"{utils.krw(kream_sell_price)}원" if kream_sell_price is not None else "",
                diff_str
            ), tags=('size',))
            displayed_sizes.add(k_size)
            if matched_m_size:
                displayed_sizes.add(matched_m_size)

        # (옵션) KREAM엔 없고 무신사에만 있는 사이즈 표시 (가격 비교 불가)
        for m_size in musinsa_available_sizes:
            if m_size not in displayed_sizes:
                 # 이미 위에서 매칭되어 displayed_sizes에 들어갔는지 확인 필요. 
                 # 위의 루프는 k_size 기준이라 m_size가 set에 없을 수 있음.
                 # 정밀한 체크 생략하고, 일단 KREAM 정보가 있는 것 위주로 보여줌.
                 pass
        
        if not self.tree.get_children(parent_iid):
             self.tree.insert(parent_iid, "end", values=("", "", "", "정보 없음", "", "", "", ""), tags=('no_data',))

    def _get_discord_alert_margin(self):
        try:
            return int(float((self.alert_margin_entry.get() or "0").strip()))
        except (ValueError, TypeError):
            return 0

    def _send_kream_margin_notification(self, item_data, kream_product_id, prices_by_size, prices_by_size2, all_kream_sizes):
        webhook_url = self.discord_webhook_url_entry.get().strip()
        if not webhook_url:
            return

        musinsa_price = item_data.get("finalPrice") or 0
        if not musinsa_price:
            return

        alert_margin = self._get_discord_alert_margin()
        combined_sizes = set(all_kream_sizes or [])
        combined_sizes.update(str(size).strip() for size in prices_by_size.keys() if str(size).strip())
        combined_sizes.update(str(size).strip() for size in prices_by_size2.keys() if str(size).strip())
        sent_count = 0

        for size in sorted(combined_sizes):
            buy_price = prices_by_size.get(size)
            if buy_price is None:
                continue

            diff = musinsa_price - buy_price
            if diff <= 0 or diff < alert_margin:
                continue

            sell_price = prices_by_size2.get(size)
            sell_text = f"{utils.krw(sell_price)}원" if sell_price is not None else "-"
            goods_no = item_data.get("goodsNo")
            goods_url = item_data.get("linkUrl") or f"https://www.musinsa.com/products/{goods_no}"
            kream_url = f"https://kream.co.kr/products/{kream_product_id}" if kream_product_id else ""
            message = (
                f"**수익 사이즈 감지**\n\n"
                f"**브랜드:** {item_data.get('brandName')}\n"
                f"**상품명:** {item_data.get('goodsName')}\n"
                f"**사이즈:** {size}\n"
                f"**무신사 최종가:** {utils.krw(musinsa_price)}원\n"
                f"**크림 즉시구매가:** {utils.krw(buy_price)}원\n"
                f"**크림 즉시판매가:** {sell_text}\n"
                f"**차액:** {utils.krw(diff)}원\n"
                f"**알림 마진 기준:** {utils.krw(alert_margin)}원\n"
                f"**무신사 링크:** {goods_url}\n"
                f"**크림 링크:** {kream_url}"
            )
            sent = webhook.send_discord_notification(webhook_url, message)
            if sent:
                sent_count += 1
                self.log(f"[DISCORD] '{item_data.get('goodsName')}' {size} 알림 전송 완료.")
            else:
                self.log(f"[DISCORD] '{item_data.get('goodsName')}' {size} 알림 전송 실패.")

        if sent_count == 0:
            self.log(f"[DISCORD] '{item_data.get('goodsName')}' 알림 조건 미충족.")

    def _change_tree_headings(self, source: str):
        if self.tree is None:
            return
        if source == "kream":
            self.tree.heading("kream_buy_price", text="크림 즉시구매가")
            self.tree.heading("kream_sell_price", text="크림 즉시판매가")
        elif source == "poizon":
            self.tree.heading("kream_buy_price", text="포이즌 구매가")
            self.tree.heading("kream_sell_price", text="") # 판매가 컬럼은 사용 안함

    def update_poizon_prices(self, parent_iid, item_data, prices_by_size, poizon_link=None):
        if self.tree is None:
            return
        
        if poizon_link:
            item_data['poizon_link'] = poizon_link

        musinsa_price = item_data.get("finalPrice") or 0
        
        musinsa_sizes_str = item_data.get("availableSizes", "")
        musinsa_available_sizes = set([s.strip() for s in musinsa_sizes_str.split(',') if s.strip()])

        if not self.tree.exists(parent_iid):
            return

        for child_iid in self.tree.get_children(parent_iid):
            self.tree.delete(child_iid)
        self.tree.item(parent_iid, open=True)

        if not prices_by_size:
            self.tree.insert(parent_iid, "end", values=("", "", "", "정보 없음", "", "", "", ""), tags=('no_poizon_data',))
            return

        # 사이즈 키를 float으로 변환하여 정렬
        sorted_sizes = sorted(prices_by_size.keys(), key=lambda s: float(s) if s.replace('.', '', 1).isdigit() else float('inf'))

        # 성별 추측 (상품명이나 카테고리에 'WOMEN', '여성' 등이 있으면 W, 아니면 M 기본값 등)
        # 현재는 UGG 예시가 WOMEN이므로 우선 W를 시도하고, 매칭 안되면 M도 시도하는 방식 등을 고려할 수 있음.
        # 여기서는 상품명에 [WOMEN] 등이 포함된 경우를 간단히 체크
        gender = 'M'
        goods_name = item_data.get("goodsName", "").upper()
        if "WOMEN" in goods_name or "여성" in goods_name or "W" in goods_name.split():
            gender = 'W'

        for size in sorted_sizes:
            # Poizon 사이즈(US)를 MM로 변환 시도
            converted_mm = get_mm_size(size, gender)
            
            # 매칭 여부 확인
            # 1. 변환된 MM가 무신사 사이즈 목록에 있는지
            # 2. 혹은 무신사 사이즈 목록이 비어있으면(정보 없음) 그냥 표시할지 결정 (여기서는 매칭된 것만 표시)
            
            is_matched = False
            display_size = size

            if converted_mm:
                if converted_mm in musinsa_available_sizes:
                    is_matched = True
                    display_size = f"{size} ({converted_mm})"
            elif size in musinsa_available_sizes: # 변환 실패했더라도 그 자체로 텍스트 일치하면 (예: 의류 S, M)
                is_matched = True

            # [Fix] 가방/지갑 등 단일 사이즈(FREE/ONE SIZE) 매칭 완화 로직
            if not is_matched and len(musinsa_available_sizes) == 1:
                m_single = list(musinsa_available_sizes)[0].upper()
                # 무신사가 범용 사이즈이고 포이즌 사이즈도 뭔가 존재한다면 매칭으로 간주
                if m_single in ["FREE", "ONE SIZE", "ONESIZE", "OS", "FR"]:
                    is_matched = True
                    display_size = f"{size} (-> {m_single})"

            # Modified logic: Difference = Musinsa Price - Poizon Price (Savings/Difference)
            # Ensure musinsa_price is a number
            if isinstance(musinsa_price, str):
                 actual_musinsa_price = utils.parse_krw(musinsa_price)
            else:
                 actual_musinsa_price = musinsa_price if musinsa_price is not None else 0
            
            poizon_price = prices_by_size.get(size)
            
            # Formula: Musinsa Final Price - Poizon Price = Difference
            # If Poizon price is missing, diff is 0 (or undefined)
            if poizon_price is not None:
                diff = actual_musinsa_price - poizon_price
            else:
                diff = 0

            diff_str = f"{utils.krw(diff)}원"
            
            if diff > 0:
                diff_str = f"▲{diff_str}" # Musinsa is more expensive (Positive difference)
            elif diff < 0:
                diff_str = f"▼{diff_str}" # Musinsa is cheaper (Negative difference)

            # Display musinsa price for all rows since it applies to all sizes
            display_musinsa_price = f"{utils.krw(musinsa_price)}원" if musinsa_price is not None else ""
            
            self.tree.insert(parent_iid, "end", values=(
                "", "", "", display_size, display_musinsa_price,
                f"{utils.krw(poizon_price)}원" if poizon_price is not None else "",
                "",  # 포이즌은 판매가가 없으므로 비워둠
                diff_str
            ), tags=('size',))

    def on_click_poizon_login_auto(self):
        email = self.poizon_id.get().strip()
        pw = self.poizon_pw.get().strip()
        if not email or not pw:
            messagebox.showwarning("입력 필요", "POIZON 아이디/비밀번호를 입력하세요.")
            return

        def _callback(success):
            if success:
                self.poizon_status_label.config(text="POIZON 로그인 성공", foreground="green")
            else:
                self.poizon_status_label.config(text="POIZON 로그인 실패", foreground="red")

        self.poizon_queue.put({
            'type': 'LOGIN',
            'email': email,
            'pw': pw,
            'callback': _callback
        })

    def on_click_poizon_login(self):
        # 수동 로그인이므로 ID/PW 입력 불필요
        self.poizon_status_label.config(text="브라우저가 열리면 로그인을 완료해주세요.", foreground="blue")
        
        def _task():
            # 인자는 placeholder로 넘김
            success = poizon_login.login_and_save_poizon_cookies("manual", "manual", log_callback=self.log)
            if success:
                self.root.after(0, lambda: self.poizon_status_label.config(text="POIZON 로그인 감지 성공!", foreground="green"))
            else:
                self.root.after(0, lambda: self.poizon_status_label.config(text="POIZON 로그인 실패/취소", foreground="red"))
        
        threading.Thread(target=_task, daemon=True).start()

    def on_click_extract_poizon_prices(self):
        if not self.items_in_tree:
            messagebox.showwarning("데이터 필요", "무신사 크롤링을 먼저 진행해주세요.")
            return
        
        email = self.poizon_id.get().strip()
        pw = self.poizon_pw.get().strip()

        # 크롤링 전 결과 창 미리 열기
        self.open_product_list_window()

        self.log("[INFO] 포이즌 가격 정보 크롤링 요청됨 (대기열 추가)")
        
        # Pass a copy of the items to avoid modification during crawl
        items_copy = self.items_in_tree.copy()
        
        self.poizon_queue.put({
            'type': 'CRAWL_LIST',
            'items': items_copy,
            'email': email,
            'pw': pw
        })




    def update_musinsa_tree_item(self, item):
        def _update_in_main_thread():
            if self.tree is None:
                return
            iid = self.tree.insert("", "end", values=(
                item.get("rank", ""),
                item.get("goodsName", ""),
                item.get("brandName", ""),
                "", # size
                f'{utils.krw(item.get("finalPrice") or 0)}원' if item.get("finalPrice") is not None else "", # musinsa_price
                "", # kream_buy_price
                "", # kream_sell_price
                ""  # diff
            ), tags=('product',))
            self.items_in_tree[iid] = item
        self.root.after(0, _update_in_main_thread)

    def on_click_get_musinsa_cookie(self):
        uid = self.musinsa_id.get().strip()
        pw  = self.musinsa_pw.get().strip()
        if not uid or not pw:
            messagebox.showwarning("입력 필요","무신사 아이디/비밀번호를 입력하세요.")
            return
        def _task():
            res = crawler.get_cookie_with_selenium(uid, pw, log_callback=self.log)
            if res.get("ok"):
                self.musinsa_cookie_header = res.get("cookie_header", "")
                self.musinsa_auth_payload = res.get("auth_payload") or {}
                crawler.set_auth_payload(self.musinsa_auth_payload)
                database.save_cookie(self.musinsa_cookie_header, self.musinsa_auth_payload)
                self.log("[OK] 무신사 쿠키 추출 및 저장 완료.")
            else:
                messagebox.showerror("로그인 실패", res.get("message", "알 수 없는 오류"))
        threading.Thread(target=_task, daemon=True).start()

    def on_click_musinsa_collect(self):
        if not self.musinsa_cookie_header:
            messagebox.showwarning("쿠키 필요", "먼저 무신사 쿠키를 얻어주세요.")
            return
        try:
            top_n = int(self.entry_topn.get().strip())
        except (ValueError, TypeError):
            messagebox.showwarning("입력 오류", "최대 상품 수에 유효한 숫자를 입력하세요.")
            return
        selected_ranking = [code for code, var in self.ranking_vars.items() if var.get()]
        selected_categories = [code for code, var in self.category_vars.items() if var.get()]
        if not selected_ranking and not selected_categories:
            messagebox.showwarning("카테고리 선택 필요", "하나 이상의 카테고리를 선택해주세요.")
            return
        self.tree.delete(*self.tree.get_children())
        self.items_in_tree.clear()
        def _task():
            self._run_musinsa_crawl_task(top_n, selected_ranking, selected_categories)
        threading.Thread(target=_task, daemon=True).start()

    def open_brand_selection_window(self):
        csv_path = self.find_csv_file("kream_overlapping_brands")
        if not csv_path:
            messagebox.showerror("파일 없음", "매칭 DB CSV 파일 (kream_musinsa_overlapping_products_view...)을 찾을 수 없습니다.")
            return

        brands = utils.get_musinsa_brands_from_csv(csv_path)
        if not brands:
            messagebox.showinfo("브랜드 없음", "CSV 파일에서 'musinsa' 브랜드를 찾을 수 없습니다.")
            return

        win = tk.Toplevel(self.root)
        win.title("브랜드 선택 및 크롤링")
        win.geometry("600x700")

        tk.Label(win, text="크롤링할 브랜드를 선택하세요:", justify=tk.LEFT, anchor="w").pack(pady=(10, 5), padx=10, fill='x')

        # --- Scrollable Checkbox Frame ---
        canvas = tk.Canvas(win)
        scrollbar = ttk.Scrollbar(win, orient="vertical", command=canvas.yview)
        scrollable_frame = ttk.Frame(canvas)

        scrollable_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(
                scrollregion=canvas.bbox("all")
            )
        )

        canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)

        canvas.pack(side="left", fill="both", expand=True, padx=(10,0), pady=5)
        scrollbar.pack(side="right", fill="y", padx=(0,10))
        # --- End Scrollable Frame ---

        brand_vars = {}
        for brand in brands:
            var = tk.BooleanVar()
            chk = ttk.Checkbutton(scrollable_frame, text=brand, variable=var)
            chk.pack(anchor="w", padx=10, pady=2)
            brand_vars[brand] = var

        bottom_frame = ttk.Frame(win)
        bottom_frame.pack(fill='x', side='bottom', pady=10, padx=10)

        input_frame = ttk.Frame(bottom_frame)
        input_frame.pack(pady=5)
        tk.Label(input_frame, text="브랜드당 수집 개수:").pack(side='left', padx=(0, 5))
        count_entry = ttk.Entry(input_frame, width=10)
        count_entry.insert(0, "30")
        count_entry.pack(side='left')

        crawl_button = ttk.Button(bottom_frame, text="선택한 브랜드 크롤링 시작", command=lambda: self.start_brand_crawling(
            win,
            [brand for brand, var in brand_vars.items() if var.get()],
            count_entry.get()
        ))
        crawl_button.pack(pady=5)

        full_crawl_button = ttk.Button(bottom_frame, text="선택 브랜드 무신사+크림", command=lambda: self.start_brand_full_crawl(
            win,
            [brand for brand, var in brand_vars.items() if var.get()],
            count_entry.get()
        ))
        full_crawl_button.pack(pady=5)

        full_poizon_crawl_button = ttk.Button(bottom_frame, text="선택 브랜드 무신사+포이즌", command=lambda: self.start_brand_full_poizon_crawl(
            win,
            [brand for brand, var in brand_vars.items() if var.get()],
            count_entry.get()
        ))
        full_poizon_crawl_button.pack(pady=5)

        auto_start_button = ttk.Button(bottom_frame, text="선택 브랜드 이벤트 자동 시작", command=lambda: self.start_brand_auto_event_mode(
            win,
            [brand for brand, var in brand_vars.items() if var.get()],
            count_entry.get()
        ))
        auto_start_button.pack(pady=5)

        auto_stop_button = ttk.Button(bottom_frame, text="자동 중지", command=self.on_click_stop_auto_crawl)
        auto_stop_button.pack(pady=5)

    def start_brand_crawling(self, window, selected_brands, count_str):
        if not self.musinsa_cookie_header:
            messagebox.showwarning("쿠키 필요", "먼저 무신사 쿠키를 얻어주세요.", parent=window)
            return
        if not selected_brands:
            messagebox.showwarning("브랜드 선택 필요", "하나 이상의 브랜드를 선택해주세요.", parent=window)
            return
        try:
            count = int(count_str)
            if count <= 0:
                raise ValueError()
        except ValueError:
            messagebox.showwarning("입력 오류", "수집 개수에 유효한 양의 정수를 입력하세요.", parent=window)
            return

        window.destroy()
        self.tree.delete(*self.tree.get_children())
        self.items_in_tree.clear()
        
        def _task():
            self._run_brand_crawl_task(selected_brands, count)
        
        threading.Thread(target=_task, daemon=True).start()

    def _run_brand_crawl_task(self, brands, count):
        self.log(f"[INFO] 선택된 {len(brands)}개 브랜드에 대해 각각 최대 {count}개 상품 크롤링을 시작합니다.")
        try:
            s = crawler.build_session(self.musinsa_cookie_header)
            all_products_map = {}

            for i, brand_name in enumerate(brands):
                self.log(f"--- ({i+1}/{len(brands)}) 브랜드 '{brand_name}' 처리 시작 ---")
                
                product_list = crawler.get_products_by_brand(brand_name, count, log_callback=self.log)
                if not product_list:
                    self.log(f"[WARN] 브랜드 '{brand_name}'에서 상품을 수집하지 못했습니다.")
                    continue

                for j, dmeta in enumerate(product_list):
                    goods_no = str(dmeta.get("goodsNo"))
                    if not goods_no or goods_no in all_products_map:
                        continue

                    link_url = dmeta.get("linkUrl", f"https://www.musinsa.com/app/goods/{goods_no}")
                    self.log(f"[INFO] ({j+1}/{len(product_list)}) 상품 처리 시작: {goods_no} ({dmeta.get('goodsName')})")
                    
                    thumbnail_url = dmeta.get("thumbnail")
                    if thumbnail_url and "imageUrl" not in dmeta:
                        dmeta["imageUrl"] = thumbnail_url

                    try:
                        row = crawler.process_one(
                            s,
                            self.musinsa_cookie_header,
                            goods_no,
                            link_url,
                            dmeta,
                            sleep_each=0.5,
                            auth_payload=self.musinsa_auth_payload,
                        )
                    except Exception as e:
                        row = {"goodsNo": goods_no, "linkUrl": link_url, "goodsName": dmeta.get("goodsName"), "brandName": dmeta.get("brandName"), "error": str(e)}
                        self.log(f"[WARN] 처리 실패 goodsNo={goods_no}: {e}")
                    
                    row["rank"] = dmeta.get("rank", "")
                    all_products_map[goods_no] = row
                    
                    if self.update_musinsa_tree_item:
                        self.update_musinsa_tree_item(row)

            final_items = list(all_products_map.values())
            if final_items:
                database.save_products(final_items)
                self.log(f"[DB] 총 {len(final_items)}개의 브랜드 기반 상품 저장 완료.")
                self.log("[OK] 브랜드 기반 크롤링 완료.")
                return True
            else:
                self.log("[WARN] 브랜드 크롤링으로 수집된 아이템이 없습니다.")
                self.log("[OK] 브랜드 기반 크롤링 완료.")
                return False
        except Exception as e:
            self.log(f"[ERROR] 브랜드 기반 크롤링 중 오류 발생: {e}")
            import traceback
            self.log(traceback.format_exc())
            return False

    def start_brand_full_crawl(self, window, selected_brands, count_str):
        window.destroy()
        self.log("[INFO] 선택 브랜드 무신사+크림 통합 크롤링 시작")
        self.tree.delete(*self.tree.get_children())
        self.items_in_tree.clear()

        def _task():
            if self._run_brand_crawl_task(selected_brands, int(count_str)):
                if self.kream_cookies:
                    self._run_extract_kream_ids_task()
                    self._run_extract_kream_prices_task()
            self.log("[OK] 선택 브랜드 통합 크롤링 완료.")
        
        threading.Thread(target=_task, daemon=True).start()

    def start_brand_full_poizon_crawl(self, window, selected_brands, count_str):
        window.destroy()
        self.log("[INFO] 선택 브랜드 무신사+포이즌 통합 크롤링 시작")
        self.tree.delete(*self.tree.get_children())
        self.items_in_tree.clear()

        def _task():
            if self._run_brand_crawl_task(selected_brands, int(count_str)):
                self.poizon_queue.put({'type': 'CRAWL_LIST', 'items': self.items_in_tree.copy()})
            self.log("[OK] 선택 브랜드 포이즌 통합 크롤링 완료.")

        threading.Thread(target=_task, daemon=True).start()

    def start_brand_auto_event_mode(self, window, selected_brands, count_str, mode_override=None):
        if not self.musinsa_cookie_header:
            messagebox.showwarning("쿠키 필요", "먼저 무신사 쿠키를 얻어주세요.", parent=window)
            return
        if not selected_brands:
            messagebox.showwarning("브랜드 선택 필요", "하나 이상의 브랜드를 선택해주세요.", parent=window)
            return

        try:
            count = int(count_str)
            if count <= 0:
                raise ValueError()
        except ValueError:
            messagebox.showwarning("입력 오류", "수집 개수는 1 이상의 정수여야 합니다.", parent=window)
            return

        window.destroy()
        self.auto_event_seed_brands = list(selected_brands)
        self.auto_event_seed_count = count
        self.auto_event_force_reseed = True
        self.auto_event_use_current_selection = False
        self.items_in_tree.clear()
        if self.tree is not None:
            self.tree.delete(*self.tree.get_children())
        self._start_auto_event_mode(mode_override=mode_override)

    def start_brand_auto_crawl(self, window, selected_brands, count_str):
        if self.auto_crawl_active.is_set():
            messagebox.showinfo("알림", "자동 크롤링이 이미 실행 중입니다.")
            return

        try:
            count = int(count_str)
            if count <= 0:
                raise ValueError()
        except ValueError:
            messagebox.showwarning("입력 오류", "수집 개수에 유효한 양의 정수를 입력하세요.", parent=window)
            return

        window.destroy()
        self.log("[INFO] 선택 브랜드 전체 자동 크롤링 시작")
        self.auto_crawl_active.set()
        self.start_auto_crawl_button.config(state="disabled")
        self.stop_auto_crawl_button.config(state="normal")


        def _task():
            while self.auto_crawl_active.is_set():
                self.log("[AUTO CRAWL - BRAND] 새 크롤링 사이클 시작.")
                self.tree.delete(*self.tree.get_children())
                self.items_in_tree.clear()

                crawl_success = self._run_brand_crawl_task(selected_brands, count)
                
                if self.auto_crawl_active.is_set() and crawl_success and self.kream_cookies:
                    self.log("[AUTO CRAWL - BRAND] 무신사 크롤링 완료. 크림 정보 추출 시작.")
                    self._run_extract_kream_ids_task()
                    if self.auto_crawl_active.is_set():
                        self._run_extract_kream_prices_task()
                
                if self.auto_crawl_active.is_set():
                    self.log("[AUTO CRAWL - BRAND] 크롤링 사이클 완료. 다음 사이클까지 대기합니다...")
                    time.sleep(5) # 5 seconds delay
                else:
                    self.log("[AUTO CRAWL - BRAND] 자동 크롤링이 중지되었습니다.")
                    break
            
            self.root.after(0, lambda: self.start_auto_crawl_button.config(state="normal"))
            self.root.after(0, lambda: self.stop_auto_crawl_button.config(state="disabled"))

        threading.Thread(target=_task, daemon=True).start()

    def _run_musinsa_crawl_task(self, top_n, selected_ranking, selected_categories):
        crawled_items = []
        try:
            ranking_map = {
                "shoes_ranking": ("[INFO] 신발 랭킹 크롤링을 시작합니다.", config.SHOES_RANKING_LAYOUT),
                "boutique_wallet_bag_ranking": ("[INFO] 부티크 - 지갑/가방 랭킹 크롤링을 시작합니다.", config.BOUTIQUE_WALLET_BAG_RANKING_LAYOUT),
                "boutique_accessory_ranking": ("[INFO] 부티크 - 악세서리 랭킹 크롤링을 시작합니다.", config.BOUTIQUE_ACCESSORY_RANKING_LAYOUT),
            }
            for ranking_code in selected_ranking:
                if ranking_code in ranking_map:
                    log_msg, layout_url = ranking_map.get(ranking_code)
                    self.log(log_msg)
                    items = crawler.run_ranking_crawl(
                        self.musinsa_cookie_header,
                        layout_url,
                        top_n,
                        log_callback=self.log,
                        progress_callback=self.update_musinsa_tree_item,
                        auth_payload=self.musinsa_auth_payload,
                    )
                    crawled_items.extend(items)
            if selected_categories:
                self.log(f"[INFO] 선택된 카테고리 {len(selected_categories)}개에 대해 각각 Top {top_n} 크롤링을 시작합니다.")
                items = crawler.run_full_crawl(
                    self.musinsa_cookie_header,
                    selected_categories,
                    top_n,
                    log_callback=self.log,
                    progress_callback=self.update_musinsa_tree_item,
                    auth_payload=self.musinsa_auth_payload,
                )
                crawled_items.extend(items)
            if crawled_items:
                unique_items = {item['goodsNo']: item for item in crawled_items}.values()
                database.save_products(list(unique_items))
                self.log(f"[DB] 총 {len(unique_items)}개의 무신사 상품 저장 완료.")
                self.populate_tree(list(unique_items))
                return True
            else:
                self.log("[WARN] 크롤링된 아이템이 없습니다.")
                return False
        except Exception as e:
            self.log(f"[ERROR] 무신사 크롤링 중 오류 발생: {e}")
            return False

    def on_click_kream_login(self):
        uid = self.kream_id.get().strip()
        pw = self.kream_pw.get().strip()
        if not uid or not pw:
            messagebox.showwarning("입력 필요", "KREAM 아이디/비밀번호를 입력하세요.")
            return
        self.kream_status_label.config(text="로그인 시도 중...", foreground="blue")
        def _task():
            success = kream_login.login_and_save_kream_cookies(uid, pw, self.KREAM_COOKIES_PATH, log_callback=self.log)
            if success:
                try:
                    with open(self.KREAM_COOKIES_PATH, 'r', encoding='utf-8') as f:
                        self.kream_cookies = json.load(f)
                    self.root.after(0, lambda: self.kream_status_label.config(text="KREAM 로그인 성공!", foreground="green"))
                except Exception as e:
                    self.root.after(0, lambda: self.log(f"[ERROR] 로그인 후 쿠키 로드 실패: {e}"))
            else:
                self.root.after(0, lambda: self.kream_status_label.config(text="KREAM 로그인 실패", foreground="red"))
        threading.Thread(target=_task, daemon=True).start()

    def on_click_extract_kream_ids(self):
        if not self.kream_cookies:
            messagebox.showwarning("로그인 필요", "크림 로그인을 먼저 진행해주세요.")
            return
        if not self.items_in_tree:
            messagebox.showwarning("데이터 필요", "무신사 크롤링을 먼저 진행해주세요.")
            return
        self.log("[INFO] 크림 상품번호 추출 시작")
        threading.Thread(target=self._run_extract_kream_ids_task, daemon=True).start()

    def _run_extract_kream_ids_task(self):
        self.log("[INFO] Kream 상품번호 추출 및 DB 업데이트를 시작합니다.")
        match_df, db_path = self._get_or_update_db_dataframe(update_prices=False, search_if_missing=True)
        if match_df is None:
            self.log("[ERROR] Kream 상품번호 추출을 중단합니다.")
            return
        self.log("[OK] 크림 상품번호 추출 및 DB 업데이트 완료.")

    def on_click_extract_kream_prices(self):
        print(">>> [DEBUG] on_click_extract_kream_prices called")
        if not self.kream_cookies:
            messagebox.showwarning("로그인 필요", "크림 로그인을 먼저 진행해주세요.")
            return
        if not self.items_in_tree:
            messagebox.showwarning("데이터 필요", "무신사 크롤링을 먼저 진행해주세요.")
            return
        self.log("[INFO] 크림 가격 정보 크롤링 시작")
        threading.Thread(target=self._run_extract_kream_prices_task, daemon=True).start()

    def _auto_login_kream(self):
        """KREAM 자동 재로그인 시도"""
        self.log("[KREAM AUTH] 토큰 만료 감지. 자동 재로그인을 시도합니다.")
        with self.kream_login_lock:
            uid = self.kream_id.get().strip()
            pw = self.kream_pw.get().strip()
            if not uid or not pw:
                self.log("[KREAM AUTH] 아이디/비밀번호가 없어 재로그인 불가.")
                return False

            self.root.after(0, lambda: self.kream_status_label.config(text="재로그인 시도 중...", foreground="orange"))
            success = kream_login.login_and_save_kream_cookies(uid, pw, self.KREAM_COOKIES_PATH, log_callback=self.log)
            
            if success:
                self.log("[KREAM AUTH] 재로그인 성공. 토큰이 갱신되었습니다.")
                try:
                    with open(self.KREAM_COOKIES_PATH, 'r', encoding='utf-8') as f:
                        self.kream_cookies = json.load(f)
                    self.root.after(0, lambda: self.kream_status_label.config(text="재로그인 성공!", foreground="green"))
                    return True
                except Exception as e:
                    self.log(f"[ERROR] 재로그인 후 쿠키 로드 실패: {e}")
            else:
                self.log("[KREAM AUTH] 재로그인 실패.")
                self.root.after(0, lambda: self.kream_status_label.config(text="재로그인 실패", foreground="red"))
            return False

    def _run_extract_kream_prices_task(self):
        print(">>> [DEBUG] _run_extract_kream_prices_task started")
        self.log("[INFO] Kream 가격 크롤링 및 UI 업데이트를 시작합니다.")
        self.last_crawl_source = 'kream'
        self.root.after(0, self._change_tree_headings, "kream")

        # 토큰 로드 (1차)
        try:
            access_token, _, _, _, sentry_trace, baggage = kream_crawler._load_tokens()
            final_token = os.getenv("KREAM_BEARER") or access_token
            if final_token:
                final_token = str(final_token).strip().strip('"').strip("'")
            else:
                self.log("[WARN] KREAM 토큰 없음. 로그인 필요.")
                if not self._auto_login_kream():
                    return
                # 재로그인 후 토큰 다시 로드
                access_token, _, _, _, sentry_trace, baggage = kream_crawler._load_tokens()
                final_token = access_token
        except Exception as e:
            self.log(f"[ERROR] 토큰 로드 중 오류: {e}")
            return
        
        match_df, db_path = self._get_or_update_db_dataframe(
            update_prices=False,
            search_if_missing=True,
        )
        if match_df is None:
            self.log("[ERROR] Kream 가격 추출을 위한 매칭 DB 준비에 실패했습니다.")
            return

        self.log("[INFO] UI를 최신 가격 정보로 업데이트합니다. (get_kream_price6.py 사용)")

        if 'other_mall_product_id' in match_df.columns:
            match_df['other_mall_product_id'] = pd.to_numeric(match_df['other_mall_product_id'], errors='coerce').fillna(0).astype(int).astype(str)

        for iid, item in self.items_in_tree.items():
            musinsa_goods_no = str(item.get("goodsNo"))
            musinsa_name = item.get("goodsName")

            if not musinsa_goods_no:
                continue

            matched_rows = match_df[match_df['other_mall_product_id'] == musinsa_goods_no]
            
            if matched_rows.empty:
                continue

            kream_id = matched_rows.iloc[0]['kream_product_id']
            if pd.notna(kream_id) and kream_id != "" and str(kream_id) != "0":
                retry_count = 0
                max_retries = 1
                
                while retry_count <= max_retries:
                    try:
                        kream_id_int = int(float(kream_id))
                        self.log(f"[KREAM] 상품 '{musinsa_name}' (ID: {kream_id_int}) 가격 조회 시도...")
                        
                        current_client_dt = kream_crawler.kst_ts()
                        current_req_key = str(uuid.uuid4())

                        print(f"[DEBUG] Calling run_logic for {kream_id_int}...") # 디버깅 로그

                        buy_price_data_list = get_kream_price6.run_logic(
                            product_id=kream_id_int,
                            picker_type="buy",
                            token=final_token,
                            device_id=kream_crawler.KREAM_DEVICE_ID,
                            web_build_version=kream_crawler.KREAM_WEB_BUILD_VER,
                            api_version=kream_crawler.KREAM_API_VERSION,
                            web_request_secret=kream_crawler.KREAM_WEB_REQUEST_SEC,
                            baggage=baggage,
                            sentry_trace=sentry_trace,
                            client_datetime=current_client_dt,
                            request_key=current_req_key
                        )

                        sell_price_data_list = get_kream_price6.run_logic(
                            product_id=kream_id_int,
                            picker_type="sell",
                            token=final_token,
                            device_id=kream_crawler.KREAM_DEVICE_ID,
                            web_build_version=kream_crawler.KREAM_WEB_BUILD_VER,
                            api_version=kream_crawler.KREAM_API_VERSION,
                            web_request_secret=kream_crawler.KREAM_WEB_REQUEST_SEC,
                            baggage=baggage,
                            sentry_trace=sentry_trace,
                            client_datetime=current_client_dt,
                            request_key=str(uuid.uuid4())
                        )
                        
                        # 성공 시 데이터 처리
                        prices1 = {}
                        prices2 = {}
                        all_kream_sizes = []
                        
                        for p_item in buy_price_data_list:
                            size = str(p_item.get("option_key")).strip()
                            if size:
                                if size not in all_kream_sizes:
                                    all_kream_sizes.append(size)
                                
                                buy_p = p_item.get("price")
                                if buy_p is not None:
                                    prices1[size] = int(buy_p)

                        for p_item in sell_price_data_list:
                            size = str(p_item.get("option_key")).strip()
                            if size:
                                if size not in all_kream_sizes:
                                    all_kream_sizes.append(size)
                                sell_p = p_item.get("price")
                                if sell_p is not None:
                                    prices2[size] = int(sell_p)

                        self.log(
                            f"[KREAM] '{musinsa_name}' 사이즈 가격 수집: "
                            f"buy={len(prices1)}, sell={len(prices2)}, sizes={len(all_kream_sizes)}"
                        )
                        
                        self._send_kream_margin_notification(
                            item,
                            kream_id_int,
                            prices1,
                            prices2,
                            all_kream_sizes,
                        )
                        self.root.after(0, self.update_kream_prices, iid, item, prices1, prices2, sorted(all_kream_sizes))
                        self.log(f"[KREAM] '{musinsa_name}' 가격 업데이트 완료.")
                        break # 성공하면 루프 탈출

                    except Exception as e:
                        if retry_count < max_retries:
                            self.log(f"[WARN] 가격 조회 실패 ({e}). 재로그인 후 1회 재시도합니다.")
                            if self._auto_login_kream():
                                # 토큰 갱신
                                access_token, _, _, _, sentry_trace, baggage = kream_crawler._load_tokens()
                                final_token = access_token
                                retry_count += 1
                                continue
                            else:
                                self.log("[ERROR] 재로그인 실패로 인해 작업을 중단하거나 건너뜁니다.")
                                break
                        else:
                            self.log(f"[ERROR] '{musinsa_name}' 가격 조회 최종 실패: {e}")
                            self.root.after(0, self.update_kream_prices, iid, item, {}, {}, [])
                            break
            else:
                self.root.after(0, self.update_kream_prices, iid, item, {}, {}, [])
        self.log("[OK] 크림 가격 정보 크롤링 및 UI 업데이트 완료.")

    def on_click_kream_filter(self):
        print("on_click_kream_filter메서드 호출")
        if not self.kream_cookies:
            messagebox.showwarning("로그인 필요", "크림 로그인을 먼저 진행해주세요.")
            return
        if not self.items_in_tree:
            messagebox.showwarning("데이터 필요", "무신사 크롤링을 먼저 진행해주세요.")
            return
        
        try:
            min_total = int(self.entry_min_total.get().strip())
        except ValueError:
            min_total = 0

        try:
            min_monthly_total = int(self.entry_min_monthly_total.get().strip())
        except ValueError:
            min_monthly_total = 0

        self.log("[INFO] 크림 필터링 시작 (체결 내역 조회)... ")
        def _task():
            def _filter_logic(access_token, ua, device_id, sentry_trace, baggage):
                print("_filter_logic메서드 호출")
                print(f"{access_token}, {ua}, {device_id}, {sentry_trace}, {baggage}")
                rotate_module.init_rotate_module(
                    ua=ua,
                    api_version=kream_crawler.KREAM_API_VERSION,
                    web_request_sec=kream_crawler.KREAM_WEB_REQUEST_SEC,
                    web_build_ver=kream_crawler.KREAM_WEB_BUILD_VER,
                    device_id=device_id,
                    access_token=access_token,
                )
                
                to_remove = []
                for iid, item_data in self.items_in_tree.items():
                    goods_no = item_data.get("goodsNo")
                    kream_product_id = None

                    cached_search = database.get_cached_kream_search(goods_no)
                    if cached_search and cached_search.get('kream_product_id'):
                        kream_product_id = cached_search.get('kream_product_id')
                    print(f"kream_product_id:{kream_product_id}")

                    if kream_product_id:
                        self.log(f"[KREAM FILTER] 상품 {item_data.get('goodsName')} (ID: {kream_product_id}) 체결 내역 조회 중...")
                        sales_data = rotate_module.fetch_sales(kream_product_id)
                        if not sales_data:
                            continue
                        
                        total_cnt = int(sales_data.get("total") or 0)

                        monthly_cnt = 0
                        one_month_ago = datetime.now(timezone.utc) - timedelta(days=30)
                        
                        for sale_item in sales_data.get("items", []):
                            try:
                                sale_date_str = sale_item.get("date_created")
                                sale_date = datetime.strptime(sale_date_str, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=timezone.utc)
                                if sale_date >= one_month_ago:
                                    monthly_cnt += 1
                            except (ValueError, TypeError):
                                continue

                        item_data['kream_total_transactions'] = total_cnt
                        item_data['kream_monthly_transactions'] = monthly_cnt
                        database.update_product_kream_transactions(goods_no, total_cnt, monthly_cnt)

                        print(f"상품명: {item_data.get('goodsName')}, 한달 거래수: {monthly_cnt}")

                        if (total_cnt < min_total) or (monthly_cnt < min_monthly_total):
                            to_remove.append((iid, item_data, total_cnt, monthly_cnt))
                            self.log(f"[FILTER] 상품 {item_data.get('goodsName')} (ID: {kream_product_id}) 거래수 {total_cnt} (한달: {monthly_cnt}) < (최소: {min_total}, 한달최소: {min_monthly_total}), 제거 목록에 추가.")
                            continue

                        items = sales_data.get("items", [])
                        self.log(f"[KREAM] 상품 {item_data.get('goodsName')} total={total_cnt} (한달: {monthly_cnt}) kept.")
                        for it in items[:5]:
                            print(f"  • option={it.get('option')} price={it.get('price')} time={it.get('date_created')}")
                    else:
                        self.log(f"[KREAM FILTER] 상품 {item_data.get('goodsName')} - KREAM product_id를 DB에서 찾을 수 없습니다. 건너뜁니다.")
                    
                    self.log(f"[KREAM FILTER] 다음 상품으로 넘어갑니다.")
                
                if not to_remove:
                    self.log("[INFO] 필터링 기준에 미달하여 제거된 상품이 없습니다.")
                else:
                    print("\n--- 필터링으로 제거된 상품 목록 ---")
                    for iid, item_data, total_cnt, monthly_cnt in to_remove:
                        self.root.after(0, lambda iid=iid: self.tree.delete(iid))
                        self.items_in_tree.pop(iid, None)
                        print(f"[제거됨] 상품명: {item_data.get('goodsName')}, 브랜드: {item_data.get('brandName')}, 전체거래수: {total_cnt}, 한달거래수: {monthly_cnt}")
                    print("-------------------------------------")

                self.log("[OK] 크림 필터링 완료.")

            access_token, refresh_token, ua_saved, device_id_saved, sentry_trace_saved, baggage_saved = kream_crawler._load_tokens(self.KREAM_TOKEN_PATH)
            
            try:
                self._execute_kream_api(
                    _filter_logic,
                    access_token=access_token,
                    ua=ua_saved,
                    device_id=device_id_saved,
                    sentry_trace=sentry_trace_saved,
                    baggage=baggage_saved
                )
            except Exception as e:
                self.log(f"[ERROR] 크림 필터링 중 오류 발생: {e}")

        threading.Thread(target=_task, daemon=True).start()

    def on_click_full_crawl(self):
        self.log("[INFO] 무신사+크림 통합 크롤링 시작")
        self.tree.delete(*self.tree.get_children())
        self.items_in_tree.clear()
        def _task():
            selected_ranking = [code for code, var in self.ranking_vars.items() if var.get()]
            selected_categories = [code for code, var in self.category_vars.items() if var.get()]
            if self._run_musinsa_crawl_task(int(self.entry_topn.get()), selected_ranking, selected_categories):
                if self.kream_cookies:
                    self._run_extract_kream_ids_task()
                    self._run_extract_kream_prices_task()
            self.log("[OK] 통합 크롤링 완료.")
        threading.Thread(target=_task, daemon=True).start()

    def on_click_full_poizon_crawl(self):
        self.log("[INFO] 무신사+포이즌 통합 크롤링 시작")
        self.tree.delete(*self.tree.get_children())
        self.items_in_tree.clear()

        email = self.poizon_id.get().strip()
        pw = self.poizon_pw.get().strip()
        
        # Check for cookies (logging purpose mainly, as the extraction task loads them)
        if os.path.exists("poizon_cookies.json"):
            self.log("[INFO] Poizon 로그인 쿠키가 확인되었습니다. 인증된 상태로 진행합니다.")
        else:
            self.log("[WARN] Poizon 로그인 쿠키가 없습니다. 비로그인 상태로 진행됩니다.")

        def _task():
            selected_ranking = [code for code, var in self.ranking_vars.items() if var.get()]
            selected_categories = [code for code, var in self.category_vars.items() if var.get()]
            if self._run_musinsa_crawl_task(int(self.entry_topn.get()), selected_ranking, selected_categories):
                self.poizon_queue.put({
                    'type': 'CRAWL_LIST', 
                    'items': self.items_in_tree.copy(),
                    'email': email,
                    'pw': pw
                })
            self.log("[OK] 포이즌 통합 크롤링 완료.")
        threading.Thread(target=_task, daemon=True).start()

    def on_click_start_auto_crawl(self):
        if self.auto_crawl_active.is_set():
            messagebox.showinfo("알림", "자동 크롤링이 이미 실행 중입니다.")
            return

        self.log("[INFO] 전체 자동 크롤링 시작")
        self.auto_crawl_active.set()
        self.start_auto_crawl_button.config(state="disabled")
        self.stop_auto_crawl_button.config(state="normal")

        def _task():
            while self.auto_crawl_active.is_set():
                self.log("[AUTO CRAWL] 새 크롤링 사이클 시작.")
                self.tree.delete(*self.tree.get_children())
                self.items_in_tree.clear()

                selected_ranking = [code for code, var in self.ranking_vars.items() if var.get()]
                selected_categories = [code for code, var in self.category_vars.items() if var.get()]

                crawl_success = self._run_musinsa_crawl_task(int(self.entry_topn.get()), selected_ranking, selected_categories)
                
                if self.auto_crawl_active.is_set() and crawl_success and self.kream_cookies:
                    self.log("[AUTO CRAWL] 무신사 크롤링 완료. 크림 정보 추출 시작.")
                    self._run_extract_kream_ids_task()
                    if self.auto_crawl_active.is_set():
                        self._run_extract_kream_prices_task()
                
                if self.auto_crawl_active.is_set():
                    self.log("[AUTO CRAWL] 크롤링 사이클 완료. 다음 사이클까지 대기합니다...")
                    time.sleep(5) # 5초 대기
                else:
                    self.log("[AUTO CRAWL] 자동 크롤링이 중지되었습니다.")
                    break
            
            self.root.after(0, lambda: self.start_auto_crawl_button.config(state="normal"))
            self.root.after(0, lambda: self.stop_auto_crawl_button.config(state="disabled"))

        threading.Thread(target=_task, daemon=True).start()

    def on_click_stop_auto_crawl(self):
        self.log("[INFO] 자동 크롤링 중지를 요청했습니다.")
        self.auto_crawl_active.clear()

    def _set_auto_event_status(self, text, color="gray"):
        if hasattr(self, "auto_event_status_label"):
            self.root.after(0, lambda: self.auto_event_status_label.config(text=text, foreground=color))

    def _get_auto_event_interval(self, entry_widget, default_value):
        try:
            return max(30, int(float((entry_widget.get() or str(default_value)).strip())))
        except (ValueError, TypeError):
            return default_value

    def _wait_with_stop(self, seconds, event_flag):
        end_time = time.time() + max(0, seconds)
        while event_flag.is_set() and time.time() < end_time:
            time.sleep(min(1.0, max(0.0, end_time - time.time())))

    def _get_auto_mode(self):
        if self.auto_event_mode_override in ("night_full", "day_event"):
            return self.auto_event_mode_override
        now = datetime.now()
        return "night_full" if 0 <= now.hour < 9 else "day_event"

    def _find_tree_iid_by_goods_no(self, goods_no):
        target = str(goods_no)
        for iid, item in self.items_in_tree.items():
            if str(item.get("goodsNo")) == target:
                return iid
        return None

    def _update_tree_parent_row(self, iid, item_data):
        if self.tree is None or not self.tree.exists(iid):
            return
        current_values = list(self.tree.item(iid, "values"))
        while len(current_values) < 8:
            current_values.append("")
        current_values[0] = item_data.get("rank", "")
        current_values[1] = item_data.get("goodsName", "")
        current_values[2] = item_data.get("brandName", "")
        current_values[4] = (
            f"{utils.krw(item_data.get('finalPrice') or 0)}원"
            if item_data.get("finalPrice") is not None else ""
        )
        self.tree.item(iid, values=tuple(current_values))

    def _replace_item_snapshot(self, iid, updated_item):
        self.items_in_tree[iid] = updated_item
        database.update_product_snapshot(updated_item)
        self.root.after(0, self._update_tree_parent_row, iid, updated_item)

    def _persist_last_observed_price(self, goods_no, current_price):
        self.auto_event_last_prices[goods_no] = int(current_price)
        database.update_last_observed_price(goods_no, int(current_price))
        iid = self._find_tree_iid_by_goods_no(goods_no)
        if iid and iid in self.items_in_tree:
            self.items_in_tree[iid]["lastObservedFinalPrice"] = int(current_price)

    def _seed_auto_event_items_if_needed(self):
        should_use_current_selection = self.auto_event_use_current_selection
        self.auto_event_use_current_selection = False
        if self.auto_event_force_reseed:
            self.auto_event_force_reseed = False
            self.items_in_tree.clear()
            if self.tree is not None:
                self.tree.delete(*self.tree.get_children())
        elif self.items_in_tree:
            return True
        if self.auto_event_seed_brands:
            self.log(
                f"[AUTO EVENT] 선택 브랜드 {len(self.auto_event_seed_brands)}개 / 브랜드당 {self.auto_event_seed_count}개로 감시 대상을 준비합니다."
            )
            return self._run_brand_crawl_task(self.auto_event_seed_brands, self.auto_event_seed_count)
        if should_use_current_selection:
            self.log("[AUTO EVENT] 현재 선택된 랭킹/카테고리 기준으로 감시 대상을 새로 크롤링합니다.")
            selected_ranking = [code for code, var in self.ranking_vars.items() if var.get()]
            selected_categories = [code for code, var in self.category_vars.items() if var.get()]
            return self._run_musinsa_crawl_task(int(self.entry_topn.get()), selected_ranking, selected_categories)
        existing_products = database.load_products()
        if existing_products:
            self.populate_tree(existing_products)
            self.log(f"[AUTO EVENT] 저장된 상품 {len(existing_products)}개를 감시 대상으로 복구했습니다.")
            return True
        self.log("[AUTO EVENT] 현재 상품이 없어 무신사 크롤링으로 기준 목록을 준비합니다.")
        selected_ranking = [code for code, var in self.ranking_vars.items() if var.get()]
        selected_categories = [code for code, var in self.category_vars.items() if var.get()]
        return self._run_musinsa_crawl_task(int(self.entry_topn.get()), selected_ranking, selected_categories)

    def _build_auto_event_baseline(self):
        self.auto_event_last_prices = {}
        for item in self.items_in_tree.values():
            goods_no = str(item.get("goodsNo") or "")
            last_observed_price = item.get("lastObservedFinalPrice")
            if goods_no and last_observed_price is not None:
                self.auto_event_last_prices[goods_no] = int(last_observed_price)
            elif goods_no and item.get("finalPrice") is not None:
                self.auto_event_last_prices[goods_no] = int(item.get("finalPrice"))

    def _load_kream_runtime_auth(self):
        access_token, _, _, _, sentry_trace, baggage = kream_crawler._load_tokens()
        final_token = os.getenv("KREAM_BEARER") or access_token
        if final_token:
            final_token = str(final_token).strip().strip('"').strip("'")
        return final_token, sentry_trace, baggage

    def _trust_existing_kream_matches(self):
        return os.getenv("KREAM_TRUST_EXISTING_MATCHES", "").strip().lower() in {"1", "true", "yes", "on"}

    def _normalize_kream_id_value(self, value):
        try:
            if value is None or pd.isna(value):
                return None
            value_str = str(value).strip()
            if value_str in ("", "0", "0.0", "nan"):
                return None
            return str(int(float(value_str)))
        except (TypeError, ValueError):
            return None

    def _is_verified_cached_kream_match(self, cached_search, kream_id=None):
        if not cached_search or not cached_search.get("kream_product_id"):
            return False
        if int(cached_search.get("match_verified") or 0) != 1:
            return False
        if kream_id is None:
            return True
        return self._normalize_kream_id_value(cached_search.get("kream_product_id")) == self._normalize_kream_id_value(kream_id)

    def _ensure_kream_id_for_item(self, item_data):
        goods_no = str(item_data.get("goodsNo") or "").strip()
        if not goods_no:
            return None

        cached_search = database.get_cached_kream_search(goods_no)
        if (
            cached_search
            and cached_search.get("kream_product_id")
            and (self._is_verified_cached_kream_match(cached_search) or self._trust_existing_kream_matches())
        ):
            try:
                return int(float(cached_search.get("kream_product_id")))
            except (TypeError, ValueError):
                pass

        db_path = self.find_csv_file("kream_musinsa_overlapping_products_view")
        if db_path:
            try:
                match_df = pd.read_csv(db_path)
                if "other_mall_product_id" in match_df.columns:
                    match_df["other_mall_product_id"] = (
                        pd.to_numeric(match_df["other_mall_product_id"], errors="coerce")
                        .fillna(0).astype(int).astype(str)
                    )
                    matched_rows = match_df[match_df["other_mall_product_id"] == goods_no]
                    if not matched_rows.empty:
                        kream_id = matched_rows.iloc[0].get("kream_product_id")
                        if self._trust_existing_kream_matches() and pd.notna(kream_id) and str(kream_id).strip() not in ("", "0", "0.0"):
                            kream_id_int = int(float(kream_id))
                            database.cache_kream_search(goods_no, str(kream_id_int), f"https://kream.co.kr/products/{kream_id_int}")
                            return kream_id_int
            except Exception as e:
                self.log(f"[AUTO EVENT] KREAM 매칭 DB 조회 실패: {e}")

        if not self.kream_cookies:
            return None

        musinsa_name = item_data.get("goodsName")
        if not musinsa_name:
            return None

        searched_id = self._execute_kream_api(
            kream_crawler.search_product,
            keyword=musinsa_name,
            brand=item_data.get("brandName"),
            serial_number=item_data.get("serial_number"),
            cookies=self.kream_cookies,
            image_url=item_data.get("imageUrl"),
        )
        if not searched_id:
            return None
        if database.is_kream_id_already_used(str(searched_id), goods_no, verified_only=not self._trust_existing_kream_matches()):
            self.log(f"[AUTO EVENT] KREAM ID 중복 감지로 건너뜁니다: goodsNo={goods_no}, kreamId={searched_id}")
            return None

        kream_id_int = int(float(searched_id))
        database.cache_kream_search(goods_no, str(kream_id_int), f"https://kream.co.kr/products/{kream_id_int}")
        database.update_product_kream_id(goods_no, str(kream_id_int))
        return kream_id_int

    def _fetch_kream_prices_for_item(self, iid, item_data):
        if not self.kream_cookies:
            self.log("[AUTO EVENT] KREAM 로그인 정보가 없어 가격 비교를 건너뜁니다.")
            return False

        try:
            final_token, sentry_trace, baggage = self._load_kream_runtime_auth()
            if not final_token:
                self.log("[AUTO EVENT] KREAM 토큰이 없어 가격 비교를 건너뜁니다.")
                return False
        except Exception as e:
            self.log(f"[AUTO EVENT] KREAM 토큰 로드 실패: {e}")
            return False

        kream_id = self._ensure_kream_id_for_item(item_data)
        if not kream_id:
            self.log(f"[AUTO EVENT] KREAM ID를 찾지 못해 비교를 건너뜁니다: {item_data.get('goodsName')}")
            return False

        current_client_dt = kream_crawler.kst_ts()
        buy_price_data_list = get_kream_price6.run_logic(
            product_id=kream_id,
            picker_type="buy",
            token=final_token,
            device_id=kream_crawler.KREAM_DEVICE_ID,
            web_build_version=kream_crawler.KREAM_WEB_BUILD_VER,
            api_version=kream_crawler.KREAM_API_VERSION,
            web_request_secret=kream_crawler.KREAM_WEB_REQUEST_SEC,
            baggage=baggage,
            sentry_trace=sentry_trace,
            client_datetime=current_client_dt,
            request_key=str(uuid.uuid4())
        )
        sell_price_data_list = get_kream_price6.run_logic(
            product_id=kream_id,
            picker_type="sell",
            token=final_token,
            device_id=kream_crawler.KREAM_DEVICE_ID,
            web_build_version=kream_crawler.KREAM_WEB_BUILD_VER,
            api_version=kream_crawler.KREAM_API_VERSION,
            web_request_secret=kream_crawler.KREAM_WEB_REQUEST_SEC,
            baggage=baggage,
            sentry_trace=sentry_trace,
            client_datetime=current_client_dt,
            request_key=str(uuid.uuid4())
        )

        prices1 = {}
        prices2 = {}
        all_kream_sizes = []

        for p_item in buy_price_data_list:
            size = str(p_item.get("option_key")).strip()
            if size:
                if size not in all_kream_sizes:
                    all_kream_sizes.append(size)
                buy_p = p_item.get("price")
                if buy_p is not None:
                    prices1[size] = int(buy_p)

        for p_item in sell_price_data_list:
            size = str(p_item.get("option_key")).strip()
            if size:
                if size not in all_kream_sizes:
                    all_kream_sizes.append(size)
                sell_p = p_item.get("price")
                if sell_p is not None:
                    prices2[size] = int(sell_p)

        self.log(f"[AUTO EVENT][KREAM] '{item_data.get('goodsName')}' buy={len(prices1)}, sell={len(prices2)}, sizes={len(all_kream_sizes)}")
        self._send_kream_margin_notification(item_data, kream_id, prices1, prices2, all_kream_sizes)
        self.root.after(0, self.update_kream_prices, iid, item_data, prices1, prices2, sorted(all_kream_sizes))
        return True

    def _run_auto_event_cycle(self, trigger_kream_on_drop=True):
        if not self.items_in_tree:
            return

        sess = crawler.build_session(self.musinsa_cookie_header, auth_payload=self.musinsa_auth_payload)
        items_snapshot = list(self.items_in_tree.items())
        for iid, item_data in items_snapshot:
            if not self.auto_event_active.is_set():
                break
            goods_no = str(item_data.get("goodsNo") or "")
            if not goods_no:
                continue
            try:
                result = crawler.process_one(
                    sess,
                    self.musinsa_cookie_header,
                    goods_no,
                    item_data.get("linkUrl") or f"https://www.musinsa.com/products/{goods_no}",
                    item_data,
                    auth_payload=self.musinsa_auth_payload,
                )
                updated_item = dict(item_data)
                updated_item.update(result)
                self._replace_item_snapshot(iid, updated_item)

                previous_price = self.auto_event_last_prices.get(goods_no)
                current_price = updated_item.get("finalPrice")
                if (
                    trigger_kream_on_drop
                    and previous_price is not None
                    and current_price is not None
                    and int(current_price) < int(previous_price)
                ):
                    self.log(f"[AUTO EVENT] 가격 하락 감지: {updated_item.get('goodsName')} ({utils.krw(previous_price)}원 -> {utils.krw(current_price)}원)")
                    try:
                        self._fetch_kream_prices_for_item(iid, updated_item)
                    except Exception as e:
                        self.log(f"[AUTO EVENT] KREAM 비교 실패: {updated_item.get('goodsName')} / {e}")

                if current_price is not None:
                    self._persist_last_observed_price(goods_no, current_price)
            except Exception as e:
                self.log(f"[AUTO EVENT] 상품 모니터링 실패: {goods_no} / {e}")

    def _run_auto_full_cycle(self):
        if not self.items_in_tree:
            return
        self.log("[AUTO EVENT] 밤 모드: 현재 상품 전체 무신사 가격을 갱신합니다.")
        self._run_auto_event_cycle(trigger_kream_on_drop=False)
        if not self.kream_cookies:
            self.log("[AUTO EVENT] KREAM 로그인 정보가 없어 밤 전체 KREAM 비교를 건너뜁니다.")
            return
        self.log("[AUTO EVENT] 밤 모드: 현재 상품 전체 KREAM 비교를 실행합니다.")
        self._run_extract_kream_ids_task()
        self._run_extract_kream_prices_task()
        self._build_auto_event_baseline()

    def _auto_event_loop(self):
        while self.auto_event_active.is_set():
            mode = self._get_auto_mode()
            if mode != self.auto_event_last_mode:
                mode_text = "밤 전체 갱신 모드" if mode == "night_full" else "낮 이벤트 감시 모드"
                self.log(f"[AUTO EVENT] 모드 전환: {mode_text}")
                self._set_auto_event_status(mode_text, "green" if mode == "day_event" else "blue")
                self.auto_event_last_mode = mode

            if not self._seed_auto_event_items_if_needed():
                self.log("[AUTO EVENT] 감시 대상 준비에 실패했습니다. 다음 주기에 다시 시도합니다.")
                self._wait_with_stop(60, self.auto_event_active)
                continue

            try:
                if mode == "night_full":
                    self._run_auto_full_cycle()
                    wait_sec = self._get_auto_event_interval(self.auto_event_night_interval_entry, 1800)
                else:
                    self._run_auto_event_cycle()
                    wait_sec = self._get_auto_event_interval(self.auto_event_day_interval_entry, 300)
            except Exception as e:
                self.log(f"[AUTO EVENT] 자동 이벤트 모드 오류: {e}")
                wait_sec = 60

            if self.auto_event_active.is_set():
                self.log(f"[AUTO EVENT] 다음 주기까지 {wait_sec}초 대기합니다.")
                self._wait_with_stop(wait_sec, self.auto_event_active)

        self._set_auto_event_status("중지됨", "gray")
        self.root.after(0, lambda: self.start_auto_event_button.config(state="normal"))
        self.root.after(0, lambda: self.start_day_event_button.config(state="normal"))
        self.root.after(0, lambda: self.start_night_event_button.config(state="normal"))
        self.root.after(0, lambda: self.stop_auto_event_button.config(state="disabled"))
        self.auto_event_mode_override = None
        self.log("[AUTO EVENT] 자동 이벤트 모드가 종료되었습니다.")

    def _start_auto_event_mode(self, mode_override=None):
        if self.auto_event_active.is_set():
            messagebox.showinfo("알림", "자동 이벤트 모드가 이미 실행 중입니다.")
            return
        if not self.musinsa_cookie_header:
            messagebox.showwarning("무신사 로그인 필요", "먼저 무신사 쿠키 얻기를 완료해주세요.")
            return

        self.auto_event_active.set()
        self.auto_event_last_mode = None
        self.auto_event_mode_override = mode_override
        self._build_auto_event_baseline()
        self._set_auto_event_status("시작 중", "orange")
        self.start_auto_event_button.config(state="disabled")
        self.start_day_event_button.config(state="disabled")
        self.start_night_event_button.config(state="disabled")
        self.stop_auto_event_button.config(state="normal")
        self.log("[AUTO EVENT] 이벤트 트리거 기반 자동화를 시작합니다.")
        self.auto_event_thread = threading.Thread(target=self._auto_event_loop, daemon=True)
        self.auto_event_thread.start()

    def on_click_start_auto_event_mode(self):
        self.auto_event_seed_brands = []
        self.auto_event_seed_count = 0
        self.auto_event_force_reseed = True
        self.auto_event_use_current_selection = True
        self._start_auto_event_mode(mode_override=None)

    def on_click_start_day_event_mode(self):
        self.auto_event_seed_brands = []
        self.auto_event_seed_count = 0
        self.auto_event_force_reseed = True
        self.auto_event_use_current_selection = True
        self._start_auto_event_mode(mode_override="day_event")

    def on_click_start_night_full_mode(self):
        self.auto_event_seed_brands = []
        self.auto_event_seed_count = 0
        self.auto_event_force_reseed = True
        self.auto_event_use_current_selection = True
        self._start_auto_event_mode(mode_override="night_full")

    def on_click_stop_auto_event_mode(self):
        self.log("[AUTO EVENT] 자동 이벤트 모드 중지를 요청했습니다.")
        self.auto_event_active.clear()

    def _get_or_update_db_dataframe(self, update_prices=False, search_if_missing=True):
        if not self.kream_cookies and search_if_missing:
            self.root.after(0, lambda: messagebox.showwarning("KREAM 로그인 필요", "DB 업데이트 및 크롤링을 위해 KREAM 로그인을 먼저 진행해주세요."))
            return None, None

        musinsa_products = list(self.items_in_tree.values())
        if not musinsa_products:
            self.log("[WARN] 무신사 데이터가 없습니다. 먼저 크롤링을 실행하세요.")
            self.root.after(0, lambda: messagebox.showwarning("데이터 필요", "무신사 데이터가 없습니다. 먼저 크롤링을 실행하세요."))
            return None, None

        expected_columns = [
            'kream_product_id', 'kream_brand_name', 'kream_product_name',
            'kream_style_code', 'kream_color', 'kream_options',
            'other_mall_product_id', 'other_mall', 'other_mall_brand_name',
            'other_mall_product_name',
        ]
        try:
            db_path = self.find_csv_file("kream_musinsa_overlapping_products_view")
            if db_path:
                self.log(f"[INFO] 매칭 DB 로드: {db_path}")
                match_df = pd.read_csv(db_path)
            else:
                db_path = utils.get_persistent_path("kream_musinsa_overlapping_products_view-generated.csv")
                self.log(f"[WARN] 매칭 DB CSV가 없어 새 파일로 시작합니다: {db_path}")
                match_df = pd.DataFrame(columns=expected_columns)

            for column in expected_columns:
                if column not in match_df.columns:
                    match_df[column] = ""

            if 'other_mall_product_id' in match_df.columns:
                match_df['other_mall_product_id'] = match_df['other_mall_product_id'].astype(str)
                match_df.loc[match_df['other_mall_product_id'] == 'nan', 'other_mall_product_id'] = ''
        except Exception as e:
            self.log(f"[ERROR] 매칭 DB 파일을 읽는 중 오류 발생: {e}")
            self.root.after(0, lambda: messagebox.showerror("파일 오류", f"매칭 DB 파일을 읽는 중 오류가 발생했습니다: {e}"))
            return None, None

        if not search_if_missing:
            self.log("[INFO] 엑셀 생성 모드: Kream ID 신규 검색을 건너뜁니다.")
            return match_df, db_path

        # --- Blacklist Logic ---
        BLACKLIST_CSV = 'kream_search_blacklist.csv'
        try:
            blacklist_df = pd.read_csv(BLACKLIST_CSV)
            blacklist_set = set(blacklist_df['musinsa_goodsNo'].astype(str))
            self.log(f"[INFO] {len(blacklist_set)}개의 상품이 포함된 검색 제외 목록을 불러왔습니다.")
        except FileNotFoundError:
            blacklist_set = set()
        # --- End Blacklist Logic ---

        new_info_added = False
        new_rows = []
        newly_blacklisted = []
        revalidate_existing_matches = os.getenv("KREAM_REVALIDATE_EXISTING_MATCHES", "").strip().lower() in {"1", "true", "yes", "on"}
        trust_existing_matches = self._trust_existing_kream_matches()

        for product in musinsa_products:
            musinsa_name = product.get("goodsName")
            goods_no = str(product.get("goodsNo"))
            if not musinsa_name or not goods_no: continue

            # --- Blacklist Check ---
            if goods_no in blacklist_set:
                self.log(f"[SKIP] 상품 ID {goods_no}는 검색 제외 목록에 있으므로 건너뜁니다.")
                continue
            # --- End Blacklist Check ---

            matched_rows = match_df[match_df['other_mall_product_name'] == musinsa_name]

            def get_and_format_kream_options(kream_id):
                print(f"크림 번호: {kream_id}, 모델번호: {product.get('serial_number')} @@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@@")
                try:
                    all_sizes = self._execute_kream_api(kream_crawler.get_all_sizes, kream_id, cookies=self.kream_cookies)
                    
                    if not all_sizes:
                        self.log(f"[INFO] Kream ID {kream_id}에 대한 사이즈 정보를 찾을 수 없습니다.")
                        return ""
                    
                    sorted_sizes = sorted(all_sizes, key=lambda s: int(s) if s.isdigit() else 999)
                    
                    formatted_sizes = [f'\"{s}\"' if s == "ONE SIZE" else str(s) for s in sorted_sizes]
                    return "{" + ",".join(formatted_sizes) + "}"
                except Exception as e:
                    self.log(f"[ERROR] Kream 옵션(사이즈) 조회 중 오류: {e}")
                    return ""

            if matched_rows.empty:
                self.log(f"[NEW] CSV에 없는 새 상품: '{musinsa_name}'. Kream ID 검색을 시작합니다.")
                try:
                    searched_id = self._execute_kream_api(
                        kream_crawler.search_product,
                        keyword=musinsa_name, brand=product.get('brandName'),
                        serial_number=product.get('serial_number'), cookies=self.kream_cookies,
                        image_url=product.get('imageUrl')
                    )
                    if searched_id:
                        if database.is_kream_id_already_used(str(searched_id), goods_no, verified_only=not trust_existing_matches):
                            self.log(f"[중복 방지] KREAM ID({searched_id})는 이미 다른 상품에 할당되어 있어 건너뜁니다. 상품: '{musinsa_name}'")
                            if goods_no not in blacklist_set:
                                newly_blacklisted.append(goods_no)
                                blacklist_set.add(goods_no)
                        else:
                            self.log(f"[NEW] ID 찾음: {searched_id}. 새 정보를 추가합니다.")
                            kream_product_url = f"https://kream.co.kr/products/{searched_id}"
                            database.cache_kream_search(goods_no, searched_id, kream_product_url)
                            self.log(f"[DB CACHE] Kream ID 캐시 저장: {goods_no} -> {searched_id}")
                            options_string = get_and_format_kream_options(searched_id)
                            new_row = {
                                'kream_product_id': searched_id, 'kream_brand_name': ' ', 'kream_product_name': ' ',
                                'kream_style_code': product.get('serial_number'), 'kream_color': '',
                                'kream_options': options_string, 'other_mall_product_id': product.get('goodsNo'),
                                'other_mall': 'musinsa', 'other_mall_brand_name': product.get('brandName'),
                                'other_mall_product_name': musinsa_name,
                            }
                            new_rows.append(new_row)
                            new_info_added = True
                    else:
                        self.log(f"[NEW] ID 검색 실패: '{musinsa_name}'. 블랙리스트에 추가합니다.")
                        if goods_no not in blacklist_set:
                            newly_blacklisted.append(goods_no)
                            blacklist_set.add(goods_no)
                except Exception as e:
                    self.log(f"[ERROR] 새 상품 KREAM ID 검색 중 오류: {e}")
            else:
                for idx, row in matched_rows.iterrows():
                    kream_id = row.get("kream_product_id")
                    cached_search = database.get_cached_kream_search(goods_no)
                    normalized_kream_id = self._normalize_kream_id_value(kream_id)
                    normalized_cached_id = self._normalize_kream_id_value(
                        cached_search.get("kream_product_id") if cached_search else None
                    )
                    has_verified_cache_for_existing_id = (
                        normalized_kream_id is not None
                        and self._is_verified_cached_kream_match(cached_search, normalized_kream_id)
                    )
                    if (
                        normalized_kream_id is not None
                        and normalized_cached_id is not None
                        and normalized_cached_id != normalized_kream_id
                        and self._is_verified_cached_kream_match(cached_search)
                    ):
                        self.log(
                            f"[DB CACHE] 검증된 KREAM ID({normalized_cached_id})가 CSV ID({normalized_kream_id})와 달라 CSV 값을 교체합니다."
                        )
                        kream_id = normalized_cached_id
                        match_df.loc[idx, "kream_product_id"] = normalized_cached_id
                        normalized_kream_id = normalized_cached_id
                        has_verified_cache_for_existing_id = True
                        new_info_added = True

                    needs_match_verification = (
                        normalized_kream_id is not None
                        and not trust_existing_matches
                        and not has_verified_cache_for_existing_id
                    )
                    is_id_missing = (
                        normalized_kream_id is None
                        or revalidate_existing_matches
                        or needs_match_verification
                    )

                    if is_id_missing:
                        if needs_match_verification and not revalidate_existing_matches:
                            self.log(f"검증되지 않은 기존 KREAM ID({kream_id})를 재검증합니다: '{musinsa_name}'")
                        else:
                            self.log(f"크림 상품 ID가 없거나 재검증이 필요하므로 검색을 시도합니다: '{musinsa_name}'")
                        try:
                            searched_id = self._execute_kream_api(
                                kream_crawler.search_product,
                                keyword=musinsa_name, brand=product.get('brandName'),
                                serial_number=product.get('serial_number'), cookies=self.kream_cookies,
                                image_url=product.get('imageUrl')
                            )
                            if searched_id:
                                if database.is_kream_id_already_used(str(searched_id), goods_no, verified_only=not trust_existing_matches):
                                    self.log(f"[중복 방지] KREAM ID({searched_id})는 이미 다른 상품에 할당되어 있어 건너뜁니다. 상품: '{musinsa_name}'")
                                    if goods_no not in blacklist_set:
                                        newly_blacklisted.append(goods_no)
                                        blacklist_set.add(goods_no)
                                else:
                                    self.log(f"[UPDATE] ID 찾음: {searched_id}. DB를 업데이트합니다.")
                                    kream_product_url = f"https://kream.co.kr/products/{searched_id}"
                                    database.cache_kream_search(goods_no, searched_id, kream_product_url)
                                    self.log(f"[DB CACHE] Kream ID 캐시 업데이트: {goods_no} -> {searched_id}")
                                    options_string = get_and_format_kream_options(searched_id)
                                    match_df.loc[idx, 'kream_product_id'] = searched_id
                                    match_df.loc[idx, 'kream_product_name'] = ''
                                    match_df.loc[idx, 'kream_style_code'] = product.get('serial_number')
                                    match_df.loc[idx, 'kream_color'] = ''
                                    match_df.loc[idx, 'kream_options'] = options_string
                                    match_df.loc[idx, 'other_mall_product_id'] = product.get('goodsNo')
                                    match_df.loc[idx, 'other_mall'] = 'musinsa'
                                    new_info_added = True
                            else:
                                self.log(f"[UPDATE] ID 검색 실패: '{musinsa_name}'. 블랙리스트에 추가합니다.")
                                if goods_no not in blacklist_set:
                                    newly_blacklisted.append(goods_no)
                                    blacklist_set.add(goods_no)
                        except Exception as e:
                            self.log(f"[ERROR] KREAM ID 검색 중 오류: {e}")
                    else:
                        self.log(f"크림 상품 ID가 이미 있으므로(ID: {kream_id}) 검색을 pass합니다: '{musinsa_name}'")
                        
                        # KREAM ID가 있어도 옵션 정보가 없는 경우, 옵션만 새로 크롤링
                        kream_options = row.get("kream_options")
                        kream_style_code = row.get("kream_style_code")
                        are_options_missing = pd.isna(kream_options) or str(kream_options).strip() in ['', '{{}}', '[]', '""']
                        are_style_missing = pd.isna(kream_style_code) or str(kream_style_code).strip() == ''

                        # 옵션 또는 스타일 코드가 비어있으면 정보 업데이트
                        if are_options_missing or are_style_missing:
                            self.log(f"[INFO] KREAM ID({kream_id})의 옵션 또는 스타일 정보가 누락되어 업데이트합니다.")
                            try:
                                kream_id_int = int(float(kream_id))
                                # 옵션 정보 가져오기
                                options_string = get_and_format_kream_options(kream_id_int)
                                
                                # 해당 행의 모든 관련 정보를 업데이트
                                match_df.loc[idx, 'kream_style_code'] = product.get('serial_number')
                                match_df.loc[idx, 'kream_options'] = options_string
                                match_df.loc[idx, 'other_mall_product_id'] = product.get('goodsNo')
                                match_df.loc[idx, 'other_mall'] = 'musinsa'
                                match_df.loc[idx, 'other_mall_brand_name'] = product.get('brandName')
                                match_df.loc[idx, 'other_mall_product_name'] = musinsa_name
                                match_df.loc[idx, 'kream_brand_name'] = ' '
                                match_df.loc[idx, 'kream_product_name'] = ' '

                                new_info_added = True
                                self.log(f"[UPDATE] KREAM ID({kream_id})의 전체 정보를 업데이트했습니다.")
                            except (ValueError, TypeError) as e:
                                self.log(f"[WARN] 정보 업데이트를 위한 KREAM ID({kream_id})가 유효하지 않습니다: {e}")

                        if pd.notna(kream_id) and str(kream_id).strip() not in ['', '0', '0.0']:
                            try:
                                kream_id_int = int(float(kream_id))
                                kream_product_url = f"https://kream.co.kr/products/{kream_id_int}"
                                database.cache_kream_search(goods_no, str(kream_id_int), kream_product_url)
                                self.log(f"[DB CACHE] 기존 Kream ID 캐시 저장: {goods_no} -> {kream_id_int}")
                            except (ValueError, TypeError):
                                self.log(f"[WARN] 캐시하려던 기존 Kream ID가 유효하지 않습니다: {kream_id}")

        if new_rows:
            new_df = pd.DataFrame(new_rows)
            match_df = pd.concat([match_df, new_df], ignore_index=True)
        
        if newly_blacklisted:
            try:
                new_blacklist_df = pd.DataFrame({'musinsa_goodsNo': newly_blacklisted})
                header = not os.path.exists(BLACKLIST_CSV)
                new_blacklist_df.to_csv(BLACKLIST_CSV, mode='a', index=False, header=header)
                self.log(f"{len(newly_blacklisted)}개 상품을 검색 제외 목록에 추가했습니다.")
            except Exception as e:
                self.log(f"[ERROR] 검색 제외 목록 파일 저장 실패: {e}")

        if new_info_added:
            try:
                self.log(f"[INFO] 원본 DB 파일 덮어쓰기: {db_path}")
                match_df.to_csv(db_path, index=False, encoding='utf-8-sig')
                self.log("[SUCCESS] Kream 정보가 업데이트된 DB 파일을 저장했습니다.")
            except Exception as e:
                self.log(f"[ERROR] 업데이트된 DB 파일 저장 실패: {e}")
        
        return match_df, db_path

    def on_click_generate_excel(self):
        self.log("[INFO] 엑셀 파일 생성을 시작합니다.")
        threading.Thread(target=self._generate_excel_task, daemon=True).start()

    def _generate_excel_task(self):
        self.log("[INFO] DB Join 방식으로 엑셀 파일 생성을 시작합니다.")
        conn = None
        try:
            margin_rate = float(self.margin_rate.get().strip()) / 100.0
            quantity_str = self.quantity.get().strip()
            deadline_str = self.deadline.get().strip()
            if not quantity_str or not deadline_str:
                self.root.after(0, lambda: messagebox.showerror("입력 오류", "수량과 입찰기한을 모두 설정해주세요."))
                return
        except (ValueError, TypeError):
            self.root.after(0, lambda: messagebox.showerror("입력 오류", "마진율과 수량에 유효한 숫자를 입력하세요."))
            return

        try:
            # 1. 매칭 CSV 파일 경로 찾기
            csv_path = self.find_csv_file("kream_musinsa_overlapping_products_view")
            if not csv_path:
                self.log("[ERROR] 매칭 DB CSV 파일을 찾을 수 없습니다.")
                self.root.after(0, lambda: messagebox.showerror("파일 없음", "매칭 DB CSV 파일을 찾을 수 없습니다."))
                return
            self.log(f"[INFO] 매칭 DB 로드: {csv_path}")

            # 2. DB 연결
            conn = database.get_db_conn()
            self.log("[INFO] 데이터베이스에 연결되었습니다.")

            # 3. CSV 데이터를 DB의 임시 테이블로 로드
            self.log("[INFO] CSV 데이터를 DB 임시 테이블 'temp_matching_data'로 로드합니다...")
            csv_df = pd.read_csv(csv_path)
            # to_sql이 kream_product_id를 float으로 읽는 것을 방지하기 위해 str으로 명시적 변환
            if 'kream_product_id' in csv_df.columns:
                csv_df['kream_product_id'] = csv_df['kream_product_id'].astype(str)
            if 'other_mall_product_id' in csv_df.columns:
                csv_df['other_mall_product_id'] = csv_df['other_mall_product_id'].astype(str)
            csv_df.to_sql('temp_matching_data', conn, if_exists='replace', index=False)
            self.log("[INFO] CSV 데이터 로드 완료.")

            # 4. JOIN 쿼리 실행 (brandName 추가)
            self.log("[INFO] DB에서 'products'와 'temp_matching_data' 테이블을 JOIN합니다.")
            join_query = """
                SELECT
                    t_match.kream_product_id,
                    t_match.kream_style_code,
                    t_match.kream_product_name,
                    t_match.other_mall_product_name,
                    p.availableSizes,
                    p.finalPrice,
                    t_match.kream_options,
                    p.brandName
                FROM
                    products AS p
                JOIN
                    temp_matching_data AS t_match ON p.goodsNo = t_match.other_mall_product_id
                WHERE
                    t_match.kream_product_id IS NOT NULL 
                    AND t_match.kream_product_id != 'nan'
                    AND t_match.kream_product_id != ''
                    AND t_match.kream_product_id != '0'
                    AND t_match.kream_product_id != '0.0'
            """
            merged_df = pd.read_sql_query(join_query, conn)
            self.log(f"[INFO] DB JOIN 완료. {len(merged_df)}개의 병합된 상품 데이터를 얻었습니다.")

            # 5. 엑셀 데이터 생성
            output_data = []
            for _, row in merged_df.iterrows():
                musinsa_price = row["finalPrice"]
                if pd.isna(musinsa_price):
                    self.log(f"[WARN] 'finalPrice'가 없어 엑셀 생성에서 제외되는 상품: {row.get('other_mall_product_name', '알 수 없음')}")
                    continue

                raw_selling_price = musinsa_price * (1 + margin_rate)
                selling_price = int(math.ceil(raw_selling_price / 1000) * 1000)
                
                brand_name = row.get("brandName", "")
                musinsa_sizes_str = row["availableSizes"]
                kream_options_str = row.get("kream_options")

                # kream_product_name이 비어있으면 other_mall_product_name(무신사 상품명)을 사용
                product_name = row["kream_product_name"]
                if not isinstance(product_name, str) or not product_name.strip():
                    product_name = row["other_mall_product_name"]

                item_data = {
                    "상품번호": str(row["kream_product_id"]).split('.')[0],
                    "모델번호": row["kream_style_code"],
                    "영문 상품명": product_name,
                    "판매 희망가": selling_price,
                    "수량": quantity_str,
                    "입찰기한": deadline_str
                }

                # Kream 옵션 파싱
                kream_sizes_list = []
                if kream_options_str and isinstance(kream_options_str, str) and kream_options_str.startswith('{') and kream_options_str.endswith('}'):
                    content = kream_options_str[1:-1]
                    kream_sizes_list = [s.strip().strip("'\"") for s in content.split(',') if s.strip()]

                # Musinsa 옵션 파싱
                musinsa_sizes_list = []
                if musinsa_sizes_str and isinstance(musinsa_sizes_str, str):
                    musinsa_sizes_list = [s.strip() for s in musinsa_sizes_str.split(',') if s.strip()]

                if not musinsa_sizes_list:
                    self.log(f"[WARN] 상품 '{item_data['영문 상품명']}'의 무신사 사이즈 정보가 없어 'ONE SIZE'로 처리합니다.")
                    found_one_size = False
                    for k_size in kream_sizes_list:
                        if "ONE SIZE" in k_size.upper():
                            new_item = item_data.copy()
                            new_item["옵션명"] = k_size
                            output_data.append(new_item)
                            found_one_size = True
                            break
                    if not found_one_size:
                        new_item = item_data.copy()
                        new_item["옵션명"] = "ONE SIZE"
                        output_data.append(new_item)
                    continue

                # --- NEW HYBRID MATCHING LOGIC ---
                matched_count = 0
                processed_k_sizes = set()

                for m_size in musinsa_sizes_list:
                    # 1. UGG 브랜드 특별 로직
                    if 'ugg' in brand_name.lower() or '어그' in brand_name:
                        is_m_size_numeric = m_size.replace('.', '', 1).isdigit()
                        if is_m_size_numeric:
                            ugg_match = find_ugg_match(m_size, kream_sizes_list)
                            if ugg_match and ugg_match not in processed_k_sizes:
                                new_item = item_data.copy()
                                new_item["옵션명"] = ugg_match
                                output_data.append(new_item)
                                self.log(f"[MATCH] UGG Musinsa '{m_size}' -> Kream '{ugg_match}'")
                                matched_count += 1
                                processed_k_sizes.add(ugg_match)
                                continue # UGG 매칭 성공 시 다음 무신사 사이즈로

                    # 2. 일반 사이즈 매칭 로직 (하이브리드)
                    for k_size in kream_sizes_list:
                        if k_size in processed_k_sizes:
                            continue

                        match_found = False
                        
                        # 전략 1: 숫자/신발 사이즈 비교
                        is_m_size_numeric = m_size.replace('.', '', 1).isdigit()
                        converted_k_size_mm = get_mm_size(k_size)
                        if is_m_size_numeric and converted_k_size_mm and m_size == converted_k_size_mm:
                            match_found = True

                        # 전략 2: 정규화된 문자열 비교 (의류 등)
                        if not match_found:
                            normalized_m_size = normalize_size(m_size)
                            normalized_k_size = normalize_size(k_size)
                            if normalized_m_size and normalized_k_size and normalized_m_size == normalized_k_size:
                                match_found = True
                        
                        if match_found:
                            final_kream_size = k_size
                            # NIKE 사이즈 강제 매칭
                            if "nike" in brand_name.lower() and k_size == "US 5.5":
                                self.log(f"[INFO] Nike 사이즈 강제 매칭: 'US 5.5' -> 'US 6' (상품: {item_data['영문 상품명']})")
                                final_kream_size = "US 6"

                            new_item = item_data.copy()
                            new_item["옵션명"] = final_kream_size
                            output_data.append(new_item)
                            matched_count += 1
                            processed_k_sizes.add(k_size)

                if matched_count == 0 and musinsa_sizes_list:
                     self.log(f"[WARN] 상품 '{item_data['영문 상품명']}'에 대해 매칭되는 사이즈가 전혀 없습니다.")

            # 중복 제거 로직 추가
            unique_output_data = []
            seen_options = set()
            for item in output_data:
                # 상품번호와 옵션명을 기준으로 중복 체크
                item_key = (item.get("상품번호"), item.get("옵션명"))
                if item_key not in seen_options:
                    unique_output_data.append(item)
                    seen_options.add(item_key)
            
            self.log(f"[INFO] 중복 제거 후 {len(unique_output_data)}개의 엑셀 행 데이터 생성을 완료했습니다.")
            self.create_excel_file(unique_output_data)

        except Exception as e:
            self.log(f"[CRITICAL ERROR] 엑셀 생성 중 심각한 오류 발생: {e}")
            import traceback
            self.log(traceback.format_exc())
            self.root.after(0, lambda: messagebox.showerror("치명적 오류", f"엑셀 생성 중 오류가 발생했습니다: {e}"))
        finally:
            if conn:
                conn.close()
                self.log("[INFO] 데이터베이스 연결을 닫았습니다.")

    def create_excel_file(self, data):
        try:
            template_path = self.find_xlsx_file("k-c46vwlba")
            if not template_path:
                messagebox.showerror("파일 없음", "엑셀 템플릿 'k-c46vwlba.xlsx'를 찾을 수 없습니다.")
                return

            base_timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            
            # 데이터를 1000개씩 나누기
            chunk_size = 1000
            data_chunks = [data[i:i + chunk_size] for i in range(0, len(data), chunk_size)]
            
            if not data_chunks:
                self.log("[WARN] 엑셀로 내보낼 데이터가 없습니다.")
                messagebox.showinfo("알림", "엑셀로 내보낼 데이터가 없습니다.")
                return

            generated_files = []
            for i, chunk in enumerate(data_chunks):
                workbook = openpyxl.load_workbook(template_path)
                sheet = workbook.active

                header_row_idx = -1
                header = []
                required_cols = ["상품번호", "옵션명", "판매 희망가", "수량", "입찰기한", "영문 상품명", "모델번호"]
                for r_idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=20)):
                    row_values = [str(cell.value).strip() if cell.value is not None else "" for cell in row]
                    if all(req_col in row_values for req_col in required_cols):
                        header_row_idx = r_idx + 1
                        header = [cell.value for cell in row]
                        break
                
                if header_row_idx == -1:
                    messagebox.showerror("템플릿 오류", "템플릿에서 필수 열을 찾을 수 없습니다.")
                    return

                data_start_row = header_row_idx + 2
                if sheet.max_row >= data_start_row:
                    sheet.delete_rows(data_start_row, sheet.max_row - data_start_row + 1)

                col_map = {str(name).strip(): c_idx for c_idx, name in enumerate(header) if name}

                for item in chunk:
                    new_row = [None] * len(header)
                    for col_name, value in item.items():
                        if col_name.strip() in col_map:
                            col_idx = col_map[col_name.strip()]
                            new_row[col_idx] = value
                    sheet.append(new_row)

                # 파일 이름 설정
                if len(data_chunks) > 1:
                    output_filename = f"output_{base_timestamp}_{i+1}.xlsx"
                else:
                    output_filename = f"output_{base_timestamp}.xlsx"
                
                workbook.save(output_filename)
                generated_files.append(output_filename)
                self.log(f"[SUCCESS] 엑셀 파일 저장 완료: {output_filename}")

            if generated_files:
                messagebox.showinfo("완료", f"엑셀 파일이 성공적으로 생성되었습니다:\n" + "\n".join(generated_files))

        except Exception as e:
            self.log(f"[ERROR] 엑셀 파일 생성 중 오류 발생: {e}")
            import traceback
            self.log(traceback.format_exc())

    def find_csv_file(self, prefix):
        """Finds the first CSV file starting with a prefix in resource/persistent directories."""
        search_dirs = []
        for candidate in [
            utils.get_persistent_path('.'),
            utils.resource_path('.'),
            os.path.abspath('.'),
        ]:
            if candidate and candidate not in search_dirs:
                search_dirs.append(candidate)

        for base_path in search_dirs:
            try:
                matches = [
                    file for file in os.listdir(base_path)
                    if file.startswith(prefix) and file.endswith('.csv')
                ]
                if not matches:
                    continue

                # Prefer the original exported matching CSV over generated fallbacks.
                matches.sort(
                    key=lambda name: (
                        'generated' in name.lower(),
                        name.lower(),
                    )
                )
                return os.path.join(base_path, matches[0])
            except FileNotFoundError:
                continue
        return None

    def find_xlsx_file(self, prefix):
        """Finds the first XLSX file starting with a prefix in the resource directory."""
        base_path = utils.resource_path('.')
        try:
            for file in os.listdir(base_path):
                if file.startswith(prefix) and file.endswith('.xlsx'):
                    return utils.resource_path(file)
        except FileNotFoundError:
            self.log(f"[ERROR] Resource directory not found: {base_path}")
        return None

    def open_selected_link(self):
        sel = self.tree.selection()
        if not sel: return

        iid = sel[0]
        parent_iid = self.tree.parent(iid)
        if not parent_iid:
            data = self.items_in_tree.get(iid)
        else:
            data = self.items_in_tree.get(parent_iid)

        if not data: return
        webbrowser.open(data.get("linkUrl") or (f"https://www.musinsa.com/products/{data.get('goodsNo')}"))

    def on_click_save_comparison_csv(self):
        self.log("[INFO] musinsa_kream_comparison.csv 파일 저장을 시작합니다.")
        threading.Thread(target=self._save_comparison_csv_task, daemon=True).start()

    def _save_comparison_csv_task(self):
        if not self.items_in_tree:
            self.log("[WARN] 저장할 데이터가 없습니다.")
            self.root.after(0, lambda: messagebox.showinfo("알림", "저장할 데이터가 없습니다."))
            return

        if self.last_crawl_source not in ['kream', 'poizon']:
            self.log("[WARN] 어떤 비교 파일을 저장할지 결정할 수 없습니다. 크림 또는 포이즌 크롤링을 먼저 실행해주세요.")
            self.root.after(0, lambda: messagebox.showwarning("작업 필요", "크림 또는 포이즌 크롤링을 먼저 실행해주세요."))
            return

        # 크롤링 소스에 따라 파일명, 필드, 링크 조회 파일 결정
        if self.last_crawl_source == 'poizon':
            fn = "musinsa_poizon_comparison.csv"
            fields = [
                "rank", "goodsName", "brandName", "size", 
                "musinsa_price", "poizon_price", "price_diff",
                "musinsa_link", "poizon_link"
            ]
            link_lookup_file = "poizon_musinsa_overlapping_products.csv"
            link_lookup_musinsa_col = 'musinsa_product_id'
            link_lookup_external_col = 'poizon_product_id'
        else: # kream
            fn = "musinsa_kream_comparison.csv"
            fields = [
                "rank", "goodsName", "brandName", "size", 
                "musinsa_price", "kream_buy_price", "kream_sell_price", "price_diff",
                "kream_total_transactions", "kream_monthly_transactions",
                "musinsa_link", "kream_link"
            ]
            link_lookup_file = self.find_csv_file("kream_musinsa_overlapping_products_view")
            link_lookup_musinsa_col = 'other_mall_product_id'
            link_lookup_external_col = 'kream_product_id'


        # 링크 조회를 위한 데이터프레임 준비
        link_map = {}
        try:
            if link_lookup_file and os.path.exists(link_lookup_file):
                lookup_df = pd.read_csv(link_lookup_file)
                # 중복 ID가 있을 경우 마지막 것만 사용
                link_map = lookup_df.drop_duplicates(subset=[link_lookup_musinsa_col], keep='last').set_index(link_lookup_musinsa_col)[link_lookup_external_col].to_dict()
            else:
                self.log(f"[WARN] 링크 조회 파일 '{link_lookup_file}'을(를) 찾을 수 없어 외부 링크는 비어있게 됩니다.")
        except Exception as e:
            self.log(f"[ERROR] 링크 조회 파일 처리 중 오류 발생: {e}")

        try:
            with open(fn, "w", newline="", encoding="utf-8-sig") as f:
                w = csv.DictWriter(f, fieldnames=fields, extrasaction='ignore')
                w.writeheader()

                for iid in self.tree.get_children(''):
                    item = self.items_in_tree.get(iid)
                    if not item: continue

                    goods_no = str(item.get("goodsNo"))
                    musinsa_link = item.get("linkUrl") or f"https://www.musinsa.com/products/{goods_no}"
                    
                    external_link = ""
                    # Prioritize link from crawl result
                    if self.last_crawl_source == 'poizon' and item.get('poizon_link'):
                        external_link = item.get('poizon_link')

                    if not external_link:
                        link_id = link_map.get(int(goods_no)) # CSV에서 ID가 float으로 읽힐 수 있어 int로 변환
                        if pd.notna(link_id):
                            if self.last_crawl_source == 'poizon':
                                try:
                                    spu_id = int(float(link_id))
                                    search_id = uuid.uuid4()
                                    external_link = f"https://seller.poizon.com/main/bid/freely/operateBidList?operateType=add&sourceName=1&globalSpuIds={spu_id}&searchId={search_id}"
                                except (ValueError, TypeError):
                                    external_link = ""
                            else: # kream
                                try:
                                    external_link = f"https://kream.co.kr/products/{int(float(link_id))}"
                                except (ValueError, TypeError):
                                    external_link = ""

                    children = self.tree.get_children(iid)
                    if not children:
                        row = {
                            'rank': item.get('rank'),
                            'goodsName': item.get('goodsName'),
                            'brandName': item.get('brandName'),
                            'musinsa_price': item.get('finalPrice'),
                            'musinsa_link': musinsa_link,
                        }
                        if self.last_crawl_source == 'poizon':
                            row['poizon_link'] = external_link
                        else:
                            row['kream_link'] = external_link
                            row['kream_total_transactions'] = item.get('kream_total_transactions')
                            row['kream_monthly_transactions'] = item.get('kream_monthly_transactions')
                        w.writerow(row)
                    else:
                        for child_iid in children:
                            child_values = self.tree.item(child_iid, 'values')
                            row = {
                                'rank': item.get('rank'),
                                'goodsName': item.get('goodsName'),
                                'brandName': item.get('brandName'),
                                'size': child_values[3],
                                'musinsa_price': utils.parse_krw(child_values[4]),
                                'price_diff': utils.parse_krw(child_values[7]),
                                'musinsa_link': musinsa_link,
                            }
                            if self.last_crawl_source == 'poizon':
                                row['poizon_price'] = utils.parse_krw(child_values[5])
                                row['poizon_link'] = external_link
                            else: # kream
                                row['kream_buy_price'] = utils.parse_krw(child_values[5])
                                row['kream_sell_price'] = utils.parse_krw(child_values[6])
                                row['kream_link'] = external_link
                                row['kream_total_transactions'] = item.get('kream_total_transactions')
                                row['kream_monthly_transactions'] = item.get('kream_monthly_transactions')
                            w.writerow(row)
            
            self.log(f"[SUCCESS] CSV 저장 완료: {fn}")
            self.root.after(0, lambda: messagebox.showinfo("완료", f"CSV 저장이 완료되었습니다: {fn}"))

        except Exception as e:
            self.log(f"[ERROR] CSV 저장 실패: {e}")
            import traceback
            self.log(traceback.format_exc())
            self.root.after(0, lambda: messagebox.showerror("오류", f"CSV 저장 중 오류가 발생했습니다: {e}"))

    def open_selected_detail(self):
        sel = self.tree.selection()
        if not sel:
            messagebox.showinfo("안내","행을 선택하세요.")
            return
        
        iid = sel[0]
        parent_iid = self.tree.parent(iid)
        data = self.items_in_tree.get(parent_iid or iid)

        if not data:
            messagebox.showinfo("안내","데이터가 없습니다.")
            return
        
        goods_no = str(data.get("goodsNo"))
        kream_sales_data = None

        if 'kream_total_transactions' in data and data['kream_total_transactions'] is not None:
            kream_sales_data = {
                "total": data['kream_total_transactions'],
                "monthly": data.get('kream_monthly_transactions', 0),
                "items": []
            }
        elif goods_no:
            pass

        DetailWindow(self.root, data, self, kream_sales_data)

    def _ensure_promo_groups(self) -> List[dict]:
        if self._promo_groups_cache is not None:
            return self._promo_groups_cache
        try:
            ev = crawler.fetch_payment_events(self.musinsa_cookie_header)
            self._promo_groups_cache = crawler.build_promotion_groups(ev)
        except Exception as e:
            self.log(f"[WARN] 결제 프로모션 로딩 실패: {e}")
            self._promo_groups_cache = []
        return self._promo_groups_cache

    def open_monitoring_window(self):
        if not self.musinsa_cookie_header:
            messagebox.showwarning("쿠키 필요", "재입고 모니터링을 위해 먼저 무신사 쿠키를 얻어주세요.")
            return
        MonitoringWindow(self.root, self)

    def open_price_monitoring_window(self):
        if not self.musinsa_cookie_header:
            messagebox.showwarning("쿠키 필요", "가격변동 모니터링을 위해 먼저 무신사 쿠키를 얻어주세요.")
            return
        PriceMonitoringWindow(self.root, self)

    def open_purchase_limit_monitoring_window(self):
        if not self.musinsa_cookie_header:
            messagebox.showwarning("쿠키 필요", "구매제한 모니터링을 위해 먼저 무신사 쿠키를 얻어주세요.")
            return
        PurchaseLimitMonitoringWindow(self.root, self)

    def open_proxy_manager(self):
        ProxyManagerWindow(self.root, self)


class ProxyManagerWindow(tk.Toplevel):
    def __init__(self, parent, app_instance):
        super().__init__(parent)
        self.app = app_instance
        self.title("프록시 IP 관리")
        self.geometry("500x400")

        self.proxy_file = "proxies.txt"

        self.build_widgets()
        self.load_proxies()

    def build_widgets(self):
        top_frame = ttk.Frame(self, padding=10)
        top_frame.pack(fill="x", side="top")

        top_frame.columnconfigure(1, weight=1) # Make the entry column expandable

        ttk.Label(top_frame, text="프록시 주소:").grid(row=0, column=0, sticky="w", padx=(0, 5))
        
        self.proxy_entry = ttk.Entry(top_frame)
        self.proxy_entry.grid(row=0, column=1, sticky="ew")

        # Changed text to "등록"
        add_button = ttk.Button(top_frame, text="등록", command=self.add_proxy)
        add_button.grid(row=0, column=2, sticky="e", padx=(5, 0))

        list_frame = ttk.Frame(self, padding=10)
        list_frame.pack(fill="both", expand=True)

        self.proxy_listbox = tk.Listbox(list_frame)
        self.proxy_listbox.pack(side="left", fill="both", expand=True)

        scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=self.proxy_listbox.yview)
        scrollbar.pack(side="right", fill="y")
        self.proxy_listbox.config(yscrollcommand=scrollbar.set)

        bottom_frame = ttk.Frame(self, padding=10)
        bottom_frame.pack(fill="x")

        delete_button = ttk.Button(bottom_frame, text="선택 삭제", command=self.delete_proxy)
        delete_button.pack(side="right")

    def load_proxies(self):
        self.proxy_listbox.delete(0, tk.END)
        try:
            with open(self.proxy_file, "r") as f:
                proxies = [line.strip() for line in f if line.strip()]
                for proxy in proxies:
                    self.proxy_listbox.insert(tk.END, proxy)
            self.app.log(f"[Proxy] {len(proxies)}개의 프록시를 불러왔습니다.")
        except FileNotFoundError:
            self.app.log("[Proxy] 프록시 파일(proxies.txt)이 없습니다. 새로 생성됩니다.")
            # Create the file if it doesn't exist
            with open(self.proxy_file, "w") as f:
                pass

    def save_proxies(self):
        proxies = self.proxy_listbox.get(0, tk.END)
        with open(self.proxy_file, "w") as f:
            for proxy in proxies:
                f.write(proxy + "\n")
        # Also update the proxy manager in kream_crawler
        if hasattr(kream_crawler, 'proxy_manager'):
            kream_crawler.proxy_manager.set_proxies(list(proxies))
        self.app.log(f"[Proxy] {len(proxies)}개의 프록시를 파일에 저장했습니다.")

    def add_proxy(self):
        proxy = self.proxy_entry.get().strip()
        if proxy:
            # Simple validation
            if not (proxy.startswith("http://") or proxy.startswith("https://")):
                messagebox.showerror("오류", "프록시 주소는 http:// 또는 https:// 로 시작해야 합니다.", parent=self)
                return
            
            current_proxies = self.proxy_listbox.get(0, tk.END)
            if proxy in current_proxies:
                messagebox.showwarning("알림", "이미 등록된 프록시 주소입니다.", parent=self)
                return

            self.proxy_listbox.insert(tk.END, proxy)
            self.proxy_entry.delete(0, tk.END)
            self.save_proxies()
        else:
            messagebox.showwarning("입력 필요", "프록시 주소를 입력하세요.", parent=self)

    def delete_proxy(self):
        selected_indices = self.proxy_listbox.curselection()
        if not selected_indices:
            messagebox.showwarning("선택 필요", "삭제할 프록시를 선택하세요.", parent=self)
            return

        # Delete from bottom to top to avoid index shifting issues
        for i in reversed(selected_indices):
            self.proxy_listbox.delete(i)
        
        self.save_proxies()


class MonitoringWindow(tk.Toplevel):
    def __init__(self, parent, app_instance):
        super().__init__(parent)
        self.app = app_instance
        self.root = parent
        self.title("재입고 모니터링")
        self.geometry("800x600")

        self.monitoring_items = {}  # goodsNo -> {iid, data, last_sizes}
        self.monitoring_thread = None
        self.monitoring_active = threading.Event()
        self.adding_items_counter = 0

        self.build_widgets()
        self.protocol("WM_DELETE_WINDOW", self.on_close)

    def build_widgets(self):
        # Top frame for controls
        control_frame = ttk.Frame(self, padding=10)
        control_frame.pack(fill="x", side="top")

        ttk.Label(control_frame, text="상품 ID:").pack(side="left", padx=(0, 5))
        self.goods_id_entry = ttk.Entry(control_frame, width=15)
        self.goods_id_entry.pack(side="left", padx=5)
        
        self.add_button = ttk.Button(control_frame, text="추가", command=self.add_item)
        self.add_button.pack(side="left", padx=5)

        self.delete_button = ttk.Button(control_frame, text="선택 삭제", command=self.delete_item)
        self.delete_button.pack(side="left", padx=5)
        self.add_current_button = ttk.Button(control_frame, text="현재 크롤링 상품 모니터링", command=self.add_current_products)
        self.add_current_button.pack(side="left", padx=5)

        ttk.Label(control_frame, text="주기(초):").pack(side="left", padx=(10, 5))
        self.add_current_button = ttk.Button(control_frame, text="?꾩옱 ?щ·留??곹뭹 紐⑤땲?곕쭅", command=self.add_current_products)
        self.add_current_button.pack(side="left", padx=5)

        self.add_current_button = ttk.Button(control_frame, text="?꾩옱 ?щ·留??곹뭹 紐⑤땲?곕쭅", command=self.add_current_products)
        self.add_current_button.pack(side="left", padx=5)

        self.add_current_button = ttk.Button(control_frame, text="?꾩옱 ?щ·留??곹뭹 紐⑤땲?곕쭅", command=self.add_current_products)
        self.add_current_button.pack(side="left", padx=5)

        self.add_current_button = ttk.Button(control_frame, text="현재 크롤링 상품 모니터링", command=self.add_current_products)
        self.add_current_button.pack(side="left", padx=5)

        self.add_current_button = ttk.Button(control_frame, text="현재 크롤링 상품 모니터링", command=self.add_current_products)
        self.add_current_button.pack(side="left", padx=5)

        self.interval_entry = ttk.Entry(control_frame, width=8)
        self.interval_entry.insert(0, "180")
        self.interval_entry.pack(side="left", padx=5)

        self.start_button = ttk.Button(control_frame, text="모니터링 시작", command=self.start_monitoring)
        self.start_button.pack(side="left", padx=5)

        self.stop_button = ttk.Button(control_frame, text="모니터링 중지", command=self.stop_monitoring, state="disabled")
        self.stop_button.pack(side="left", padx=5)

        # Treeview for monitored items
        self.add_current_button = ttk.Button(control_frame, text="?꾩옱 ?щ·留??곹뭹 紐⑤땲?곕쭅", command=self.add_current_products)
        self.add_current_button.pack(side="left", padx=5)

        self.add_current_button = ttk.Button(control_frame, text="현재 크롤링 상품 모니터링", command=self.add_current_products)
        self.add_current_button.pack(side="left", padx=5)

        tree_frame = ttk.Frame(self, padding=10)
        tree_frame.pack(fill="both", expand=True)

        cols = ("id", "brand", "name", "status", "sizes")
        self.tree = ttk.Treeview(tree_frame, columns=cols, show="headings")
        
        headings = {"id": "상품 ID", "brand": "브랜드", "name": "상품명", "status": "상태", "sizes": "재고 사이즈"}
        widths = {"id": 80, "brand": 100, "name": 250, "status": 100, "sizes": 150}
        for col in cols:
            self.tree.heading(col, text=headings.get(col, col))
            self.tree.column(col, width=widths.get(col, 100), anchor="w")

        vsb = ttk.Scrollbar(tree_frame, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        self.tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")

        self.tree.tag_configure("RESTOCKED", background="lightgreen")
        self.tree.tag_configure("SOLD_OUT", foreground="gray")

    def add_item(self):
        goods_no = self.goods_id_entry.get().strip()
        if not goods_no.isdigit():
            messagebox.showerror("입력 오류", "유효한 상품 ID(숫자)를 입력하세요.", parent=self)
            return
        
        if goods_no in self.monitoring_items:
            messagebox.showinfo("알림", "이미 모니터링 목록에 있는 상품입니다.", parent=self)
            return

        def _task():
            self.app.log(f"[Monitor] 상품 정보 및 초기 재고 조회 중: {goods_no}")
            try:
                s = crawler.build_session(self.app.musinsa_cookie_header)
                details_map = crawler.fetch_details(s, [goods_no])
                if not details_map:
                    self.app.log(f"[Monitor] 상품 정보를 찾을 수 없습니다: {goods_no}")
                    self.root.after(0, lambda: messagebox.showerror("조회 실패", f"상품 ID {goods_no} 정보를 찾을 수 없습니다.", parent=self))
                    return

                detail_meta = details_map.get(goods_no)
                brand = detail_meta.get("brandName", "N/A")
                name = detail_meta.get("goodsName", "N/A")

                # Fetch initial stock
                opts = crawler.fetch_options_any(s, goods_no)
                initial_sizes = set(crawler.extract_available_sizes(opts))
                self.app.log(f"[Monitor] 초기 재고 확인: {goods_no}, 사이즈: {', '.join(initial_sizes) if initial_sizes else '품절'}")

                self.root.after(0, self._add_item_to_tree, goods_no, brand, name, detail_meta, initial_sizes)

            except Exception as e:
                self.app.log(f"[Monitor] 상품 정보 조회 실패: {e}")
                self.root.after(0, lambda: messagebox.showerror("조회 오류", f"상품 정보 조회 중 오류 발생: {e}", parent=self))

        threading.Thread(target=_task, daemon=True).start()

    def _add_item_to_tree(self, goods_no, brand, name, data, initial_sizes):
        sizes_str = ", ".join(sorted(initial_sizes, key=lambda x: (not x.isdigit(), int(x) if x.isdigit() else x)))
        status = "모니터링 시작" if initial_sizes else "품절 상태로 시작"
        
        iid = self.tree.insert("", "end", values=(goods_no, brand, name, status, sizes_str))
        
        # Store the initial sizes in last_sizes
        self.monitoring_items[goods_no] = {"iid": iid, "data": data, "last_sizes": initial_sizes}
        
        self.goods_id_entry.delete(0, "end")
        self.app.log(f"[Monitor] 상품 추가됨: {name} ({goods_no})")

    def delete_item(self):
        selected = self.tree.selection()
        if not selected:
            return
        
        for iid in selected:
            values = self.tree.item(iid, "values")
            goods_no = values[0]
            if goods_no in self.monitoring_items:
                del self.monitoring_items[goods_no]
                self.tree.delete(iid)
                self.app.log(f"[Monitor] 상품 삭제됨: {goods_no}")

    def add_current_products(self):
        if not self.app.items_in_tree:
            messagebox.showwarning("?뚮┝", "硫붿씤 紐⑸줉???곹뭹???놁뒿?덈떎.", parent=self)
            return

        self.app.log("[Monitor] ?꾩옱 ?щ·留곷맂 紐⑤뱺 ?곹뭹??紐⑤땲?곕쭅 紐⑸줉??異붽??⑸땲??")
        items_to_add = list(self.app.items_in_tree.values())
        self.add_current_button.config(state="disabled")
        self.adding_items_counter = len(items_to_add)

        for item_data in items_to_add:
            goods_no = item_data.get("goodsNo")
            if not goods_no or goods_no in self.monitoring_items:
                self.item_add_completed()
                continue
            self.add_item_from_data(item_data)

    def add_item_from_data(self, item_data):
        goods_no = item_data.get("goodsNo")
        brand = item_data.get("brandName", "N/A")
        name = item_data.get("goodsName", "N/A")
        available_sizes = set(
            s.strip() for s in str(item_data.get("availableSizes") or "").split(",") if s.strip()
        )
        self.root.after(0, self._add_item_to_tree, goods_no, brand, name, item_data, available_sizes)
        self.item_add_completed()

    def item_add_completed(self):
        self.adding_items_counter -= 1
        if self.adding_items_counter <= 0:
            self.add_current_button.config(state="normal")
            self.app.log("[Monitor] 紐⑤뱺 ?곹뭹 異붽? ?꾨즺.")

    def add_current_products(self):
        if not self.app.items_in_tree:
            messagebox.showwarning("?뚮┝", "硫붿씤 紐⑸줉???곹뭹???놁뒿?덈떎.", parent=self)
            return

        self.app.log("[LimitMon] ?꾩옱 ?щ·留곷맂 紐⑤뱺 ?곹뭹??紐⑤땲?곕쭅 紐⑸줉??異붽??⑸땲??")
        items_to_add = list(self.app.items_in_tree.values())
        self.add_current_button.config(state="disabled")
        self.adding_items_counter = len(items_to_add)

        for item_data in items_to_add:
            goods_no = item_data.get("goodsNo")
            if not goods_no or goods_no in self.monitoring_items:
                self.item_add_completed()
                continue
            self.add_item_from_data(item_data)

    def add_item_from_data(self, item_data):
        goods_no = item_data.get("goodsNo")

        def _task():
            try:
                s = crawler.build_session(self.app.musinsa_cookie_header)
                limit_info = crawler.fetch_purchase_limit_info(s, goods_no)
                if limit_info:
                    brand = limit_info.get("brandName", item_data.get("brandName", "N/A"))
                    name = limit_info.get("goodsName", item_data.get("goodsName", "N/A"))
                    is_limited = limit_info.get("isLimitedQuantity")
                    self.root.after(0, self._add_item_to_tree, goods_no, brand, name, is_limited, limit_info)
                else:
                    self.app.log(f"[LimitMon] ?곹뭹 ?뺣낫瑜?李얠쓣 ???놁뼱 嫄대꼫?곷땲?? {goods_no}")
            except Exception as e:
                self.app.log(f"[LimitMon] ?곹뭹 ?붽? 議고쉶 ?ㅽ뙣: {goods_no} / {e}")
            finally:
                self.item_add_completed()

        threading.Thread(target=_task, daemon=True).start()

    def item_add_completed(self):
        self.adding_items_counter -= 1
        if self.adding_items_counter <= 0:
            self.add_current_button.config(state="normal")
            self.app.log("[LimitMon] 紐⑤뱺 ?곹뭹 異붽? ?꾨즺.")

    def add_current_products(self):
        if not self.app.items_in_tree:
            messagebox.showwarning("?뚮┝", "硫붿씤 紐⑸줉???곹뭹???놁뒿?덈떎.", parent=self)
            return

        self.app.log("[LimitMon] ?꾩옱 ?щ·留곷맂 紐⑤뱺 ?곹뭹??紐⑤땲?곕쭅 紐⑸줉??異붽??⑸땲??")
        items_to_add = list(self.app.items_in_tree.values())
        self.add_current_button.config(state="disabled")
        self.adding_items_counter = len(items_to_add)

        for item_data in items_to_add:
            goods_no = item_data.get("goodsNo")
            if not goods_no or goods_no in self.monitoring_items:
                self.item_add_completed()
                continue
            self.add_item_from_data(item_data)

    def add_item_from_data(self, item_data):
        goods_no = item_data.get("goodsNo")

        def _task():
            try:
                s = crawler.build_session(self.app.musinsa_cookie_header)
                limit_info = crawler.fetch_purchase_limit_info(s, goods_no)
                if limit_info:
                    brand = limit_info.get("brandName", item_data.get("brandName", "N/A"))
                    name = limit_info.get("goodsName", item_data.get("goodsName", "N/A"))
                    is_limited = limit_info.get("isLimitedQuantity")
                    self.root.after(0, self._add_item_to_tree, goods_no, brand, name, is_limited, limit_info)
                else:
                    self.app.log(f"[LimitMon] ?곹뭹 ?뺣낫瑜?李얠쓣 ???놁뼱 嫄대꼫?곷땲?? {goods_no}")
            except Exception as e:
                self.app.log(f"[LimitMon] ?곹뭹 ?붽? 議고쉶 ?ㅽ뙣: {goods_no} / {e}")
            finally:
                self.item_add_completed()

        threading.Thread(target=_task, daemon=True).start()

    def item_add_completed(self):
        self.adding_items_counter -= 1
        if self.adding_items_counter <= 0:
            self.add_current_button.config(state="normal")
            self.app.log("[LimitMon] 紐⑤뱺 ?곹뭹 異붽? ?꾨즺.")

    def start_monitoring(self):
        if not self.monitoring_items:
            messagebox.showwarning("알림", "모니터링할 상품을 먼저 추가해주세요.", parent=self)
            return
        
        try:
            interval_sec = int(self.interval_entry.get().strip())
            if interval_sec <= 0:
                raise ValueError
        except ValueError:
            messagebox.showerror("입력 오류", "모니터링 주기에 유효한 양의 정수를 입력하세요.", parent=self)
            return

        if self.monitoring_thread and self.monitoring_thread.is_alive():
            messagebox.showinfo("알림", "이미 모니터링이 실행 중입니다.", parent=self)
            return

        self.monitoring_active.set()
        self.monitoring_thread = threading.Thread(target=self._monitoring_loop, args=(interval_sec,), daemon=True)
        self.monitoring_thread.start()

        self.start_button.config(state="disabled")
        self.stop_button.config(state="normal")
        self.add_button.config(state="disabled")
        self.delete_button.config(state="disabled")
        self.add_current_button.config(state="disabled")
        self.add_current_button.config(state="disabled")
        self.add_current_button.config(state="disabled")
        self.app.log(f"[Monitor] 재입고 모니터링을 시작합니다. (주기: {interval_sec}초)")

    def stop_monitoring(self):
        self.monitoring_active.clear()
        self.start_button.config(state="normal")
        self.stop_button.config(state="disabled")
        self.add_button.config(state="normal")
        self.delete_button.config(state="normal")
        self.add_current_button.config(state="normal")
        self.add_current_button.config(state="normal")
        self.add_current_button.config(state="normal")
        self.app.log("[Monitor] 재입고 모니터링을 중지합니다.")
        for goods_no, item in self.monitoring_items.items():
            self.update_tree_item(goods_no, "중지됨", item.get("last_sizes", set()))


    def _monitoring_loop(self, interval_sec):
        s = crawler.build_session(self.app.musinsa_cookie_header)
        webhook_url = self.app.discord_webhook_url_entry.get().strip()

        while self.monitoring_active.is_set():
            self.app.log("[Monitor] 재고 상태 확인 시작...")
            
            current_items = list(self.monitoring_items.items())

            for goods_no, item_info in current_items:
                if not self.monitoring_active.is_set():
                    break
                
                try:
                    self.app.log(f"[Monitor] 상품 확인: {item_info['data'].get('goodsName')}")
                    opts = crawler.fetch_options_any(s, goods_no)
                    available_sizes = set(crawler.extract_available_sizes(opts))
                    
                    last_sizes = item_info.get("last_sizes", set())
                    
                    status = "품절" if not available_sizes else "재고 있음"
                    tags = ()
                    notification_needed = False

                    if not last_sizes and available_sizes:
                        status = "✨ 재입고!"
                        tags = ("RESTOCKED",)
                        self.app.log(f"[Monitor] 재입고 감지! 상품: {item_info['data'].get('goodsName')}, 사이즈: {', '.join(available_sizes)}")
                        notification_needed = True
                    elif last_sizes and not available_sizes:
                        status = "품절"
                        tags = ("SOLD_OUT",)
                    elif available_sizes:
                        newly_stocked = available_sizes - last_sizes
                        if newly_stocked:
                            status = f"✨ {', '.join(newly_stocked)} 재입고!"
                            tags = ("RESTOCKED",)
                            self.app.log(f"[Monitor] 부분 재입고 감지! 상품: {item_info['data'].get('goodsName')}, 새 사이즈: {', '.join(newly_stocked)}")
                            notification_needed = True

                    if notification_needed and webhook_url:
                        message = (
                            f"**🚨 재입고 알림 🚨**\n\n"
                            f"**상품 ID:** {goods_no}\n"
                            f"**브랜드:** {item_info['data'].get('brandName')}\n"
                            f"**상품명:** {item_info['data'].get('goodsName')}\n"
                            f"**재입고 사이즈:** {', '.join(sorted(available_sizes))}\n"
                            f"[상품 바로가기](https://www.musinsa.com/app/goods/{goods_no})"
                        )
                        webhook.send_discord_notification(webhook_url, message)

                    self.root.after(0, self.update_tree_item, goods_no, status, available_sizes, tags)
                    
                    self.monitoring_items[goods_no]["last_sizes"] = available_sizes

                except Exception as e:
                    self.app.log(f"[Monitor] 상품 {goods_no} 확인 중 오류: {e}")
                    self.root.after(0, self.update_tree_item, goods_no, "오류", set())
                
                if self.monitoring_active.is_set():
                    time.sleep(random.uniform(3, 7))

            self.app.log(f"[Monitor] 모든 상품 확인 완료. {interval_sec}초 후 다시 시작합니다.")
            self.monitoring_active.wait(interval_sec)

        self.app.log("[Monitor] 모니터링 루프가 종료되었습니다.")


    def update_tree_item(self, goods_no, status, sizes, tags=()):
        if goods_no in self.monitoring_items:
            iid = self.monitoring_items[goods_no]["iid"]
            if self.tree.exists(iid):
                self.tree.item(iid, values=(
                    goods_no,
                    self.tree.item(iid, "values")[1],
                    self.tree.item(iid, "values")[2],
                    status,
                    ", ".join(sorted(sizes, key=lambda x: (not x.isdigit(), int(x) if x.isdigit() else x)))
                ), tags=tags)

    # Keep these definitions near the end of the class so they override any earlier duplicates.
    def add_current_products(self):
        if not self.app.items_in_tree:
            messagebox.showwarning("?뚮┝", "硫붿씤 紐⑸줉???곹뭹???놁뒿?덈떎.", parent=self)
            return

        self.app.log("[Monitor] ?꾩옱 ?щ·留곷맂 紐⑤뱺 ?곹뭹??紐⑤땲?곕쭅 紐⑸줉??異붽??⑸땲??")
        items_to_add = list(self.app.items_in_tree.values())
        self.add_current_button.config(state="disabled")
        self.adding_items_counter = len(items_to_add)

        for item_data in items_to_add:
            goods_no = item_data.get("goodsNo")
            if not goods_no or goods_no in self.monitoring_items:
                self.item_add_completed()
                continue
            self.add_item_from_data(item_data)

    def add_item_from_data(self, item_data):
        goods_no = item_data.get("goodsNo")
        brand = item_data.get("brandName", "N/A")
        name = item_data.get("goodsName", "N/A")
        available_sizes = set(s.strip() for s in str(item_data.get("availableSizes") or "").split(",") if s.strip())
        self.root.after(0, self._add_item_to_tree, goods_no, brand, name, item_data, available_sizes)
        self.item_add_completed()

    def item_add_completed(self):
        self.adding_items_counter -= 1
        if self.adding_items_counter <= 0:
            self.add_current_button.config(state="normal")
            self.app.log("[Monitor] 紐⑤뱺 ?곹뭹 異붽? ?꾨즺.")

    def on_close(self):
        if self.monitoring_thread and self.monitoring_thread.is_alive():
            self.stop_monitoring()
        self.destroy()

class PriceMonitoringWindow(tk.Toplevel):
    def __init__(self, parent, app_instance):
        super().__init__(parent)
        self.app = app_instance
        self.root = parent
        self.title("가격 변동 모니터링")
        self.geometry("900x600")

        self.monitoring_items = {}  # goodsNo -> {iid, data, original_price, last_price}
        self.monitoring_thread = None
        self.monitoring_active = threading.Event()
        self.adding_items_counter = 0

        self.build_widgets()
        self.protocol("WM_DELETE_WINDOW", self.on_close)

    def build_widgets(self):
        control_frame = ttk.Frame(self, padding=10)
        control_frame.pack(fill="x", side="top")

        ttk.Label(control_frame, text="상품 ID:").pack(side="left", padx=(0, 5))
        self.goods_id_entry = ttk.Entry(control_frame, width=15)
        self.goods_id_entry.pack(side="left", padx=5)
        
        self.add_button = ttk.Button(control_frame, text="추가", command=self.add_item)
        self.add_button.pack(side="left", padx=5)

        self.delete_button = ttk.Button(control_frame, text="선택 삭제", command=self.delete_item)
        self.delete_button.pack(side="left", padx=5)

        self.add_current_button = ttk.Button(control_frame, text="현재 크롤링 상품 모니터링", command=self.add_current_products)
        self.add_current_button.pack(side="left", padx=5)

        ttk.Label(control_frame, text="주기(초):").pack(side="left", padx=(10, 5))
        self.interval_entry = ttk.Entry(control_frame, width=8)
        self.interval_entry.insert(0, "300")
        self.interval_entry.pack(side="left", padx=5)

        self.start_button = ttk.Button(control_frame, text="모니터링 시작", command=self.start_monitoring)
        self.start_button.pack(side="left", padx=5)

        self.stop_button = ttk.Button(control_frame, text="모니터링 중지", command=self.stop_monitoring, state="disabled")
        self.stop_button.pack(side="left", padx=5)

        tree_frame = ttk.Frame(self, padding=10)
        tree_frame.pack(fill="both", expand=True)

        cols = ("id", "brand", "name", "original_price", "current_price", "status")
        self.tree = ttk.Treeview(tree_frame, columns=cols, show="headings")
        
        headings = {"id": "상품 ID", "brand": "브랜드", "name": "상품명", "original_price": "최초가", "current_price": "현재가", "status": "상태"}
        widths = {"id": 80, "brand": 100, "name": 250, "original_price": 80, "current_price": 80, "status": 120}
        for col in cols:
            self.tree.heading(col, text=headings.get(col, col))
            self.tree.column(col, width=widths.get(col, 100), anchor="w")

        vsb = ttk.Scrollbar(tree_frame, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        self.tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")

        self.tree.tag_configure("PRICE_UP", foreground="red")
        self.tree.tag_configure("PRICE_DOWN", foreground="blue")

    def add_item(self):
        goods_no = self.goods_id_entry.get().strip()
        if not goods_no.isdigit():
            messagebox.showerror("입력 오류", "유효한 상품 ID(숫자)를 입력하세요.", parent=self)
            return
        
        if goods_no in self.monitoring_items:
            messagebox.showinfo("알림", "이미 모니터링 목록에 있는 상품입니다.", parent=self)
            return

        def _task():
            self.app.log(f"[PriceMon] 상품 정보 및 최초 가격 조회 중: {goods_no}")
            try:
                s = crawler.build_session(self.app.musinsa_cookie_header)
                details_map = crawler.fetch_details(s, [goods_no])
                if not details_map:
                    self.app.log(f"[PriceMon] 상품 정보를 찾을 수 없습니다: {goods_no}")
                    self.root.after(0, lambda: messagebox.showerror("조회 실패", f"상품 ID {goods_no} 정보를 찾을 수 없습니다.", parent=self))
                    return
                
                detail_meta = details_map.get(goods_no)
                link_url = f"https://www.musinsa.com/app/goods/{goods_no}"
                processed_data = crawler.process_one(
                    s,
                    self.app.musinsa_cookie_header,
                    goods_no,
                    link_url,
                    detail_meta,
                    auth_payload=self.app.musinsa_auth_payload,
                )
                
                initial_price = processed_data.get("finalPrice")
                
                brand = detail_meta.get("brandName", "N/A")
                name = detail_meta.get("goodsName", "N/A")

                self.root.after(0, self._add_item_to_tree, goods_no, brand, name, detail_meta, initial_price)

            except Exception as e:
                self.app.log(f"[PriceMon] 상품 정보 조회 실패: {e}")
                self.root.after(0, lambda: messagebox.showerror("조회 오류", f"상품 정보 조회 중 오류 발생: {e}", parent=self))

        threading.Thread(target=_task, daemon=True).start()

    def _add_item_to_tree(self, goods_no, brand, name, data, original_price):
        # Display "확인 불가" if price is None, otherwise display the price.
        price_str = f"{utils.krw(original_price)}원" if original_price is not None else "확인 불가"
        
        iid = self.tree.insert("", "end", values=(
            goods_no, 
            brand, 
            name, 
            price_str, # 최초가
            price_str, # 현재가
            "모니터링 시작"
        ))
        
        self.monitoring_items[goods_no] = {
            "iid": iid, 
            "data": data, 
            "original_price": original_price, 
            "last_price": original_price
        }
        self.app.log(f"[Price Monitor] 상품 추가됨: {name} ({goods_no}), 최초가: {price_str}")

    def delete_item(self):
        selected = self.tree.selection()
        if not selected: return
        for iid in selected:
            goods_no = self.tree.item(iid, "values")[0]
            if goods_no in self.monitoring_items:
                del self.monitoring_items[goods_no]
                self.tree.delete(iid)
                self.app.log(f"[PriceMon] 상품 삭제됨: {goods_no}")

    def add_current_products(self):
        if not self.app.items_in_tree:
            messagebox.showwarning("알림", "메인 목록에 상품이 없습니다.", parent=self)
            return

        self.app.log("[Price Monitor] 현재 크롤링된 모든 상품을 모니터링 목록에 추가합니다.")
        
        items_to_add = list(self.app.items_in_tree.values())
        
        self.add_current_button.config(state="disabled")
        self.adding_items_counter = len(items_to_add)

        for item_data in items_to_add:
            goods_no = item_data.get("goodsNo")
            if not goods_no or goods_no in self.monitoring_items:
                self.item_add_completed()
                continue
            
            self.add_item_from_data(item_data)

    def add_item_from_data(self, item_data):
        goods_no = item_data.get("goodsNo")
        brand = item_data.get("brandName", "N/A")
        name = item_data.get("goodsName", "N/A")
        
        initial_price = item_data.get("finalPrice")

        self.root.after(0, self._add_item_to_tree, goods_no, brand, name, item_data, initial_price)
        
        self.item_add_completed()

    def item_add_completed(self):
        self.adding_items_counter -= 1
        if self.adding_items_counter <= 0:
            self.add_current_button.config(state="normal")
            self.app.log("[Price Monitor] 모든 상품 추가 완료.")

    def start_monitoring(self):
        if not self.monitoring_items:
            messagebox.showwarning("알림", "모니터링할 상품을 먼저 추가해주세요.", parent=self)
            return
        
        try:
            interval_sec = int(self.interval_entry.get().strip())
            if interval_sec <= 0:
                raise ValueError
        except ValueError:
            messagebox.showerror("입력 오류", "모니터링 주기에 유효한 양의 정수를 입력하세요.", parent=self)
            return

        if self.monitoring_thread and self.monitoring_thread.is_alive():
            messagebox.showinfo("알림", "이미 모니터링이 실행 중입니다.", parent=self)
            return

        self.monitoring_active.set()
        self.monitoring_thread = threading.Thread(target=self._monitoring_loop, args=(interval_sec,), daemon=True)
        self.monitoring_thread.start()

        self.start_button.config(state="disabled")
        self.stop_button.config(state="normal")
        self.add_button.config(state="disabled")
        self.delete_button.config(state="disabled")
        self.add_current_button.config(state="disabled")
        self.app.log(f"[Price Monitor] 가격 변동 모니터링을 시작합니다. (주기: {interval_sec}초)")

    def stop_monitoring(self):
        self.monitoring_active.clear()
        self.start_button.config(state="normal")
        self.stop_button.config(state="disabled")
        self.add_button.config(state="normal")
        self.delete_button.config(state="normal")
        self.add_current_button.config(state="normal")
        self.app.log("[Price Monitor] 가격 변동 모니터링을 중지합니다.")
        for goods_no, item in self.monitoring_items.items():
            self.update_tree_item(goods_no, "중지됨", item.get("last_price"))

    def _monitoring_loop(self, interval_sec):
        s = crawler.build_session(self.app.musinsa_cookie_header)
        webhook_url = self.app.discord_webhook_url_entry.get().strip()

        while self.monitoring_active.is_set():
            self.app.log("[Price Monitor] 가격 상태 확인 시작...")
            
            current_items = list(self.monitoring_items.items())

            for goods_no, item_info in current_items:
                if not self.monitoring_active.is_set():
                    break
                
                try:
                    detail_meta = item_info['data']
                    self.app.log(f"[Price Monitor] 상품 확인: {detail_meta.get('goodsName')}")
                    
                    # process_one을 사용하여 정확한 최종 가격을 가져옵니다.
                    link_url = detail_meta.get("linkUrl", f"https://www.musinsa.com/app/goods/{goods_no}")
                    processed_data = crawler.process_one(
                        s,
                        self.app.musinsa_cookie_header,
                        goods_no,
                        link_url,
                        detail_meta,
                        sleep_each=0.2,
                        auth_payload=self.app.musinsa_auth_payload,
                    )
                    
                    current_price = processed_data.get("finalPrice")
                    original_price = item_info.get("original_price")
                    
                    status = "변동 없음"
                    tags = ()
                    notification_needed = False

                    if current_price is None:
                        status = "확인 불가"
                    elif original_price is not None and current_price < original_price:
                        status = "가격 변동 감지"
                        tags = ("PRICE_DOWN",)
                        notification_needed = True
                    # 'else'는 가격이 같거나 오른 경우로, 모두 "변동 없음" 처리
                    
                    if notification_needed and webhook_url:
                        message = self.get_price_notification_message(goods_no, item_info, original_price, current_price, status)
                        webhook.send_discord_notification(webhook_url, message)

                    self.root.after(0, self.update_tree_item, goods_no, status, current_price, tags)
                    
                    self.monitoring_items[goods_no]["last_price"] = current_price

                except Exception as e:
                    self.app.log(f"[Price Monitor] 상품 {goods_no} 확인 중 오류: {e}")
                    self.root.after(0, self.update_tree_item, goods_no, "오류", None)
                
                    self.app.log(f"[Price Monitor] 모든 상품 확인 완료. {interval_sec}초 후 다시 시작합니다.")
                    if self.monitoring_active.is_set():
                        # Pausable sleep for interval_sec seconds
                        for _ in range(interval_sec):
                            if not self.monitoring_active.is_set():
                                break
                            time.sleep(1)
        self.app.log("[Price Monitor] 모니터링 루프가 종료되었습니다.")

    def get_price_notification_message(self, goods_no, item_info, original_price, current_price, status):
        initial_price_from_dict = item_info.get("original_price")
        original_price_str = f"{utils.krw(initial_price_from_dict)}원" if initial_price_from_dict else "품절"
        current_price_str = f"{utils.krw(current_price)}원" if current_price else "품절"
        
        price_diff = 0
        if initial_price_from_dict and current_price:
            price_diff = current_price - initial_price_from_dict
        
        diff_str = ""
        if price_diff > 0:
            diff_str = f" (▲ {utils.krw(price_diff)}원)"
        elif price_diff < 0:
            diff_str = f" (▼ {utils.krw(abs(price_diff))}원)"

        return (
            f"**🚨 가격 변동 알림: {status} 🚨**\n\n"
            f"**상품 ID:** {goods_no}\n"
            f"**브랜드:** {item_info['data'].get('brandName')}\n"
            f"**상품명:** {item_info['data'].get('goodsName')}\n"
            f"**최초가:** {original_price_str}\n"
            f"**현재가:** {current_price_str}{diff_str}\n"
            f"[상품 바로가기](https://www.musinsa.com/app/goods/{goods_no})"
        )

    def update_tree_item(self, goods_no, status, current_price, tags=()):
        if goods_no in self.monitoring_items:
            iid = self.monitoring_items[goods_no]["iid"]
            if self.tree.exists(iid):
                values = list(self.tree.item(iid, "values"))
                
                # Display "확인 불가" if price is None, otherwise display the price.
                current_price_str = f"{utils.krw(current_price)}원" if current_price is not None else "확인 불가"
                
                values[4] = current_price_str # 현재가
                values[5] = status
                
                self.tree.item(iid, values=tuple(values), tags=tags)

    def on_close(self):
        if self.monitoring_thread and self.monitoring_thread.is_alive():
            self.stop_monitoring()
        self.destroy()

class PurchaseLimitMonitoringWindow(tk.Toplevel):
    def __init__(self, parent, app_instance):
        super().__init__(parent)
        self.app = app_instance
        self.root = parent
        self.title("구매 제한 해제 모니터링")
        self.geometry("800x600")

        self.monitoring_items = {}  # goodsNo -> {iid, data, last_limit_status}
        self.monitoring_thread = None
        self.monitoring_active = threading.Event()
        self.adding_items_counter = 0

        self.build_widgets()
        self.protocol("WM_DELETE_WINDOW", self.on_close)

    def build_widgets(self):
        control_frame = ttk.Frame(self, padding=10)
        control_frame.pack(fill="x", side="top")

        ttk.Label(control_frame, text="상품 ID:").pack(side="left", padx=(0, 5))
        self.goods_id_entry = ttk.Entry(control_frame, width=15)
        self.goods_id_entry.pack(side="left", padx=5)
        
        self.add_button = ttk.Button(control_frame, text="추가", command=self.add_item)
        self.add_button.pack(side="left", padx=5)

        self.delete_button = ttk.Button(control_frame, text="선택 삭제", command=self.delete_item)
        self.delete_button.pack(side="left", padx=5)

        ttk.Label(control_frame, text="주기(초):").pack(side="left", padx=(10, 5))
        self.interval_entry = ttk.Entry(control_frame, width=8)
        self.interval_entry.insert(0, "240")
        self.interval_entry.pack(side="left", padx=5)

        self.start_button = ttk.Button(control_frame, text="모니터링 시작", command=self.start_monitoring)
        self.start_button.pack(side="left", padx=5)

        self.stop_button = ttk.Button(control_frame, text="모니터링 중지", command=self.stop_monitoring, state="disabled")
        self.stop_button.pack(side="left", padx=5)

        tree_frame = ttk.Frame(self, padding=10)
        tree_frame.pack(fill="both", expand=True)

        cols = ("id", "brand", "name", "status")
        self.tree = ttk.Treeview(tree_frame, columns=cols, show="headings")
        
        headings = {"id": "상품 ID", "brand": "브랜드", "name": "상품명", "status": "구매 제한 상태"}
        widths = {"id": 80, "brand": 120, "name": 300, "status": 150}
        for col in cols:
            self.tree.heading(col, text=headings.get(col, col))
            self.tree.column(col, width=widths.get(col, 100), anchor="w")

        vsb = ttk.Scrollbar(tree_frame, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        self.tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")

        self.tree.tag_configure("UNLIMITED", background="lightgreen")
        self.tree.tag_configure("LIMITED", foreground="orange")

    def add_item(self):
        goods_no = self.goods_id_entry.get().strip()
        if not goods_no.isdigit():
            messagebox.showerror("입력 오류", "유효한 상품 ID(숫자)를 입력하세요.", parent=self)
            return
        
        if goods_no in self.monitoring_items:
            messagebox.showinfo("알림", "이미 모니터링 목록에 있는 상품입니다.", parent=self)
            return

        def _task():
            self.app.log(f"[LimitMon] 상품 정보 조회 중: {goods_no}")
            try:
                s = crawler.build_session(self.app.musinsa_cookie_header)
                limit_info = crawler.fetch_purchase_limit_info(s, goods_no)
                
                if not limit_info:
                    self.app.log(f"[LimitMon] 상품 정보를 찾을 수 없습니다: {goods_no}")
                    self.root.after(0, lambda: messagebox.showerror("조회 실패", f"상품 ID {goods_no} 정보를 찾을 수 없습니다.", parent=self))
                    return

                brand = limit_info.get("brandName", "N/A")
                name = limit_info.get("goodsName", "N/A")
                is_limited = limit_info.get("isLimitedQuantity")

                self.root.after(0, self._add_item_to_tree, goods_no, brand, name, is_limited, limit_info)

            except Exception as e:
                self.app.log(f"[LimitMon] 상품 정보 조회 실패: {e}")
                self.root.after(0, lambda: messagebox.showerror("조회 오류", f"상품 정보 조회 중 오류 발생: {e}", parent=self))

        threading.Thread(target=_task, daemon=True).start()

    def _add_item_to_tree(self, goods_no, brand, name, is_limited, data):
        status = "구매 제한" if is_limited else "제한 없음"
        iid = self.tree.insert("", "end", values=(goods_no, brand, name, status))
        self.monitoring_items[goods_no] = {"iid": iid, "data": data, "last_limit_status": is_limited}
        self.goods_id_entry.delete(0, "end")
        self.app.log(f"[LimitMon] 상품 추가됨: {name} ({goods_no}), 상태: {status}")

    def delete_item(self):
        selected = self.tree.selection()
        if not selected: return
        for iid in selected:
            goods_no = self.tree.item(iid, "values")[0]
            if goods_no in self.monitoring_items:
                del self.monitoring_items[goods_no]
                self.tree.delete(iid)
                self.app.log(f"[LimitMon] 상품 삭제됨: {goods_no}")

    def start_monitoring(self):
        if not self.monitoring_items:
            messagebox.showwarning("알림", "모니터링할 상품을 먼저 추가해주세요.", parent=self)
            return

        try:
            interval_sec = int(self.interval_entry.get().strip())
            if interval_sec <= 0:
                raise ValueError
        except ValueError:
            messagebox.showerror("입력 오류", "모니터링 주기에 유효한 양의 정수를 입력하세요.", parent=self)
            return

        if self.monitoring_thread and self.monitoring_thread.is_alive():
            messagebox.showinfo("알림", "이미 모니터링이 실행 중입니다.", parent=self)
            return

        self.monitoring_active.set()
        self.monitoring_thread = threading.Thread(target=self._purchase_limit_monitoring_loop, args=(interval_sec,), daemon=True)
        self.monitoring_thread.start()

        self.start_button.config(state="disabled")
        self.stop_button.config(state="normal")
        self.add_button.config(state="disabled")
        self.delete_button.config(state="disabled")
        self.app.log(f"[LimitMon] 구매 제한 모니터링을 시작합니다. (주기: {interval_sec}초)")

    def stop_monitoring(self):
        self.monitoring_active.clear()
        self.start_button.config(state="normal")
        self.stop_button.config(state="disabled")
        self.add_button.config(state="normal")
        self.delete_button.config(state="normal")
        self.app.log("[LimitMon] 구매 제한 모니터링을 중지합니다.")
        for goods_no, item in self.monitoring_items.items():
            self.update_tree_item(goods_no, "중지됨", item['last_limit_status'])

    def add_current_products(self):
        if not self.app.items_in_tree:
            messagebox.showwarning("?뚮┝", "硫붿씤 紐⑸줉???곹뭹???놁뒿?덈떎.", parent=self)
            return

        self.app.log("[LimitMon] ?꾩옱 ?щ·留곷맂 紐⑤뱺 ?곹뭹??紐⑤땲?곕쭅 紐⑸줉??異붽??⑸땲??")
        items_to_add = list(self.app.items_in_tree.values())
        self.add_current_button.config(state="disabled")
        self.adding_items_counter = len(items_to_add)

        for item_data in items_to_add:
            goods_no = item_data.get("goodsNo")
            if not goods_no or goods_no in self.monitoring_items:
                self.item_add_completed()
                continue
            self.add_item_from_data(item_data)

    def add_item_from_data(self, item_data):
        goods_no = item_data.get("goodsNo")

        def _task():
            try:
                s = crawler.build_session(self.app.musinsa_cookie_header)
                limit_info = crawler.fetch_purchase_limit_info(s, goods_no)
                if limit_info:
                    brand = limit_info.get("brandName", item_data.get("brandName", "N/A"))
                    name = limit_info.get("goodsName", item_data.get("goodsName", "N/A"))
                    is_limited = limit_info.get("isLimitedQuantity")
                    self.root.after(0, self._add_item_to_tree, goods_no, brand, name, is_limited, limit_info)
                else:
                    self.app.log(f"[LimitMon] ?곹뭹 ?뺣낫瑜?李얠쓣 ???놁뼱 嫄대꼫?곷땲?? {goods_no}")
            except Exception as e:
                self.app.log(f"[LimitMon] ?곹뭹 ?붽? 議고쉶 ?ㅽ뙣: {goods_no} / {e}")
            finally:
                self.item_add_completed()

        threading.Thread(target=_task, daemon=True).start()

    def item_add_completed(self):
        self.adding_items_counter -= 1
        if self.adding_items_counter <= 0:
            self.add_current_button.config(state="normal")
            self.app.log("[LimitMon] 紐⑤뱺 ?곹뭹 異붽? ?꾨즺.")

    def _purchase_limit_monitoring_loop(self, interval_sec):
        s = crawler.build_session(self.app.musinsa_cookie_header)
        webhook_url = self.app.discord_webhook_url_entry.get().strip()

        while self.monitoring_active.is_set():
            self.app.log("[LimitMon] 구매 제한 상태 확인 시작...")
            current_items = list(self.monitoring_items.items())

            for goods_no, item_info in current_items:
                if not self.monitoring_active.is_set(): break

                try:
                    self.app.log(f"[LimitMon] 상품 확인: {item_info['data'].get('goodsName')}")
                    limit_info = crawler.fetch_purchase_limit_info(s, goods_no)
                    if not limit_info:
                        self.app.log(f"[LimitMon] 상품 정보를 찾을 수 없어 건너뜁니다: {goods_no}")
                        continue

                    is_limited = limit_info.get("isLimitedQuantity")
                    last_limit_status = item_info.get("last_limit_status")

                    status = "구매 제한" if is_limited else "제한 없음"
                    tags = ("LIMITED",) if is_limited else ()
                    notification_needed = False

                    if last_limit_status and not is_limited:
                        status = "✅ 구매제한 해제!"
                        tags = ("UNLIMITED",)
                        self.app.log(f"[LimitMon] 구매 제한 해제 감지! 상품: {item_info['data'].get('goodsName')}")
                        notification_needed = True

                    if notification_needed and webhook_url:
                        message = (
                            f"**🔓 구매제한 해제 알림 🔓**\n\n"
                            f"**상품 ID:** {goods_no}\n"
                            f"**브랜드:** {item_info['data'].get('brandName')}\n"
                            f"**상품명:** {item_info['data'].get('goodsName')}\n"
                            f"[상품 바로가기](https://www.musinsa.com/app/goods/{goods_no})"
                        )
                        webhook.send_discord_notification(webhook_url, message)

                    self.root.after(0, self.update_tree_item, goods_no, status, is_limited, tags)
                    self.monitoring_items[goods_no]["last_limit_status"] = is_limited

                except Exception as e:
                    self.app.log(f"[LimitMon] 상품 {goods_no} 확인 중 오류: {e}")
                    self.root.after(0, self.update_tree_item, goods_no, "오류", item_info.get('last_limit_status'), ())

                if self.monitoring_active.is_set():
                    time.sleep(random.uniform(5, 10))

            self.app.log(f"[LimitMon] 모든 상품 확인 완료. {interval_sec}초 후 다시 시작합니다.")
            self.monitoring_active.wait(interval_sec)

        self.app.log("[LimitMon] 구매 제한 모니터링 루프가 종료되었습니다.")

    def update_tree_item(self, goods_no, status, is_limited, tags=()):
        if goods_no in self.monitoring_items:
            iid = self.monitoring_items[goods_no]["iid"]
            if self.tree.exists(iid):
                self.tree.item(iid, values=(
                    goods_no,
                    self.tree.item(iid, "values")[1],
                    self.tree.item(iid, "values")[2],
                    status
                ), tags=tags)

    def on_close(self):
        if self.monitoring_thread and self.monitoring_thread.is_alive():
            self.stop_monitoring()
        self.destroy()

def _patched_monitoring_build_widgets(self):
    control_frame = ttk.Frame(self, padding=10)
    control_frame.pack(fill="x", side="top")

    ttk.Label(control_frame, text="상품 ID:").pack(side="left", padx=(0, 5))
    self.goods_id_entry = ttk.Entry(control_frame, width=15)
    self.goods_id_entry.pack(side="left", padx=5)

    self.add_button = ttk.Button(control_frame, text="추가", command=self.add_item)
    self.add_button.pack(side="left", padx=5)

    self.delete_button = ttk.Button(control_frame, text="선택 삭제", command=self.delete_item)
    self.delete_button.pack(side="left", padx=5)

    self.add_current_button = ttk.Button(control_frame, text="현재 크롤링 상품 모니터링", command=self.add_current_products)
    self.add_current_button.pack(side="left", padx=5)

    ttk.Label(control_frame, text="주기(초):").pack(side="left", padx=(10, 5))
    self.interval_entry = ttk.Entry(control_frame, width=8)
    self.interval_entry.insert(0, "180")
    self.interval_entry.pack(side="left", padx=5)

    self.start_button = ttk.Button(control_frame, text="모니터링 시작", command=self.start_monitoring)
    self.start_button.pack(side="left", padx=5)

    self.stop_button = ttk.Button(control_frame, text="모니터링 중지", command=self.stop_monitoring, state="disabled")
    self.stop_button.pack(side="left", padx=5)

    tree_frame = ttk.Frame(self, padding=10)
    tree_frame.pack(fill="both", expand=True)

    cols = ("id", "brand", "name", "status", "sizes")
    self.tree = ttk.Treeview(tree_frame, columns=cols, show="headings")
    headings = {"id": "상품 ID", "brand": "브랜드", "name": "상품명", "status": "상태", "sizes": "재고 사이즈"}
    widths = {"id": 80, "brand": 100, "name": 250, "status": 100, "sizes": 150}
    for col in cols:
        self.tree.heading(col, text=headings.get(col, col))
        self.tree.column(col, width=widths.get(col, 100), anchor="w")

    vsb = ttk.Scrollbar(tree_frame, orient="vertical", command=self.tree.yview)
    self.tree.configure(yscrollcommand=vsb.set)
    self.tree.pack(side="left", fill="both", expand=True)
    vsb.pack(side="right", fill="y")

    self.tree.tag_configure("RESTOCKED", background="lightgreen")
    self.tree.tag_configure("SOLD_OUT", foreground="gray")


def _patched_monitoring_add_current_products(self):
    if not self.app.items_in_tree:
        messagebox.showwarning("알림", "메인 목록에 상품이 없습니다.", parent=self)
        return

    self.app.log("[Monitor] 현재 크롤링 상품을 재입고 모니터링 목록에 추가합니다.")
    items_to_add = list(self.app.items_in_tree.values())
    self.add_current_button.config(state="disabled")
    self.adding_items_counter = len(items_to_add)

    for item_data in items_to_add:
        goods_no = item_data.get("goodsNo")
        if not goods_no or goods_no in self.monitoring_items:
            _patched_monitoring_item_add_completed(self)
            continue
        _patched_monitoring_add_item_from_data(self, item_data)


def _patched_monitoring_add_item_from_data(self, item_data):
    goods_no = item_data.get("goodsNo")
    brand = item_data.get("brandName", "N/A")
    name = item_data.get("goodsName", "N/A")
    sizes = set(s.strip() for s in str(item_data.get("availableSizes") or "").split(",") if s.strip())
    self.root.after(0, self._add_item_to_tree, goods_no, brand, name, item_data, sizes)
    _patched_monitoring_item_add_completed(self)


def _patched_monitoring_item_add_completed(self):
    self.adding_items_counter -= 1
    if self.adding_items_counter <= 0:
        self.add_current_button.config(state="normal")
        self.app.log("[Monitor] 모든 상품 추가 완료.")


def _patched_purchase_limit_build_widgets(self):
    control_frame = ttk.Frame(self, padding=10)
    control_frame.pack(fill="x", side="top")

    ttk.Label(control_frame, text="상품 ID:").pack(side="left", padx=(0, 5))
    self.goods_id_entry = ttk.Entry(control_frame, width=15)
    self.goods_id_entry.pack(side="left", padx=5)

    self.add_button = ttk.Button(control_frame, text="추가", command=self.add_item)
    self.add_button.pack(side="left", padx=5)

    self.delete_button = ttk.Button(control_frame, text="선택 삭제", command=self.delete_item)
    self.delete_button.pack(side="left", padx=5)

    self.add_current_button = ttk.Button(control_frame, text="현재 크롤링 상품 모니터링", command=self.add_current_products)
    self.add_current_button.pack(side="left", padx=5)

    ttk.Label(control_frame, text="주기(초):").pack(side="left", padx=(10, 5))
    self.interval_entry = ttk.Entry(control_frame, width=8)
    self.interval_entry.insert(0, "240")
    self.interval_entry.pack(side="left", padx=5)

    self.start_button = ttk.Button(control_frame, text="모니터링 시작", command=self.start_monitoring)
    self.start_button.pack(side="left", padx=5)

    self.stop_button = ttk.Button(control_frame, text="모니터링 중지", command=self.stop_monitoring, state="disabled")
    self.stop_button.pack(side="left", padx=5)

    tree_frame = ttk.Frame(self, padding=10)
    tree_frame.pack(fill="both", expand=True)

    cols = ("id", "brand", "name", "status")
    self.tree = ttk.Treeview(tree_frame, columns=cols, show="headings")
    headings = {"id": "상품 ID", "brand": "브랜드", "name": "상품명", "status": "구매 제한 상태"}
    widths = {"id": 80, "brand": 120, "name": 300, "status": 150}
    for col in cols:
        self.tree.heading(col, text=headings.get(col, col))
        self.tree.column(col, width=widths.get(col, 100), anchor="w")

    vsb = ttk.Scrollbar(tree_frame, orient="vertical", command=self.tree.yview)
    self.tree.configure(yscrollcommand=vsb.set)
    self.tree.pack(side="left", fill="both", expand=True)
    vsb.pack(side="right", fill="y")

    self.tree.tag_configure("UNLIMITED", background="lightgreen")
    self.tree.tag_configure("LIMITED", foreground="orange")


def _patched_purchase_limit_add_current_products(self):
    if not self.app.items_in_tree:
        messagebox.showwarning("알림", "메인 목록에 상품이 없습니다.", parent=self)
        return

    self.app.log("[LimitMon] 현재 크롤링 상품을 구매제한 모니터링 목록에 추가합니다.")
    items_to_add = list(self.app.items_in_tree.values())
    self.add_current_button.config(state="disabled")
    self.adding_items_counter = len(items_to_add)

    for item_data in items_to_add:
        goods_no = item_data.get("goodsNo")
        if not goods_no or goods_no in self.monitoring_items:
            _patched_purchase_limit_item_add_completed(self)
            continue
        _patched_purchase_limit_add_item_from_data(self, item_data)


def _patched_purchase_limit_add_item_from_data(self, item_data):
    goods_no = item_data.get("goodsNo")

    def _task():
        try:
            s = crawler.build_session(self.app.musinsa_cookie_header)
            limit_info = crawler.fetch_purchase_limit_info(s, goods_no)
            if limit_info:
                brand = limit_info.get("brandName", item_data.get("brandName", "N/A"))
                name = limit_info.get("goodsName", item_data.get("goodsName", "N/A"))
                is_limited = limit_info.get("isLimitedQuantity")
                self.root.after(0, self._add_item_to_tree, goods_no, brand, name, is_limited, limit_info)
            else:
                self.app.log(f"[LimitMon] 상품 정보를 찾을 수 없어 건너뜁니다: {goods_no}")
        except Exception as e:
            self.app.log(f"[LimitMon] 상품 추가 조회 실패: {goods_no} / {e}")
        finally:
            _patched_purchase_limit_item_add_completed(self)

    threading.Thread(target=_task, daemon=True).start()


def _patched_purchase_limit_item_add_completed(self):
    self.adding_items_counter -= 1
    if self.adding_items_counter <= 0:
        self.add_current_button.config(state="normal")
        self.app.log("[LimitMon] 모든 상품 추가 완료.")


def _monitoring_sort_sizes(values):
    values = [str(v) for v in values if str(v).strip()]
    try:
        return sorted(values, key=lambda x: (not x.isdigit(), int(x) if x.isdigit() else x))
    except Exception:
        return sorted(values)


def _monitoring_stock_sets(stock_info):
    if isinstance(stock_info, dict):
        available = set(str(v) for v in (stock_info.get("available") or []) if str(v).strip())
        unavailable = set(str(v) for v in (stock_info.get("unavailable") or []) if str(v).strip())
        return available, unavailable
    values = set(str(v) for v in (stock_info or []) if str(v).strip())
    return values, set()


def _monitoring_stock_text(stock_info):
    available, unavailable = _monitoring_stock_sets(stock_info)
    parts = []
    if available:
        parts.append(f"재고있음: {', '.join(_monitoring_sort_sizes(available))}")
    if unavailable:
        parts.append(f"품절: {', '.join(_monitoring_sort_sizes(unavailable))}")
    return " / ".join(parts) if parts else "정보 없음"


def _patched_monitoring_add_item_from_data(self, item_data):
    goods_no = item_data.get("goodsNo")
    brand = item_data.get("brandName", "N/A")
    name = item_data.get("goodsName", "N/A")

    def _task():
        try:
            stock_info = {"available": [], "unavailable": []}
            if self.app.musinsa_cookie_header:
                s = crawler.build_session(self.app.musinsa_cookie_header)
                opts = crawler.fetch_options_any(s, goods_no)
                stock_info = crawler.extract_stock_size_sets(opts)
            else:
                stock_info["available"] = [
                    s.strip() for s in str(item_data.get("availableSizes") or "").split(",") if s.strip()
                ]
            self.root.after(0, self._add_item_to_tree, goods_no, brand, name, item_data, stock_info)
        except Exception as e:
            self.app.log(f"[Monitor] 상품 추가 조회 실패: {goods_no} / {e}")
            fallback_sizes = [s.strip() for s in str(item_data.get("availableSizes") or "").split(",") if s.strip()]
            self.root.after(
                0,
                self._add_item_to_tree,
                goods_no,
                brand,
                name,
                item_data,
                {"available": fallback_sizes, "unavailable": []},
            )
        finally:
            _patched_monitoring_item_add_completed(self)

    threading.Thread(target=_task, daemon=True).start()


def _patched_monitoring_add_item_to_tree(self, goods_no, brand, name, data, initial_stock):
    available_sizes, unavailable_sizes = _monitoring_stock_sets(initial_stock)
    sizes_str = _monitoring_stock_text(initial_stock)
    status = "모니터링 시작" if available_sizes else "품절 상태로 시작"

    iid = self.tree.insert("", "end", values=(goods_no, brand, name, status, sizes_str))
    self.monitoring_items[goods_no] = {
        "iid": iid,
        "data": data,
        "last_sizes": available_sizes,
        "last_unavailable_sizes": unavailable_sizes,
    }

    self.goods_id_entry.delete(0, "end")
    self.app.log(f"[Monitor] 상품 추가: {name} ({goods_no})")


def _patched_monitoring_loop(self, interval_sec):
    while self.monitoring_active.is_set():
        webhook_url = self.app.discord_webhook_url_entry.get().strip()

        for goods_no, item_info in list(self.monitoring_items.items()):
            if not self.monitoring_active.is_set():
                break

            try:
                s = crawler.build_session(self.app.musinsa_cookie_header)
                self.app.log(f"[Monitor] 상품 확인: {item_info['data'].get('goodsName')}")
                opts = crawler.fetch_options_any(s, goods_no)
                stock_info = crawler.extract_stock_size_sets(opts)
                available_sizes, unavailable_sizes = _monitoring_stock_sets(stock_info)
                last_sizes = set(item_info.get("last_sizes", set()))

                if available_sizes and unavailable_sizes:
                    status = "일부 재고 있음"
                elif available_sizes:
                    status = "재고 있음"
                elif unavailable_sizes:
                    status = "품절"
                else:
                    status = "정보 없음"

                tags = ()
                notification_needed = False

                if not last_sizes and available_sizes:
                    status = "재입고"
                    tags = ("RESTOCKED",)
                    self.app.log(
                        f"[Monitor] 재입고 감지! 상품: {item_info['data'].get('goodsName')}, 사이즈 {', '.join(_monitoring_sort_sizes(available_sizes))}"
                    )
                    notification_needed = True
                elif last_sizes and not available_sizes:
                    status = "품절"
                    tags = ("SOLD_OUT",)
                elif available_sizes:
                    newly_stocked = available_sizes - last_sizes
                    if newly_stocked:
                        status = f"새 입고: {', '.join(_monitoring_sort_sizes(newly_stocked))}"
                        tags = ("RESTOCKED",)
                        self.app.log(
                            f"[Monitor] 부분 재입고 감지! 상품: {item_info['data'].get('goodsName')}, 신규 사이즈 {', '.join(_monitoring_sort_sizes(newly_stocked))}"
                        )
                        notification_needed = True

                if notification_needed and webhook_url:
                    message = (
                        f"**무신사 재입고 알림**\n\n"
                        f"**상품 ID:** {goods_no}\n"
                        f"**브랜드:** {item_info['data'].get('brandName')}\n"
                        f"**상품명:** {item_info['data'].get('goodsName')}\n"
                        f"**재고 상태:** {_monitoring_stock_text(stock_info)}\n"
                        f"[상품 바로가기](https://www.musinsa.com/app/goods/{goods_no})"
                    )
                    webhook.send_discord_notification(webhook_url, message)

                self.root.after(0, self.update_tree_item, goods_no, status, stock_info, tags)
                self.monitoring_items[goods_no]["last_sizes"] = available_sizes
                self.monitoring_items[goods_no]["last_unavailable_sizes"] = unavailable_sizes

            except Exception as e:
                self.app.log(f"[Monitor] 상품 {goods_no} 확인 중 오류: {e}")
                self.root.after(0, self.update_tree_item, goods_no, "오류", {"available": [], "unavailable": []})

        for _ in range(interval_sec):
            if not self.monitoring_active.is_set():
                break
            time.sleep(1)
    self.app.log("[Monitor] 모니터링 루프가 종료되었습니다.")


def _patched_monitoring_update_tree_item(self, goods_no, status, stock_info, tags=()):
    if goods_no in self.monitoring_items:
        iid = self.monitoring_items[goods_no]["iid"]
        if self.tree.exists(iid):
            self.tree.item(iid, values=(
                goods_no,
                self.tree.item(iid, "values")[1],
                self.tree.item(iid, "values")[2],
                status,
                _monitoring_stock_text(stock_info),
            ), tags=tags)

def _patched_monitoring_resolve_stock_info_v2(self, item_data):
    goods_no = item_data.get("goodsNo")
    stock_info = {"available": [], "unavailable": []}
    if self.app.musinsa_cookie_header:
        s = crawler.build_session(self.app.musinsa_cookie_header)
        opts = crawler.fetch_options_any(s, goods_no)
        stock_info = crawler.extract_stock_size_sets(opts)
    else:
        stock_info["available"] = [
            s.strip() for s in str(item_data.get("availableSizes") or "").split(",") if s.strip()
        ]
    return stock_info


def _patched_monitoring_add_current_products_v2(self):
    if not self.app.items_in_tree:
        messagebox.showwarning("?뚮┝", "硫붿씤 紐⑸줉???곹뭹???놁뒿?덈떎.", parent=self)
        return

    self.app.log("[Monitor] ?꾩옱 ?щ·留??곹뭹???ъ엯怨?紐⑤땲?곕쭅 紐⑸줉??異붽??⑸땲??")
    items_to_add = list(self.app.items_in_tree.values())
    self.add_current_button.config(state="disabled")
    self.adding_items_counter = len(items_to_add)

    def _task():
        for item_data in items_to_add:
            goods_no = item_data.get("goodsNo")
            if not goods_no or goods_no in self.monitoring_items:
                self.root.after(0, _patched_monitoring_item_add_completed, self)
                continue

            brand = item_data.get("brandName", "N/A")
            name = item_data.get("goodsName", "N/A")
            try:
                stock_info = _patched_monitoring_resolve_stock_info_v2(self, item_data)
            except Exception as e:
                self.app.log(f"[Monitor] ?곹뭹 異붽? 議고쉶 ?ㅽ뙣: {goods_no} / {e}")
                fallback_sizes = [s.strip() for s in str(item_data.get("availableSizes") or "").split(",") if s.strip()]
                stock_info = {"available": fallback_sizes, "unavailable": []}

            self.root.after(0, self._add_item_to_tree, goods_no, brand, name, item_data, stock_info)
            self.root.after(0, _patched_monitoring_item_add_completed, self)

    threading.Thread(target=_task, daemon=True).start()


def _patched_monitoring_add_item_from_data_v2(self, item_data):
    goods_no = item_data.get("goodsNo")
    brand = item_data.get("brandName", "N/A")
    name = item_data.get("goodsName", "N/A")

    def _task():
        try:
            stock_info = _patched_monitoring_resolve_stock_info_v2(self, item_data)
            self.root.after(0, self._add_item_to_tree, goods_no, brand, name, item_data, stock_info)
        except Exception as e:
            self.app.log(f"[Monitor] ?곹뭹 異붽? 議고쉶 ?ㅽ뙣: {goods_no} / {e}")
            fallback_sizes = [s.strip() for s in str(item_data.get("availableSizes") or "").split(",") if s.strip()]
            self.root.after(
                0,
                self._add_item_to_tree,
                goods_no,
                brand,
                name,
                item_data,
                {"available": fallback_sizes, "unavailable": []},
            )
        finally:
            _patched_monitoring_item_add_completed(self)

    threading.Thread(target=_task, daemon=True).start()


MonitoringWindow.build_widgets = _patched_monitoring_build_widgets
MonitoringWindow.add_current_products = _patched_monitoring_add_current_products_v2
MonitoringWindow.add_item_from_data = _patched_monitoring_add_item_from_data_v2
MonitoringWindow.item_add_completed = _patched_monitoring_item_add_completed
MonitoringWindow._add_item_to_tree = _patched_monitoring_add_item_to_tree
MonitoringWindow._monitoring_loop = _patched_monitoring_loop
MonitoringWindow.update_tree_item = _patched_monitoring_update_tree_item

PurchaseLimitMonitoringWindow.build_widgets = _patched_purchase_limit_build_widgets
PurchaseLimitMonitoringWindow.add_current_products = _patched_purchase_limit_add_current_products
PurchaseLimitMonitoringWindow.add_item_from_data = _patched_purchase_limit_add_item_from_data
PurchaseLimitMonitoringWindow.item_add_completed = _patched_purchase_limit_item_add_completed


class DetailWindow(tk.Toplevel):
    def __init__(self, parent, data: dict, app_instance: App, kream_sales_data: dict | None = None):
        super().__init__(parent)
        self.data = data
        self.app = app_instance
        self._img_cache = None
        self.kream_sales_data = kream_sales_data

        title = f"[{data.get('rank')}위] {data.get('goodsName') or ''}"
        self.title(title)
        self.geometry("900x620")

        self.build_widgets()

    def build_widgets(self):
        left = ttk.Frame(self, padding=10, relief="solid", borderwidth=1)
        left.grid(row=0, column=0, sticky="nsew", padx=(10,5), pady=10)
        right = ttk.Frame(self, padding=10, relief="solid", borderwidth=1)
        right.grid(row=0, column=1, sticky="nsew", padx=(5,10), pady=10)

        self.grid_rowconfigure(0, weight=1)
        self.grid_columnconfigure(0, weight=0)
        self.grid_columnconfigure(1, weight=1)

        img_label = ttk.Label(left)
        img_label.pack(fill="both", expand=True)

        def load_image():
            try:
                url = self.data.get("imageUrl")
                if not url: 
                    img_label.configure(text="이미지 없음"); return
                r = requests.get(url, timeout=10)
                r.raise_for_status()
                im = Image.open(io.BytesIO(r.content))
                im.thumbnail((360, 360))
                self._img_cache = ImageTk.PhotoImage(im)
                img_label.configure(image=self._img_cache)
            except Exception as e:
                img_label.configure(text=f"이미지 로드 실패: {e}")
        threading.Thread(target=load_image, daemon=True).start()

        grid = ttk.Frame(right)
        grid.pack(fill="x", anchor="n")

        row_idx = 0
        def add_row(k, v):
            nonlocal row_idx
            kl = ttk.Label(grid, text=k, width=16, anchor="e")
            vl = ttk.Label(grid, text=str(v), width=60, anchor="w", wraplength=400)
            kl.grid(row=row_idx, column=0, sticky="e", padx=4, pady=2)
            vl.grid(row=row_idx, column=1, sticky="w", padx=4, pady=2)
            sep = ttk.Separator(grid, orient="horizontal")
            sep.grid(row=row_idx+1, column=0, columnspan=2, sticky="ew", pady=(0,4))
            row_idx += 2

        add_row("브랜드", self.data.get("brandName", ""))
        add_row("상품번호", self.data.get("goodsNo", ""))
        add_row("품번", self.data.get("serial_number", ""))
        add_row("링크", self.data.get("linkUrl", ""))
        add_row("카테고리", self.data.get("baseCategoryFullName", ""))
        add_row("판매가(최초)", f"{utils.krw(self.data.get('salePrice') or 0)}원" if self.data.get("salePrice") is not None else "")
        add_row("쿠폰할인", f"{utils.krw(self.data.get('couponDiscount') or 0)}원")
        add_row("최종가", f"{utils.krw(self.data.get('finalPrice') or 0)}원" if self.data.get("finalPrice") is not None else "")
        add_row("멤버할인%", str(self.data.get("memberDiscountRate", "")))
        add_row("적립률%", str(self.data.get("memberSavePointRate", "")))
        add_row("멤버 할인 가능", "Y" if self.data.get("isMemberDiscount") else "N")
        add_row("적립 선할인", "Y" if self.data.get("isUsePointDiscount") else "N")
        add_row("goodsSavePoint", f"{utils.krw(self.data.get('goodsSavePoint') or 0)}원")
        add_row("가능한 사이즈", self.data.get("availableSizes", ""))
        add_row("선택 옵션코드", self.data.get("managedCode", ""))

        if self.kream_sales_data:
            total_cnt = int(self.kream_sales_data.get("total") or 0)
            monthly_cnt = 0
            if 'monthly' in self.kream_sales_data:
                monthly_cnt = self.kream_sales_data['monthly']
            else:
                one_month_ago = datetime.now(timezone.utc) - timedelta(days=30)
                for sale_item in self.kream_sales_data.get("items", []):
                    try:
                        sale_date_str = sale_item.get("date_created")
                        sale_date = datetime.strptime(sale_date_str, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=timezone.utc)
                        if sale_date >= one_month_ago:
                            monthly_cnt += 1
                    except (ValueError, TypeError):
                        continue
            add_row("크림 전체 거래건수", f"{total_cnt}건")
            add_row("크림 한달간 거래수", f"{monthly_cnt}건")
        elif 'kream_total_transactions' in self.data and 'kream_monthly_transactions' in self.data:
            add_row("크림 전체 거래건수", f"{self.data['kream_total_transactions']}건")
            add_row("크림 한달간 거래수", f"{self.data['kream_monthly_transactions']}건")

if __name__ == '__main__':
    root = tk.Tk()
    app = App(root)
    root.mainloop()

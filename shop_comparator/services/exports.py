"""Downloads: comparison XLSX/CSV (web), GUI-format CSV and the KREAM bid sheet."""
import csv
import io
import math
import os
from datetime import datetime
from decimal import Decimal

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from shop_comparator.services import comparison_service as cmp
from shop_comparator.storage import store
from shop_comparator.core.crawl_errors import CrawlError
from shop_comparator.core.utils import resource_path

COLUMNS = {
    "rank": "랭킹", "goodsNo": "무신사 번호", "goodsName": "상품명", "brandName": "브랜드", "serial": "품번",
    "size": "무신사 사이즈(원본)", "sizeKey": "정규화 사이즈", "optionNo": "옵션 번호",
    "salePrice": "무신사 판매가", "coupon": "최대 할인 쿠폰", "couponDiscount": "쿠폰 할인액",
    "memberDiscount": "회원 등급 할인액", "pointDiscount": "적립금 선할인액", "musinsaPrice": "무신사 회원가",
    "priceSource": "회원가 출처", "musinsaCheckedAt": "무신사 조회 시각(UTC)", "musinsaError": "무신사 오류",
    "kreamId": "KREAM 번호", "kreamSize": "KREAM 사이즈(원본)", "kreamBuy": "KREAM 즉시구매 최저가",
    "kreamBuyDelivery": "즉시구매 배송 유형", "kreamSell": "KREAM 즉시판매가(최고 구매입찰)",
    "kreamCheckedAt": "KREAM 조회 시각(UTC)", "kreamError": "KREAM 오류",
    "difference": "가격 차이 (KREAM 즉시구매가-무신사)", "sellDifference": "판매 차익 (KREAM 즉시판매가-무신사)",
    "poizonPrice": "POIZON 가격", "poizonDifference": "POIZON 차이 (POIZON-무신사)",
    "matchStatus": "상품 매칭", "matchConfidence": "매칭 신뢰도", "matchMethod": "매칭 방식", "status": "상태",
    "musinsaUrl": "무신사 URL", "kreamUrl": "KREAM URL",
}
MONEY = {"salePrice", "couponDiscount", "memberDiscount", "pointDiscount", "musinsaPrice", "kreamBuy", "kreamSell",
         "difference", "sellDifference", "poizonPrice", "poizonDifference"}


def comparison_workbook(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "사이즈별 비교"
    keys = list(COLUMNS)
    ws.append([COLUMNS[k] for k in keys])
    for row in rows:
        ws.append([row.get(key) for key in keys])
        for cell in ws[ws.max_row]:
            # Retailer-provided text must remain text, including a leading '='.
            if isinstance(cell.value, str):
                cell.data_type = "s"
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="163A35")
    for index, key in enumerate(keys, 1):
        letter = openpyxl.utils.get_column_letter(index)
        if key in MONEY:
            for cell in ws[letter][1:]:
                cell.number_format = "#,##0;[Red]-#,##0"
        ws.column_dimensions[letter].width = min(46, max(14, max(len(str(c.value or "")) for c in ws[letter]) + 3))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    info = wb.create_sheet("기준")
    for line in [
        ["생성 시각(UTC)", store.now()],
        ["가격 기준", cmp.PRICE_NOTE],
        ["가격 차이", "KREAM 즉시구매 최저가 − 무신사 회원가 (양수면 KREAM이 더 비쌈)"],
        ["KREAM 즉시구매 최저가", "새 상품 기준 즉시 구매 가능한 가격 중 배송 유형을 통틀어 가장 낮은 값. 95점·보관판매 제외. 구매 입찰만 가능한 사이즈는 빈칸."],
        ["사이즈 매칭", "같은 사이즈 체계끼리만 매칭(mm↔mm, US↔US 등). 체계가 다르면 환산하지 않고 빈칸."],
        ["빈칸", "조회하지 못했거나 값이 없는 경우. 추정값으로 채우지 않음."],
    ]:
        info.append(line)
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 110
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def comparison_csv(rows):
    text = io.StringIO()
    writer = csv.writer(text)
    keys = list(COLUMNS)
    writer.writerow([COLUMNS[k] for k in keys])
    for row in rows:
        writer.writerow([_csv_safe(row.get(k)) for k in keys])
    return io.BytesIO(text.getvalue().encode("utf-8-sig"))


def _csv_safe(value):
    """Neutralise spreadsheet formulas in shop-provided text."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return "" if value is None else value


def gui_comparison_csv(rows, source="kream"):
    """Same columns and sign as gui._save_comparison_csv_task (차액 = 무신사 − 상대)."""
    products = {str(p["goodsNo"]): p for p in store.products(active_only=False)}
    states = products
    if source == "poizon":
        fields = ["rank", "goodsName", "brandName", "size", "musinsa_price", "poizon_price", "price_diff", "musinsa_link", "poizon_link"]
    else:
        fields = ["rank", "goodsName", "brandName", "size", "musinsa_price", "kream_buy_price", "kream_sell_price", "price_diff",
                  "kream_total_transactions", "kream_monthly_transactions", "musinsa_link", "kream_link"]
    text = io.StringIO()
    writer = csv.DictWriter(text, fieldnames=fields, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        product = products.get(row["goodsNo"], {})
        other = row["poizonPrice"] if source == "poizon" else row["kreamBuy"]
        size = row["size"] or row["kreamSize"]
        if row["size"] and row["kreamSize"] and row["size"] != row["kreamSize"]:
            size = f"{row['kreamSize']} ({row['size']})"
        writer.writerow({
            "rank": row["rank"], "goodsName": _csv_safe(row["goodsName"]), "brandName": _csv_safe(row["brandName"]), "size": size,
            "musinsa_price": row["musinsaPrice"], "kream_buy_price": row["kreamBuy"], "kream_sell_price": row["kreamSell"],
            "poizon_price": row["poizonPrice"],
            "price_diff": row["musinsaPrice"] - other if row["musinsaPrice"] is not None and other is not None else None,
            "kream_total_transactions": product.get("kream_total_transactions"),
            "kream_monthly_transactions": product.get("kream_monthly_transactions"),
            "musinsa_link": product.get("linkUrl") or f"https://www.musinsa.com/products/{row['goodsNo']}",
            "kream_link": f"https://kream.co.kr/products/{row['kreamId']}" if row["kreamId"] else "",
            "poizon_link": states.get(row["goodsNo"], {}).get("poizonLink") or "",
        })
    return io.BytesIO(text.getvalue().encode("utf-8-sig"))


def bid_rows(margin_percent, quantity, deadline, goods_nos=None):
    """Rows of the KREAM bid sheet (GUI _generate_excel_task).

    판매 희망가 = 무신사 회원가 × (1 + 마진율), 1,000원 단위 올림.  Unlike the
    GUI this uses each size's own member price and only sizes that match a
    KREAM option of the same size system.
    """
    out, seen = [], set()
    for row in cmp.build(goods_nos, include_kream_only=False):
        if not row["kreamId"] or not row["kreamSize"] or row["musinsaPrice"] is None or row["soldout"]:
            continue
        if row["matchStatus"] not in ("직접 지정", "자동 검증", "이미지 검토", "모델번호 일치") and not cmp.trust_unverified():
            continue
        key = (row["kreamId"], row["kreamSize"])
        if key in seen:
            continue
        seen.add(key)
        # Decimal: 90,000 x 1.1 must be 99,000, not 99,000.00000000001 rounded up to 100,000.
        exact = Decimal(row["musinsaPrice"]) * (Decimal(100) + Decimal(str(margin_percent))) / Decimal(100)
        price = int(math.ceil(exact / 1000) * 1000)
        out.append({"상품번호": row["kreamId"], "모델번호": row["serial"], "영문 상품명": row["goodsName"],
                    "옵션명": row["kreamSize"], "판매 희망가": price, "수량": str(quantity), "입찰기한": deadline})
    return out


def find_template():
    for base in (resource_path("."), os.path.abspath(".")):
        try:
            for name in os.listdir(base):
                if name.startswith("k-c46vwlba") and name.endswith(".xlsx"):
                    return os.path.join(base, name)
        except OSError:
            continue
    return None


def bid_workbooks(rows, chunk_size=1000):
    """Fill the KREAM template; one workbook per 1,000 rows like the GUI."""
    if not rows:
        raise CrawlError("엑셀", "prerequisite", "입찰 엑셀로 내보낼 행이 없습니다. 회원가와 KREAM 사이즈가 매칭된 상품이 필요합니다.")
    template = find_template()
    if not template:
        raise CrawlError("엑셀", "prerequisite", "엑셀 템플릿 k-c46vwlba.xlsx를 프로젝트 폴더에서 찾을 수 없습니다.")
    required = ["상품번호", "옵션명", "판매 희망가", "수량", "입찰기한", "영문 상품명", "모델번호"]
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    chunks = [rows[i:i + chunk_size] for i in range(0, len(rows), chunk_size)]
    files = []
    for index, chunk in enumerate(chunks, 1):
        workbook = openpyxl.load_workbook(template)
        sheet = workbook.active
        header_row, header = None, []
        for r_index, row in enumerate(sheet.iter_rows(min_row=1, max_row=20), 1):
            values = [str(cell.value).strip() if cell.value is not None else "" for cell in row]
            if all(name in values for name in required):
                header_row, header = r_index, [cell.value for cell in row]
                break
        if header_row is None:
            raise CrawlError("엑셀", "schema", "템플릿에서 필수 열을 찾을 수 없습니다.")
        start = header_row + 2
        if sheet.max_row >= start:
            sheet.delete_rows(start, sheet.max_row - start + 1)
        col = {str(name).strip(): i for i, name in enumerate(header) if name}
        for item in chunk:
            line = [None] * len(header)
            for name, value in item.items():
                if name in col:
                    line[col[name]] = value
            sheet.append(line)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        data = io.BytesIO()
        workbook.save(data)
        data.seek(0)
        name = f"output_{stamp}_{index}.xlsx" if len(chunks) > 1 else f"output_{stamp}.xlsx"
        files.append((name, data))
    return files
